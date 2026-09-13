from __future__ import annotations

from datetime import datetime, timezone
from decimal import Decimal
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from uuid import uuid4

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table

from impodo.adapters.artifacts.fallout_workbook import (
    FALLOUT_FILL,
    FalloutWorkbookCell,
    build_fallout_workbook,
)
from impodo.domain.reconciliation import (
    ReconciliationRow,
    ReconciliationRowStatus,
    ReconciliationRun,
    ReconciliationRunStatus,
)


HASH = "sha256:" + "1" * 64
TARGET_HASH = "sha256:" + "2" * 64


def _report() -> ReconciliationRun:
    return ReconciliationRun(
        reconciliation_id=str(uuid4()),
        workspace_id=str(uuid4()),
        execution_run_id=str(uuid4()),
        snapshot_hash=HASH,
        target_hash=TARGET_HASH,
        target_database="target",
        status=ReconciliationRunStatus.FALLOUT,
        verified_at=datetime.now(timezone.utc),
        verified_by="Local operator",
        unchanged_count=0,
        rows=(
            ReconciliationRow(
                row_id="product-1",
                dataset="products",
                source_row=2,
                target_model="product.template",
                operation="CREATE",
                execution_status="COMMITTED",
                status=ReconciliationRowStatus.DIFFERENT,
                odoo_id=42,
                differing_fields=("weight",),
            ),
        ),
    )


class FalloutWorkbookTests(unittest.TestCase):
    def test_derivative_preserves_original_and_marks_exact_cell(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            source = Path(temporary) / "products.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "Products"
            worksheet.append(("Code", "Weight", "Calculation"))
            worksheet.append(("P1", Decimal("0.003"), "=B2*1000"))
            worksheet.add_table(Table(displayName="ProductsTable", ref="A1:C2"))
            workbook.save(source)
            workbook.close()
            original_hash = sha256(source.read_bytes()).hexdigest()

            content = build_fallout_workbook(
                source,
                source_display_name="products.xlsx",
                source_header_row=1,
                report=_report(),
                cells=(
                    FalloutWorkbookCell(
                        source_file="products.xlsx",
                        worksheet="Products",
                        coordinate="B2",
                        business_key="P1",
                        odoo_record="product.template · 42",
                        field="Weight",
                        source_value=None,
                        prepared_value=Decimal("0.003"),
                        odoo_value=0.0,
                        result="Different after Odoo read-back",
                        reason_code="TARGET_NUMERIC_PRECISION_LOSS",
                        recommended_action="Review precision.",
                    ),
                    FalloutWorkbookCell(
                        source_file="products.xlsx",
                        worksheet="",
                        coordinate="",
                        business_key="=FORMULA",
                        odoo_record="product.template · 42",
                        field="Generated value",
                        source_value=None,
                        prepared_value="=CMD()",
                        odoo_value="changed",
                        result="Different after Odoo read-back",
                        reason_code="VALUE_DIFFERENT",
                        recommended_action="Review rule.",
                    ),
                ),
            )

            self.assertEqual(sha256(source.read_bytes()).hexdigest(), original_hash)
            derivative = load_workbook(BytesIO(content), data_only=False)
            self.assertEqual(derivative.sheetnames[0], "Impodo Fallout")
            self.assertEqual(derivative["Products"]["C2"].value, "=B2*1000")
            self.assertIn("ProductsTable", derivative["Products"].tables)
            marked = derivative["Products"]["B2"]
            self.assertEqual(marked.fill.fgColor.rgb[-6:], FALLOUT_FILL)
            self.assertIn("Prepared value: 0.003", marked.comment.text)
            self.assertEqual(
                derivative["Impodo Fallout"]["C8"].hyperlink.target,
                "#'Products'!B2",
            )
            self.assertEqual(derivative["Impodo Fallout"]["D9"].value, "'=FORMULA")
            self.assertEqual(derivative["Impodo Fallout"]["H9"].value, "'=CMD()")
            derivative.close()


if __name__ == "__main__":
    unittest.main()
