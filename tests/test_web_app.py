from __future__ import annotations

from dataclasses import asdict, replace
from contextlib import contextmanager
from datetime import date, datetime, timedelta, timezone
from html import unescape
from io import BytesIO
import json
import multiprocessing
from pathlib import Path
import re
import tempfile
import time
from types import SimpleNamespace
import unittest
from unittest.mock import MagicMock, patch
from urllib.parse import parse_qs, urlsplit
from uuid import uuid4

import duckdb
from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl.worksheet.table import Table

from impodo.access import Actor, ActorIdentity, Capability
from impodo.connectors import (
    ConnectorAuthenticationError,
    ConnectorAuthorizationError,
    ConnectorError,
    ConnectorTransportError,
    MetadataSnapshot,
    RecordSnapshot,
)
from impodo.local_odoo_reader import LocalOdooMetadataReader
from impodo.local_stack import (
    LocalStackCheck,
    LocalStackService,
    LocalStackStartResult,
    LocalStackStatus,
    ReadinessLevel,
)
from impodo.domain.schema.governance import (
    BusinessKeyDefinition,
    BusinessKeyStatus,
    SchemaGovernance,
)
from impodo.domain.mapping.contracts import (
    CategoricalCoveragePolicy,
    DatasetMapping,
    IdentityComponentMapping,
    MappingTargetMode,
    ReferenceKeyMapping,
    RelationshipMapping,
    RelationshipResolver,
    ResolverOrigin,
    ScalarFieldMapping,
    ScalarValueSource,
    TargetFieldHandling,
    ValueMapping,
)
from impodo.domain.mapping.validation.evidence import (
    MappingValidationStatus,
    mapping_issue_fingerprint,
)
from impodo.domain.source_binding import FileSourceBinding
from impodo.data_version_sources import (
    DataVersionSourcePackage,
    SourcePackageCatalog,
    SourcePackageConfiguration,
    SourcePackageDataset,
    SourcePackageFile,
    SourcePackageOrigin,
    SourcePackageState,
    source_column_contract_hash,
)
from impodo.inspection import (
    CATALOG_CONTRACT_VERSION,
    SourceColumnProfile,
    SourceFileCatalog,
    SourceTableCatalog,
)
from impodo.domain.execution import ExecutionRowStatus, ExecutionRunStatus
from impodo.domain.odoo_source_capture import (
    OdooCaptureAccounting,
    OdooCapturePage,
    OdooCaptureValueColumn,
)
from impodo.domain.odoo_comparison import OdooComparisonOutcome
from impodo.application.preflight_service import MANIFEST_NAME
from impodo.application.odoo_connection_service import OdooConnectionTestService
from impodo.application.load_job_service import LoadJobResult
from impodo.application.odoo_read_failures import (
    OdooReadCredentialMissingError,
    OdooReadFailureCode,
    OdooReadWorkflowError,
)
from impodo.models import (
    FieldMetadata,
    ModelMetadata,
    OdooReadIdentity,
    OdooWriteIdentity,
    ProtectedOdooReadContext,
    TargetFingerprint,
    TargetRecord,
    target_identity_hash,
)
from impodo.models import canonical_json_text
from impodo.odoo_readback import ReadbackRecord
from impodo.workspace_state import OdooConnectionMode, WorkspaceStatus, SourceMode
from impodo.preparation_jobs import PreparationJobStatus, PreparationWorkspace
from impodo.planner import PreflightRequirementPlan
from impodo.quality import (
    QualityOutcomePolicy,
    QualityOwnerRole,
    QualityRuleFamily,
)
from impodo.domain.staging.transformation_impact import (
    TransformationImpactReport,
    TransformationImpactRow,
)
from impodo.secrets import MemorySecretStore, SecretStoreError
from impodo.staging_contracts import CanonicalControlTotal
from impodo.web.app import create_local_app
from impodo.web.target_credentials import (
    TargetCredentialRole,
    get_target_credential,
    store_target_credential,
)
from impodo.web.target_readers import _source_value_choices
from impodo.workspace_errors import WorkspaceError
from impodo.workspace_contracts import (
    OdooSchemaCatalog,
    SchemaField,
    SchemaModel,
    SchemaOrigin,
    SourceDataset,
    SourceDatasetColumn,
    SourceSelection,
)
from impodo.domain.odoo_source_policy import ODOO_SOURCE_POLICY_HASH


ROOT = Path(__file__).resolve().parents[1]
POST_HEADERS = {
    "Origin": "http://testserver",
    "Sec-Fetch-Site": "same-origin",
}


def _source_column_profile(
    ordinal: int,
    name: str,
    value: str,
) -> SourceColumnProfile:
    return SourceColumnProfile(
        ordinal=ordinal,
        name=name,
        candidate_type="string",
        null_count=0,
        non_null_count=1,
        distinct_count=1,
        distinct_count_is_exact=True,
        duplicate_count=0,
        minimum=value,
        maximum=value,
        minimum_length=len(value),
        maximum_length=len(value),
    )


def _replace_run_target_setup(
    context,
    workspace_id: str,
    *,
    connection_mode: OdooConnectionMode,
    base_url: str,
    database: str,
    intended_applications: tuple[str, ...] = (),
) -> None:
    workspace = context.migration_workspaces.get(
        workspace_id,
        actor=context.actor,
    )
    current = context.migration_run_target_setup.get(
        workspace.migration_run_id,
        actor=context.actor,
    )
    context.migration_run_target_setup.replace(
        workspace.migration_run_id,
        actor=context.actor,
        expected_revision=current.revision if current is not None else None,
        connection_mode=connection_mode.value,
        base_url=base_url,
        database=database,
        intended_applications=intended_applications,
    )


def _hold_duckdb_files(
    paths: tuple[str, ...],
    ready,
    release,
    status,
) -> None:
    """Own real cross-process DuckDB locks until the test releases them."""

    connections = []
    try:
        connections = [duckdb.connect(path) for path in paths]
        status.put(("ready", ""))
        ready.set()
        release.wait(60)
    except Exception as error:
        status.put(("error", f"{type(error).__name__}: {error}"))
        ready.set()
    finally:
        for connection in reversed(connections):
            connection.close()


@contextmanager
def _spawned_duckdb_locks(*paths: Path):
    """Hold database files from another process using DuckDB's real locks."""

    process_context = multiprocessing.get_context("spawn")
    ready = process_context.Event()
    release = process_context.Event()
    status = process_context.Queue()
    process = process_context.Process(
        target=_hold_duckdb_files,
        args=(tuple(str(path) for path in paths), ready, release, status),
        name="impodo-test-duckdb-lock-holder",
    )
    process.start()
    try:
        if not ready.wait(10):
            raise AssertionError("DuckDB lock holder did not start")
        state, message = status.get(timeout=2)
        if state != "ready":
            raise AssertionError(f"DuckDB lock holder failed: {message}")
        yield
    finally:
        release.set()
        process.join(timeout=10)
        if process.is_alive():
            process.terminate()
            process.join(timeout=5)
        status.close()


def _wait_for_preparation(
    client: TestClient,
    progress_url: str,
    *,
    timeout: float = 20.0,
) -> dict[str, object]:
    deadline = time.monotonic() + timeout
    status_url = f"{progress_url}/status"
    while time.monotonic() < deadline:
        response = client.get(status_url)
        if response.status_code == 200:
            payload = response.json()
            if payload["status"] not in {"QUEUED", "RUNNING"}:
                return payload
        time.sleep(0.05)
    raise AssertionError("background preparation did not finish in time")


def _wait_for_odoo_capture(
    client: TestClient,
    progress_url: str,
    *,
    timeout: float = 5.0,
) -> dict[str, object]:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        response = client.get(f"{progress_url}/status")
        if response.status_code == 200:
            payload = response.json()
            if payload["status"] not in {"QUEUED", "RUNNING"}:
                return payload
        time.sleep(0.02)
    raise AssertionError("background Odoo capture did not finish in time")


def _wait_for_load(
    client: TestClient,
    progress_url: str,
    *,
    timeout: float = 10.0,
) -> dict[str, object]:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        response = client.get(f"{progress_url}/status")
        if response.status_code == 200:
            payload = response.json()
            if payload["status"] not in {"QUEUED", "RUNNING"}:
                return payload
        time.sleep(0.02)
    raise AssertionError("background Odoo load did not finish in time")


def _created_workspace_id(app, response) -> str:
    parts = urlsplit(response.headers["location"]).path.strip("/").split("/")
    if len(parts) >= 2 and parts[0] == "workspaces":
        return parts[1]
    if len(parts) < 2 or parts[0] != "projects":
        raise AssertionError("Project creation returned an unexpected location")
    workspaces = app.state.context.migration_workspaces.list_for_project(
        parts[1],
        actor=app.state.context.actor,
    )
    if len(workspaces) != 1:
        raise AssertionError("New Project did not create one authoring workspace")
    return workspaces[0].workspace_id


def _workspace_data_version_id(context, workspace_id: str) -> str:
    return context.migration_workspaces.get(
        workspace_id,
        actor=context.actor,
    ).data_version_id


def _create_workspace(
    context,
    *,
    name: str,
    source_system: str,
    source_mode: str = "FILE",
):
    return context.project_authoring.create(
        actor=context.actor,
        display_name=name,
        source_mode=source_mode,
        creation_request_id=str(uuid4()),
        source_system_identity=source_system,
    ).workspace_state


class LocalBrowserSecurityTests(unittest.TestCase):
    def setUp(self) -> None:
        (ROOT / ".tmp").mkdir(exist_ok=True)
        self.temporary = tempfile.TemporaryDirectory(dir=ROOT / ".tmp")
        self.project_root = Path(self.temporary.name) / "impodo-data"
        self.app = create_local_app(
            self.project_root,
            launch_token="launch-secret",
            session_secret="session-secret",
            secret_store=MemorySecretStore(),
        )
        self.client = TestClient(self.app)

    def tearDown(self) -> None:
        self.client.close()
        self.temporary.cleanup()

    def test_launch_session_host_and_origin_controls(self) -> None:
        unauthenticated = self.client.get("/projects")
        self.assertEqual(unauthenticated.status_code, 401)

        wrong_host = self.client.get(
            "/launch?token=launch-secret",
            headers={"Host": "attacker.example"},
        )
        self.assertEqual(wrong_host.status_code, 400)

        launched = self.client.get(
            "/launch?token=launch-secret",
            follow_redirects=False,
        )
        self.assertEqual(launched.status_code, 303)
        cookie = launched.headers["set-cookie"].casefold()
        self.assertIn("httponly", cookie)
        self.assertIn("samesite=strict", cookie)
        reused = self.client.get(
            "/launch?token=launch-secret",
            follow_redirects=False,
        )
        self.assertEqual(reused.status_code, 401)

        projects = self.client.get("/projects")
        self.assertEqual(projects.status_code, 200)
        self.assertIn(
            '<span class="brand-tagline">Prepare clean data for Odoo</span>',
            projects.text,
        )
        self.assertIn('id="app-sidebar"', projects.text)
        self.assertIn("data-sidebar-toggle", projects.text)
        self.assertIn('aria-label="Impodo workflow"', projects.text)
        self.assertIn("bootstrap-icons.svg#folder", projects.text)
        self.assertIn("Data remains on this computer.", projects.text)
        self.assertNotIn("Customer data remains on this computer.", projects.text)
        self.assertIn('class="creator-credit"', projects.text)
        self.assertIn("Made in", projects.text)
        self.assertIn("flag-luxembourg.svg", projects.text)
        self.assertEqual(projects.headers["x-frame-options"], "DENY")
        self.assertIn("frame-ancestors 'none'", projects.headers["content-security-policy"])
        self.assertEqual(self.client.get("/projects").status_code, 200)
        self.assertEqual(self.client.get("/projects/new").status_code, 200)

        csrf = _csrf(projects.text)
        missing_origin = self.client.post(
            "/projects/new",
            data={
                "csrf_token": csrf,
                "display_name": "Blocked",
                "source_mode": "FILE",
                "source_system_identity": "Other",
            },
        )
        self.assertEqual(missing_origin.status_code, 403)

        origin_fallback = self.client.post(
            "/projects/new",
            data={
                "csrf_token": csrf,
                "creation_request_id": str(uuid4()),
                "display_name": "Origin fallback",
                "source_mode": "FILE",
                "source_system_identity": "Other",
            },
            headers={"Origin": "http://testserver"},
            follow_redirects=False,
        )
        self.assertEqual(origin_fallback.status_code, 303)

        referer_fallback = self.client.post(
            "/projects/new",
            data={
                "csrf_token": csrf,
                "creation_request_id": str(uuid4()),
                "display_name": "Referer fallback",
                "source_mode": "FILE",
                "source_system_identity": "Other",
            },
            headers={"Referer": "http://testserver/projects/new"},
            follow_redirects=False,
        )
        self.assertEqual(referer_fallback.status_code, 303)

        cross_site = self.client.post(
            "/projects/new",
            data={
                "csrf_token": csrf,
                "display_name": "Cross-site",
                "source_mode": "FILE",
                "source_system_identity": "Other",
            },
            headers={
                "Origin": "http://testserver",
                "Sec-Fetch-Site": "cross-site",
            },
        )
        self.assertEqual(cross_site.status_code, 403)

        hostile_referer = self.client.post(
            "/projects/new",
            data={
                "csrf_token": csrf,
                "display_name": "Hostile",
                "source_mode": "FILE",
                "source_system_identity": "Other",
            },
            headers={"Referer": "http://testserver.attacker.example/projects/new"},
        )
        self.assertEqual(hostile_referer.status_code, 403)


class LocalStackBrowserTests(unittest.TestCase):
    def setUp(self) -> None:
        (ROOT / ".tmp").mkdir(exist_ok=True)
        self.temporary = tempfile.TemporaryDirectory(dir=ROOT / ".tmp")
        self.workspace = Path(self.temporary.name) / "odoo_ve"
        self.config = self.workspace / "config" / "odoo.conf"
        self.config.parent.mkdir(parents=True)
        self.config.write_text(
            "\n".join(
                (
                    "[options]",
                    "http_interface = 127.0.0.1",
                    "http_port = 18069",
                    "db_host = 127.0.0.1",
                    "db_port = 5544",
                    "db_user = odoo",
                    "db_name = odoo19_local",
                    "db_password = postgres-secret",
                    "admin_passwd = master-secret",
                )
            ),
            encoding="utf-8",
        )
        for relative_path in (
            "tools/postgresql/pgsql/bin/pg_isready.exe",
            "tools/postgresql/pgsql/bin/pg_ctl.exe",
            "venv/Scripts/python.exe",
            "odoo/odoo-bin",
        ):
            candidate = self.workspace / relative_path
            candidate.parent.mkdir(parents=True, exist_ok=True)
            candidate.touch()
        (self.workspace / "pgdata").mkdir()
        (self.workspace / "logs").mkdir()
        self.picker_calls = 0
        self.start_calls = 0
        self.stack_running = False
        self.started_processes = []

        def pick_config():
            self.picker_calls += 1
            return self.config

        def probe(profile):
            odoo_ready = self.stack_running
            return LocalStackStatus(
                config_path=str(profile.config_path),
                base_url=profile.base_url,
                database_hint=profile.database_hint,
                checks=(
                    LocalStackCheck(
                        "configuration",
                        "Configuration",
                        ReadinessLevel.READY,
                        "Valid loopback Odoo configuration.",
                    ),
                    LocalStackCheck(
                        "postgresql",
                        "PostgreSQL",
                        ReadinessLevel.READY,
                        "PostgreSQL is accepting connections.",
                    ),
                    LocalStackCheck(
                        "odoo",
                        "Odoo server",
                        (
                            ReadinessLevel.READY
                            if odoo_ready
                            else ReadinessLevel.ACTION
                        ),
                        (
                            "Odoo 19.0 is responding."
                            if odoo_ready
                            else "Odoo is not responding yet."
                        ),
                    ),
                    LocalStackCheck(
                        "api",
                        "Database access (read-only)",
                        ReadinessLevel.UNKNOWN,
                        "Use Save and test connection.",
                    ),
                ),
                profile=profile,
            )

        def starter(profile):
            self.start_calls += 1
            self.stack_running = True
            process = MagicMock()
            process.poll.return_value = None
            process.terminate.side_effect = lambda: setattr(
                self,
                "stack_running",
                False,
            )
            self.started_processes.append(process)
            return LocalStackStartResult(
                status=LocalStackStatus(
                    config_path=str(profile.config_path),
                    base_url=profile.base_url,
                    database_hint=profile.database_hint,
                    checks=(
                        LocalStackCheck(
                            "configuration",
                            "Configuration",
                            ReadinessLevel.READY,
                            "Valid loopback Odoo configuration.",
                        ),
                        LocalStackCheck(
                            "postgresql",
                            "PostgreSQL",
                            ReadinessLevel.READY,
                            "PostgreSQL is accepting connections.",
                        ),
                        LocalStackCheck(
                            "odoo",
                            "Odoo server",
                            ReadinessLevel.READY,
                            "Odoo 19.0 is responding.",
                        ),
                        LocalStackCheck(
                            "api",
                            "Database access (read-only)",
                            ReadinessLevel.UNKNOWN,
                            "Use Save and test connection.",
                        ),
                    ),
                    profile=profile,
                ),
                odoo_process=process,
                postgresql_pid=None,
            )

        local_stack = LocalStackService(
            config_picker=pick_config,
            probe=probe,
            starter=starter,
        )
        self.local_odoo_reader = MagicMock(spec=LocalOdooMetadataReader)
        self.local_odoo_reader.get_target_fingerprint.side_effect = (
            lambda workspace_state, _profile: _browser_schema(workspace_state).fingerprint
        )
        self.project_root = Path(self.temporary.name) / "impodo-data"
        self.app = create_local_app(
            self.project_root,
            launch_token="launch-secret",
            session_secret="session-secret",
            secret_store=MemorySecretStore(),
            local_stack_service=local_stack,
            local_odoo_reader=self.local_odoo_reader,
        )
        self.client = TestClient(self.app)
        launched = self.client.get(
            "/launch?token=launch-secret",
            follow_redirects=False,
        )
        self.assertEqual(launched.status_code, 303)
        projects = self.client.get("/projects")
        self.csrf = _csrf(projects.text)
        created = self.client.post(
            "/projects/new",
            data={
                "csrf_token": self.csrf,
                "creation_request_id": str(uuid4()),
                "display_name": "Local readiness",
                "source_mode": "FILE",
                "source_system_identity": "Other",
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )
        self.workspace_id = _created_workspace_id(self.app, created)
        context = self.app.state.context
        workspace_state = context.queries.get(self.workspace_id)
        context.intake.accept(
            self.workspace_id,
            actor=context.actor,
            expected_revision=workspace_state.revision,
            display_name="local-readiness.csv",
            stream=BytesIO(b"code,name\nP001,Example\n"),
        )
        registered = context.workspace_states.register(
            self.workspace_id,
            actor=context.actor,
            expected_revision=context.queries.get(self.workspace_id).revision,
        )
        self.workspace_revision = registered.revision

    def tearDown(self) -> None:
        self.client.close()
        self.temporary.cleanup()

    def test_selects_config_checks_status_and_keeps_profile_session_only(self) -> None:
        target = self.client.get(f"/workspaces/{self.workspace_id}/target")
        self.assertIn("Find local Odoo", target.text)
        self.assertIn("Choose local Odoo setup", target.text)

        selected = self.client.post(
            f"/workspaces/{self.workspace_id}/local-stack/select-config",
            data={"csrf_token": self.csrf},
            headers=POST_HEADERS,
            follow_redirects=False,
        )
        self.assertEqual(selected.status_code, 303)
        self.assertEqual(
            selected.headers["location"],
            f"/workspaces/{self.workspace_id}/target?local_stack=1",
        )
        self.assertEqual(self.picker_calls, 1)

        refreshed = self.client.get(selected.headers["location"])
        self.assertIn("PostgreSQL is accepting connections", refreshed.text)
        self.assertIn("Odoo is not responding yet", refreshed.text)
        self.assertIn("status-ready", refreshed.text)
        self.assertIn("status-action", refreshed.text)
        self.assertIn('value="http://127.0.0.1:18069"', refreshed.text)
        self.assertIn('value="odoo19_local"', refreshed.text)
        self.assertNotIn("postgres-secret", refreshed.text)
        self.assertNotIn("master-secret", refreshed.text)
        self.assertIn("Start local Odoo", refreshed.text)

        workspace_state = self.app.state.context.workspace_states.repository.get(self.workspace_id)
        self.assertEqual(workspace_state.odoo_base_url, "")
        self.assertEqual(workspace_state.odoo_database, "")
        config_bytes = str(self.config).encode()
        for path in self.app.state.context.workspace_states.repository.workspace_directory(
            self.workspace_id
        ).rglob("*"):
            if path.is_file():
                self.assertNotIn(config_bytes, path.read_bytes())

    def test_start_requires_confirmation_and_updates_readiness(self) -> None:
        self.client.post(
            f"/workspaces/{self.workspace_id}/local-stack/select-config",
            data={"csrf_token": self.csrf},
            headers=POST_HEADERS,
        )

        unconfirmed = self.client.post(
            f"/workspaces/{self.workspace_id}/local-stack/start",
            data={"csrf_token": self.csrf},
            headers=POST_HEADERS,
        )
        self.assertEqual(unconfirmed.status_code, 422)
        self.assertIn("Confirm the detected paths", unconfirmed.text)
        self.assertEqual(self.start_calls, 0)

        started = self.client.post(
            f"/workspaces/{self.workspace_id}/local-stack/start",
            data={
                "csrf_token": self.csrf,
                "confirm_start": "1",
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )
        self.assertEqual(started.status_code, 303)
        self.assertEqual(self.start_calls, 1)
        page = self.client.get(started.headers["location"])
        self.assertIn("Odoo 19.0 is responding", page.text)
        self.assertIn("Control local services started by Impodo", page.text)
        self.assertIn("Stop managed services", page.text)

    def test_local_connection_test_opens_all_green_results(self) -> None:
        self._select_and_start_stack()

        tested = self.client.post(
            f"/workspaces/{self.workspace_id}/target",
            data={
                "csrf_token": self.csrf,
                "revision": str(self.workspace_revision),
                "odoo_connection_mode": "LOCAL",
                "odoo_base_url": "http://127.0.0.1:18069",
                "odoo_database": "odoo19_local",
                "action": "test",
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )

        self.assertEqual(tested.status_code, 303)
        self.assertEqual(
            tested.headers["location"],
            f"/workspaces/{self.workspace_id}/target?local_stack=1",
        )
        results = self.client.get(tested.headers["location"])
        self.assertIn('data-auto-open="true"', results.text)
        self.assertEqual(results.text.count("status-ready"), 4)
        self.assertIn("Odoo data access", results.text)
        self.assertIn("Read-only database access succeeded", results.text)
        self.local_odoo_reader.get_target_fingerprint.assert_called_once()

    def test_local_connection_test_opens_mixed_failure_results(self) -> None:
        self.client.post(
            f"/workspaces/{self.workspace_id}/local-stack/select-config",
            data={"csrf_token": self.csrf},
            headers=POST_HEADERS,
        )

        tested = self.client.post(
            f"/workspaces/{self.workspace_id}/target",
            data={
                "csrf_token": self.csrf,
                "revision": str(self.workspace_revision),
                "odoo_connection_mode": "LOCAL",
                "odoo_base_url": "http://127.0.0.1:18069",
                "odoo_database": "odoo19_local",
                "action": "test",
            },
            headers=POST_HEADERS,
        )

        self.assertEqual(tested.status_code, 422)
        self.assertIn('data-auto-open="true"', tested.text)
        self.assertEqual(tested.text.count("status-ready"), 2)
        self.assertEqual(tested.text.count("status-error"), 2)
        self.assertIn("Local connection checks failed: Odoo server", tested.text)
        self.assertIn("Read-only database access failed", tested.text)
        self.local_odoo_reader.get_target_fingerprint.assert_not_called()

    def test_local_connection_test_marks_database_access_failure(self) -> None:
        self._select_and_start_stack()
        self.local_odoo_reader.get_target_fingerprint.side_effect = ConnectorError(
            "The configured database could not be opened."
        )

        tested = self.client.post(
            f"/workspaces/{self.workspace_id}/target",
            data={
                "csrf_token": self.csrf,
                "revision": str(self.workspace_revision),
                "odoo_connection_mode": "LOCAL",
                "odoo_base_url": "http://127.0.0.1:18069",
                "odoo_database": "odoo19_local",
                "action": "test",
            },
            headers=POST_HEADERS,
        )

        self.assertEqual(tested.status_code, 422)
        self.assertIn('data-auto-open="true"', tested.text)
        self.assertEqual(tested.text.count("status-ready"), 3)
        self.assertEqual(tested.text.count("status-error"), 1)
        self.assertIn("The configured database could not be opened", tested.text)
        self.assertIn("Read-only database access failed", tested.text)

    def test_stop_requires_confirmation_and_stops_only_managed_process(self) -> None:
        self._select_and_start_stack()
        process = self.started_processes[0]

        unconfirmed = self.client.post(
            f"/workspaces/{self.workspace_id}/local-stack/control",
            data={
                "csrf_token": self.csrf,
                "action": "stop",
            },
            headers=POST_HEADERS,
        )
        self.assertEqual(unconfirmed.status_code, 422)
        self.assertIn("Confirm control", unconfirmed.text)
        process.terminate.assert_not_called()

        stopped = self.client.post(
            f"/workspaces/{self.workspace_id}/local-stack/control",
            data={
                "csrf_token": self.csrf,
                "confirm_control": "1",
                "action": "stop",
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )
        self.assertEqual(stopped.status_code, 303)
        process.terminate.assert_called_once_with()
        page = self.client.get(stopped.headers["location"])
        self.assertIn("Start local Odoo", page.text)
        self.assertNotIn("Control local services started by Impodo", page.text)

    def test_restart_stops_the_owned_process_then_starts_a_new_one(self) -> None:
        self._select_and_start_stack()
        first_process = self.started_processes[0]

        restarted = self.client.post(
            f"/workspaces/{self.workspace_id}/local-stack/control",
            data={
                "csrf_token": self.csrf,
                "confirm_control": "1",
                "action": "restart",
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )

        self.assertEqual(restarted.status_code, 303)
        first_process.terminate.assert_called_once_with()
        self.assertEqual(self.start_calls, 2)
        self.assertEqual(len(self.started_processes), 2)
        page = self.client.get(restarted.headers["location"])
        self.assertIn("Control local services started by Impodo", page.text)

    def test_stop_requires_its_explicit_capability(self) -> None:
        self._select_and_start_stack()
        process = self.started_processes[0]
        self.app.state.context.actor = Actor(
            identity=ActorIdentity(
                issuer="https://identity.example.test",
                subject_id="stack-starter",
                display_name="Stack starter",
            ),
            capabilities=frozenset(
                {
                    Capability.PROJECT_VIEW,
                    Capability.PROJECT_EDIT,
                    Capability.LOCAL_STACK_INSPECT,
                    Capability.LOCAL_STACK_START,
                }
            ),
        )

        blocked = self.client.post(
            f"/workspaces/{self.workspace_id}/local-stack/control",
            data={
                "csrf_token": self.csrf,
                "confirm_control": "1",
                "action": "stop",
            },
            headers=POST_HEADERS,
        )

        self.assertEqual(blocked.status_code, 403)
        process.terminate.assert_not_called()

    def test_stop_never_controls_a_stack_started_outside_impodo(self) -> None:
        self.client.post(
            f"/workspaces/{self.workspace_id}/local-stack/select-config",
            data={"csrf_token": self.csrf},
            headers=POST_HEADERS,
        )

        blocked = self.client.post(
            f"/workspaces/{self.workspace_id}/local-stack/control",
            data={
                "csrf_token": self.csrf,
                "confirm_control": "1",
                "action": "stop",
            },
            headers=POST_HEADERS,
        )

        self.assertEqual(blocked.status_code, 422)
        self.assertIn("does not own any running service", blocked.text)
        self.assertEqual(self.start_calls, 0)

    def test_remote_project_cannot_open_local_assistant(self) -> None:
        saved = self.client.post(
            f"/workspaces/{self.workspace_id}/target",
            data={
                "csrf_token": self.csrf,
                "revision": str(self.workspace_revision),
                "odoo_connection_mode": "REMOTE",
                "odoo_base_url": "https://odoo.example.com",
                "odoo_database": "odoo_review",
                "action": "save",
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )
        self.assertEqual(saved.status_code, 422)
        self.assertIn("Check the Odoo connection before continuing", saved.text)

        blocked = self.client.post(
            f"/workspaces/{self.workspace_id}/local-stack/select-config",
            data={"csrf_token": self.csrf},
            headers=POST_HEADERS,
        )
        self.assertEqual(blocked.status_code, 422)
        self.assertIn(
            "available only in Local Odoo mode",
            blocked.text,
        )
        self.assertEqual(self.picker_calls, 0)

    def test_registered_project_reconnects_and_returns_to_comparison(self) -> None:
        self._register_local_project()
        self.stack_running = True

        selected = self.client.post(
            f"/workspaces/{self.workspace_id}/local-stack/select-config",
            data={
                "csrf_token": self.csrf,
                "return_to": "summary_compare",
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )

        self.assertEqual(selected.status_code, 303)
        self.assertEqual(
            selected.headers["location"],
            (
                f"/workspaces/{self.workspace_id}/summary?local_stack=1"
                "#compare-with-odoo"
            ),
        )
        self.assertEqual(self.picker_calls, 1)
        self.local_odoo_reader.get_target_fingerprint.assert_called_once()
        status = self.app.state.context.local_stack.get(self.workspace_id)
        self.assertTrue(status.metadata_ready)

        summary = self.client.get(selected.headers["location"])
        self.assertEqual(summary.status_code, 200)
        self.assertIn('data-auto-open="true"', summary.text)
        self.assertIn("Reconnect local Odoo", summary.text)
        self.assertIn("Continue comparison", summary.text)
        self.assertRegex(
            summary.text,
            r'<button class="button primary" type="submit"\s*>\s*Continue comparison',
        )

    def test_registered_recovery_rejects_another_local_database(self) -> None:
        self._register_local_project(database="another_database")
        self.stack_running = True

        rejected = self.client.post(
            f"/workspaces/{self.workspace_id}/local-stack/select-config",
            data={
                "csrf_token": self.csrf,
                "return_to": "summary_compare",
            },
            headers=POST_HEADERS,
        )

        self.assertEqual(rejected.status_code, 422)
        self.assertIn("Local Odoo is not ready yet", rejected.text)
        self.assertIn("points to database odoo19_local", rejected.text)
        self.assertIn("Choose another local setup", rejected.text)
        self.local_odoo_reader.get_target_fingerprint.assert_not_called()

    def test_local_assistant_requires_its_explicit_capability(self) -> None:
        self.app.state.context.actor = Actor(
            identity=ActorIdentity(
                issuer="https://identity.example.test",
                subject_id="target-editor",
                display_name="Target editor",
            ),
            capabilities=frozenset(
                {
                    Capability.PROJECT_VIEW,
                    Capability.PROJECT_EDIT,
                }
            ),
        )

        blocked = self.client.post(
            f"/workspaces/{self.workspace_id}/local-stack/select-config",
            data={"csrf_token": self.csrf},
            headers=POST_HEADERS,
        )

        self.assertEqual(blocked.status_code, 403)
        self.assertEqual(self.picker_calls, 0)

    def test_start_requires_its_explicit_capability(self) -> None:
        self.client.post(
            f"/workspaces/{self.workspace_id}/local-stack/select-config",
            data={"csrf_token": self.csrf},
            headers=POST_HEADERS,
        )
        self.app.state.context.actor = Actor(
            identity=ActorIdentity(
                issuer="https://identity.example.test",
                subject_id="stack-inspector",
                display_name="Stack inspector",
            ),
            capabilities=frozenset(
                {
                    Capability.PROJECT_VIEW,
                    Capability.PROJECT_EDIT,
                    Capability.LOCAL_STACK_INSPECT,
                }
            ),
        )

        blocked = self.client.post(
            f"/workspaces/{self.workspace_id}/local-stack/start",
            data={
                "csrf_token": self.csrf,
                "confirm_start": "1",
            },
            headers=POST_HEADERS,
        )

        self.assertEqual(blocked.status_code, 403)
        self.assertEqual(self.start_calls, 0)

    def _select_and_start_stack(self) -> None:
        selected = self.client.post(
            f"/workspaces/{self.workspace_id}/local-stack/select-config",
            data={"csrf_token": self.csrf},
            headers=POST_HEADERS,
        )
        self.assertEqual(selected.status_code, 200)
        started = self.client.post(
            f"/workspaces/{self.workspace_id}/local-stack/start",
            data={
                "csrf_token": self.csrf,
                "confirm_start": "1",
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )
        self.assertEqual(started.status_code, 303)

    def _register_local_project(self, *, database: str = "odoo19_local"):
        context = self.app.state.context
        workspace_state = context.workspace_states.repository.get(self.workspace_id)
        now = datetime.now(timezone.utc)
        registered = replace(
            workspace_state,
            odoo_connection_mode=OdooConnectionMode.LOCAL,
            odoo_base_url="http://127.0.0.1:18069",
            odoo_database=database,
            status=WorkspaceStatus.REGISTERED,
            revision=workspace_state.revision + 1,
            updated_at=now,
            registered_at=now,
        )
        context.workspace_states.repository.save(
            registered,
            expected_revision=workspace_state.revision,
            event_type="WORKSPACE_REGISTERED",
            event_detail="",
            actor=context.actor,
        )
        _replace_run_target_setup(
            context,
            self.workspace_id,
            connection_mode=OdooConnectionMode.LOCAL,
            base_url="http://127.0.0.1:18069",
            database=database,
        )
        return registered


class ProjectSetupWizardTests(unittest.TestCase):
    def setUp(self) -> None:
        (ROOT / ".tmp").mkdir(exist_ok=True)
        self.temporary = tempfile.TemporaryDirectory(dir=ROOT / ".tmp")
        self.secrets = MemorySecretStore()
        self.connection_calls: list[tuple[str, str, OdooConnectionMode | None]] = []
        self.schema_calls: list[tuple[str, str]] = []
        self.model_catalog_calls: list[tuple[str, str]] = []
        self.read_identity_calls: list[tuple[str, str, tuple[str, ...]]] = []
        self.readiness_calls = []
        self.local_odoo_reader = MagicMock(spec=LocalOdooMetadataReader)

        def connection_tester(workspace_state, api_key):
            self.connection_calls.append(
                (workspace_state.workspace_id, api_key, workspace_state.odoo_connection_mode)
            )
            return _browser_schema(workspace_state).fingerprint

        def schema_reader(workspace_state, api_key):
            self.schema_calls.append((workspace_state.workspace_id, api_key))
            return _browser_schema(workspace_state)

        def model_catalog_reader(workspace_state, api_key):
            self.model_catalog_calls.append((workspace_state.workspace_id, api_key))
            return _browser_model_catalog(workspace_state)

        def read_identity_probe(workspace_state, api_key, models):
            normalized_models = tuple(sorted(models))
            self.read_identity_calls.append(
                (workspace_state.workspace_id, api_key, normalized_models)
            )
            return OdooReadIdentity(
                target_hash=target_identity_hash(
                    connection_mode=workspace_state.odoo_connection_mode.value,
                    base_url=workspace_state.odoo_base_url,
                    database=workspace_state.odoo_database,
                ),
                principal_hash="sha256:" + "1" * 64,
                permission_hash=(
                    "sha256:"
                    + ("2" if normalized_models == ("ir.model",) else "3") * 64
                ),
                context_hash="sha256:" + "4" * 64,
                readable_models=normalized_models,
                observed_at="2026-08-12T00:00:00Z",
            )

        def readiness_reader(workspace_state, metadata_requests, record_requests):
            self.readiness_calls.append(
                (workspace_state.workspace_id, metadata_requests, record_requests)
            )
            available_metadata = _browser_schema(workspace_state)
            metadata = replace(
                available_metadata,
                models={
                    request.model: replace(
                        available_metadata.models[request.model],
                        fields={
                            field: available_metadata.models[
                                request.model
                            ].fields[field]
                            for field in request.fields
                            if field
                            in available_metadata.models[request.model].fields
                        },
                    )
                    for request in metadata_requests
                    if request.model in available_metadata.models
                },
            )
            records = RecordSnapshot(
                fingerprint=metadata.fingerprint,
                records={item.model: () for item in record_requests},
                requested_fields={
                    item.model: item.fields for item in record_requests
                },
            )
            return metadata, records

        self.app = create_local_app(
            self.temporary.name,
            launch_token="launch-secret",
            session_secret="session-secret",
            secret_store=self.secrets,
            connection_tester=connection_tester,
            read_identity_probe=read_identity_probe,
            schema_reader=schema_reader,
            model_catalog_reader=model_catalog_reader,
            readiness_reader=readiness_reader,
            local_odoo_reader=self.local_odoo_reader,
        )
        self.client = TestClient(self.app)
        launched = self.client.get(
            "/launch?token=launch-secret",
            follow_redirects=False,
        )
        self.assertEqual(launched.status_code, 303)
        self.csrf = _csrf(self.client.get("/projects").text)

    def tearDown(self) -> None:
        self.client.close()
        self.temporary.cleanup()

    def _complete_setup_before_target(self, workspace_id: str):
        """Reach the deferred Odoo destination step for a file project."""

        context = self.app.state.context
        workspace_state = context.queries.get(workspace_id)
        context.intake.accept(
            workspace_id,
            actor=context.actor,
            expected_revision=workspace_state.revision,
            display_name="target-test.csv",
            stream=BytesIO(b"code,name\nP001,Example\n"),
        )
        return context.workspace_states.register(
            workspace_id,
            actor=context.actor,
            expected_revision=context.queries.get(workspace_id).revision,
        )

    def _replace_connection_probes(
        self,
        *,
        fingerprint_probe=None,
        identity_probe=None,
    ) -> None:
        """Keep the shared connection service aligned with replaced test seams."""

        context = self.app.state.context
        if fingerprint_probe is not None:
            context.connection_tester = fingerprint_probe
        if identity_probe is not None:
            context.read_identity_probe = identity_probe
        context.odoo_connection_tests = OdooConnectionTestService(
            context.connection_tester,
            context.read_identity_probe,
        )

    def _registered_remote_schema_workspace(self):
        context = self.app.state.context
        created = _create_workspace(
            context,
            name="Remote summary recovery",
            source_system="Odoo",
            source_mode="ODOO",
        )
        now = datetime.now(timezone.utc)
        workspace_state = replace(
            created,
            source_mode=SourceMode.ODOO,
            odoo_connection_mode=OdooConnectionMode.REMOTE,
            odoo_base_url="https://remote.example.test",
            odoo_database="production",
            intended_models=("res.partner",),
            status=WorkspaceStatus.REGISTERED,
            revision=created.revision + 1,
            updated_at=now,
            registered_at=now,
        )
        context.workspace_states.repository.save(
            workspace_state,
            expected_revision=created.revision,
            event_type="WORKSPACE_REGISTERED",
            event_detail="",
            actor=context.actor,
        )
        _replace_run_target_setup(
            context,
            workspace_state.workspace_id,
            connection_mode=OdooConnectionMode.REMOTE,
            base_url="https://remote.example.test",
            database="production",
        )
        snapshot = _browser_schema(workspace_state)
        model = snapshot.models["res.partner"]
        schema = OdooSchemaCatalog(
            workspace_id=workspace_state.workspace_id,
            policy_hash=ODOO_SOURCE_POLICY_HASH,
            captured_at=now,
            captured_by=context.actor.identity.display_name,
            connection_mode=workspace_state.odoo_connection_mode.value,
            database=workspace_state.odoo_database,
            odoo_version=snapshot.fingerprint.odoo_version,
            models=(
                SchemaModel(
                    model.model,
                    model.description or model.model,
                    tuple(
                        SchemaField(
                            name=field.name,
                            label=field.label or field.name,
                            type=field.type,
                            required=field.required,
                            readonly=field.readonly,
                            relation=field.relation,
                            relation_field=field.relation_field,
                            selection=field.selection,
                            stored=field.stored,
                            computed=field.computed,
                            has_inverse=field.has_inverse,
                            related=field.related,
                            translated=field.translated,
                            company_dependent=field.company_dependent,
                            searchable=field.searchable,
                            sortable=field.sortable,
                            exportable=field.exportable,
                            digits=field.digits,
                            currency_field=field.currency_field,
                        )
                        for _, field in sorted(model.fields.items())
                    ),
                    model.unique_constraints,
                ),
            ),
            content_hash="sha256:" + "5" * 64,
            origin=SchemaOrigin.LIVE_API,
            read_credential_binding_hash="sha256:" + "6" * 64,
            read_principal_hash="sha256:" + "1" * 64,
            read_permission_hash="sha256:" + "3" * 64,
            read_context_hash="sha256:" + "4" * 64,
            connection_target_hash=snapshot.fingerprint.target_hash,
        )
        context.schema_workspace.schemas.save_odoo_schema_catalog(
            workspace_state.workspace_id,
            schema,
            actor=context.actor,
        )
        return workspace_state, schema

    def test_new_project_asks_only_for_name_and_source_mode(self) -> None:
        new_page = self.client.get("/projects/new")
        self.assertIn("Data project name", new_page.text)
        self.assertIn(">Files<", new_page.text)
        self.assertIn(">Data already in Odoo<", new_page.text)
        self.assertNotIn('name="source_system"', new_page.text)
        self.assertNotIn('name="export_date"', new_page.text)
        self.assertNotIn("Data manager", new_page.text)
        self.assertNotIn("Odoo web address", new_page.text)

        created = self._post(
            "/projects/new",
            {
                "csrf_token": self.csrf,
                "display_name": "Customer migration",
                "source_mode": "FILE",
            },
        )
        workspace_id = _created_workspace_id(self.app, created)
        data_project_id = self.app.state.context.migration_workspaces.get(
            workspace_id,
            actor=self.app.state.context.actor,
        ).project_id
        self.assertEqual(
            created.headers["location"],
            f"/projects/{data_project_id}",
        )
        workspace_state = self.app.state.context.queries.get(workspace_id)
        self.assertEqual(workspace_state.source_system, "Uploaded files")
        self.assertEqual(workspace_state.odoo_base_url, "")

        removed_governance = self.client.get(
            f"/workspaces/{workspace_id}/governance",
            follow_redirects=False,
        )
        self.assertEqual(removed_governance.status_code, 404)

    def test_setup_shows_each_blocker_before_the_user_can_move_forward(
        self,
    ) -> None:
        created = self._post(
            "/projects/new",
            {
                "csrf_token": self.csrf,
                "display_name": "Products migration",
                "source_mode": "FILE",
                "source_system_identity": "Dynamics AX 2012",
            },
        )
        workspace_id = _created_workspace_id(self.app, created)
        context = self.app.state.context
        data_project_id = context.migration_workspaces.get(
            workspace_id,
            actor=context.actor,
        ).project_id

        files_page = self.client.get(f"/workspaces/{workspace_id}/files")
        self.assertIn("Add files to this Data version", files_page.text)
        self.assertIn("Add at least one source file.", files_page.text)
        self.assertIn('class="setup-step attention current"', files_page.text)
        self.assertNotIn('class="setup-step locked"', files_page.text)
        self.assertIn('class="sidebar-current-project-link"', files_page.text)
        self.assertIn(
            f'href="/projects/{data_project_id}"',
            files_page.text,
        )
        self.assertIn(
            'aria-label="Open data project overview: Products migration"',
            files_page.text,
        )
        self.assertNotIn("Governance", files_page.text)
        self.assertNotIn("Export date", files_page.text)
        self.assertNotIn("Odoo destination", files_page.text)

        registration_fallback = self._post(
            f"/workspaces/{workspace_id}/register",
            {"csrf_token": self.csrf, "revision": "1"},
        )
        self.assertEqual(registration_fallback.status_code, 422)
        self.assertIn("At least one source file is required", registration_fallback.text)
        self.assertNotIn("Source export", registration_fallback.text)
        self.assertNotIn("Responsible data manager", registration_fallback.text)

        blocked_target = self.client.get(
            f"/workspaces/{workspace_id}/target",
            follow_redirects=False,
        )
        self.assertEqual(blocked_target.status_code, 303)
        self.assertEqual(
            blocked_target.headers["location"],
            f"/workspaces/{workspace_id}/files",
        )

        uploaded = self.client.post(
            f"/workspaces/{workspace_id}/files",
            data={"csrf_token": self.csrf, "revision": "1"},
            files={
                "source_file": (
                    "products.csv",
                    b"code,name\nP001,Example\n",
                    "text/csv",
                )
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )
        self.assertEqual(uploaded.status_code, 303)
        ready_files_page = self.client.get(uploaded.headers["location"])
        self.assertNotIn("Complete this step to continue", ready_files_page.text)
        self.assertIn("Use these files and continue", ready_files_page.text)
        self.assertNotIn("Odoo web address", ready_files_page.text)

    def test_incomplete_file_project_can_be_left_and_resumed(self) -> None:
        created = self._post(
            "/projects/new",
            {
                "csrf_token": self.csrf,
                "display_name": "Waiting for files",
                "source_mode": "FILE",
                "source_system_identity": "Dynamics AX 2012",
            },
        )
        workspace_id = _created_workspace_id(self.app, created)
        context = self.app.state.context
        data_project_id = context.migration_workspaces.get(
            workspace_id,
            actor=context.actor,
        ).project_id

        project_list = self.client.get("/projects")
        self.assertIn("Waiting for files", project_list.text)
        project_page = self.client.get(f"/projects/{data_project_id}")
        self.assertIn("Continue preparing and matching", project_page.text)
        self.assertIn(
            f'href="/workspaces/{workspace_id}/overview"',
            project_page.text,
        )

        workspace_state = self.app.state.context.queries.get(workspace_id)
        self.assertEqual(workspace_state.status, WorkspaceStatus.DRAFT)
        resumed = self.client.get(
            f"/workspaces/{workspace_id}",
            follow_redirects=False,
        )
        self.assertEqual(
            resumed.headers["location"],
            f"/workspaces/{workspace_id}/files",
        )

    def test_source_files_can_change_only_before_table_choices_are_saved(
        self,
    ) -> None:
        context = self.app.state.context
        workspace_state = _create_workspace(
            context,
            name="Correctable source files",
            source_system="CSV",
        )
        kept = context.intake.accept(
            workspace_state.workspace_id,
            actor=context.actor,
            expected_revision=workspace_state.revision,
            display_name="customers.csv",
            stream=BytesIO(b"code,name\nC1,Kept\n"),
        )
        current = context.queries.get(workspace_state.workspace_id)
        wrong = context.intake.accept(
            workspace_state.workspace_id,
            actor=context.actor,
            expected_revision=current.revision,
            display_name="wrong.csv",
            stream=BytesIO(b"code,name\nBAD,Wrong\n"),
        )
        current = context.queries.get(workspace_state.workspace_id)
        files_page = self.client.get(f"/workspaces/{workspace_state.workspace_id}/files")
        self.assertEqual(
            files_page.text.count("data-source-file-remove-form"),
            2,
        )
        self.assertIn("data-source-file-remove-dialog", files_page.text)

        wrong_path = (
            context.workspace_states.repository.workspace_directory(workspace_state.workspace_id)
            / "inbox"
            / wrong.stored_name
        )
        removed_draft = self._post(
            f"/workspaces/{workspace_state.workspace_id}/files/{wrong.file_id}/remove",
            {
                "csrf_token": self.csrf,
                "revision": str(current.revision),
                "return_to": "files",
            },
        )
        self.assertEqual(removed_draft.status_code, 303)
        self.assertFalse(wrong_path.exists())
        current = context.queries.get(workspace_state.workspace_id)
        registered = context.workspace_states.register(
            workspace_state.workspace_id,
            actor=context.actor,
            expected_revision=current.revision,
        )

        source_page = self.client.get(f"/workspaces/{workspace_state.workspace_id}/sources")
        self.assertEqual(source_page.status_code, 200)
        self.assertEqual(
            source_page.text.count("data-source-file-remove-form"),
            1,
        )
        self.assertIn(
            f'action="/workspaces/{workspace_state.workspace_id}/sources/files"',
            source_page.text,
        )
        replacement_upload = self.client.post(
            f"/workspaces/{workspace_state.workspace_id}/sources/files",
            data={
                "csrf_token": self.csrf,
                "revision": str(registered.revision),
            },
            files={
                "source_file": (
                    "corrected.csv",
                    b"code,name\nC2,Corrected\n",
                    "text/csv",
                )
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )
        self.assertEqual(replacement_upload.status_code, 303)
        current = context.queries.get(workspace_state.workspace_id)
        corrected = next(
            item
            for item in current.source_files
            if item.display_name == "corrected.csv"
        )
        removed_registered = self._post(
            f"/workspaces/{workspace_state.workspace_id}/files/{corrected.file_id}/remove",
            {
                "csrf_token": self.csrf,
                "revision": str(current.revision),
                "return_to": "sources",
            },
        )
        self.assertEqual(removed_registered.status_code, 303)
        self.assertEqual(
            removed_registered.headers["location"],
            f"/workspaces/{workspace_state.workspace_id}/sources#source-files",
        )
        removed_page = self.client.get(removed_registered.headers["location"])
        self.assertIn(
            "Removed corrected.csv from this Data version.",
            removed_page.text,
        )
        datasets_page = self.client.get(f"/workspaces/{workspace_state.workspace_id}/datasets")
        self.assertEqual(
            datasets_page.text.count("data-source-file-remove-form"),
            1,
        )

        current = context.queries.get(workspace_state.workspace_id)
        now = datetime.now(timezone.utc)
        context.sources.sources.save_source_selection(
            workspace_state.workspace_id,
            SourceSelection(
                selection_id=str(uuid4()),
                version=1,
                data_version_id=_workspace_data_version_id(
                    context,
                    workspace_state.workspace_id,
                ),
                created_at=now,
                created_by=context.actor.identity.display_name,
                datasets=(),
                content_hash="sha256:" + "a" * 64,
            ),
            actor=context.actor,
        )
        blocked = self._post(
            f"/workspaces/{workspace_state.workspace_id}/files/{kept.file_id}/remove",
            {
                "csrf_token": self.csrf,
                "revision": str(current.revision),
                "return_to": "datasets",
            },
        )
        self.assertEqual(blocked.status_code, 422)
        self.assertIn(
            "Source files cannot be changed after table choices are saved",
            blocked.text,
        )
        self.assertNotIn("data-source-file-remove-form", blocked.text)

    def test_remote_connection_status_is_visible_persistent_and_target_bound(
        self,
    ) -> None:
        created = self._post(
            "/projects/new",
            {
                "csrf_token": self.csrf,
                "display_name": "Remote connection feedback",
                "source_mode": "FILE",
                "source_system_identity": "Other",
            },
        )
        workspace_id = _created_workspace_id(self.app, created)
        workspace_state = self._complete_setup_before_target(workspace_id)
        target_form = self.client.get(f"/workspaces/{workspace_id}/target")
        self.assertIn('name="read_api_key"', target_form.text)
        self.assertIn('name="read_api_key_storage"', target_form.text)
        self.assertIn('name="keep_api_key_for_loading"', target_form.text)
        self.assertIn("Use this key for checking and loading", target_form.text)
        self.assertIn('value="vault"', target_form.text)
        self.assertIn("Not available", target_form.text)
        self.assertNotIn('name="write_api_key"', target_form.text)

        tested = self.client.post(
            f"/workspaces/{workspace_id}/target",
            data={
                "csrf_token": self.csrf,
                "revision": str(workspace_state.revision),
                "odoo_connection_mode": "REMOTE",
                "odoo_base_url": "https://edu-ucaps.odoo.com",
                "odoo_database": "edu-ucaps",
                "read_api_key": "remote-secret-key",
                "action": "test",
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )

        self.assertEqual(tested.status_code, 303)
        self.assertEqual(
            tested.headers["location"],
            f"/workspaces/{workspace_id}/target#remote-connection-status",
        )
        result = self.client.get(tested.headers["location"])
        self.assertIn("connection-state-ready", result.text)
        self.assertIn("The Odoo connection is ready.", result.text)
        self.assertIn("Read-only access to edu-ucaps succeeded.", result.text)
        self.assertIn("Supported Odoo version 19.0.", result.text)
        self.assertIn(
            "The authenticated read-only Odoo principal was identified.",
            result.text,
        )
        self.assertIn("Checked during this Impodo session.", result.text)
        self.assertIn("Available this session", result.text)
        self.assertIn("Impodo will forget this key when Impodo closes.", result.text)
        self.assertIn("Forget checking key", result.text)
        self.assertIn(">Check again</button>", result.text)
        self.assertRegex(result.text, r"data-local-stack-entry\s+hidden")
        self.assertNotIn("remote-secret-key", result.text)
        self.assertEqual(
            self.read_identity_calls,
            [(workspace_id, "remote-secret-key", ("res.users",))],
        )
        workspace_state = self.app.state.context.workspace_states.repository.get(workspace_id)
        self.assertEqual(
            get_target_credential(
                self.secrets,
                workspace_state,
                TargetCredentialRole.READ,
            ).secret,
            "remote-secret-key",
        )
        self.assertIsNone(
            get_target_credential(
                self.secrets,
                workspace_state,
                TargetCredentialRole.WRITE,
            )
        )

        refreshed = self.client.get(f"/workspaces/{workspace_id}/target")
        self.assertIn("The Odoo connection is ready.", refreshed.text)
        self.assertEqual(len(self.connection_calls), 1)

        forgotten = self.client.post(
            f"/workspaces/{workspace_id}/target/read-credential/delete",
            data={"csrf_token": self.csrf},
            headers=POST_HEADERS,
            follow_redirects=False,
        )
        self.assertEqual(forgotten.status_code, 303)
        forgotten_page = self.client.get(forgotten.headers["location"])
        self.assertIn("The read-only Odoo key was forgotten.", forgotten_page.text)
        self.assertIn("Not available", forgotten_page.text)
        self.assertNotIn("The Odoo connection is ready.", forgotten_page.text)
        self.assertEqual(self.secrets.values, {})

        changed = self.client.post(
            f"/workspaces/{workspace_id}/target",
            data={
                "csrf_token": self.csrf,
                "revision": str(workspace_state.revision),
                "odoo_connection_mode": "REMOTE",
                "odoo_base_url": "https://other.example.com",
                "odoo_database": "other_database",
                "action": "save",
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )
        self.assertEqual(changed.status_code, 422)
        changed_target = changed
        self.assertIn("connection-state-unknown", changed_target.text)
        self.assertIn(
            "The Odoo connection has not been checked.",
            changed_target.text,
        )
        self.assertNotIn("Read-only access to edu-ucaps succeeded.", changed_target.text)
        self.assertEqual(self.secrets.values, {})

        script = self.client.get("/static/app.js")
        self.assertIn("resetRemoteConnectionStatus", script.text)
        self.assertIn('window.location.hash === "#remote-connection-status"', script.text)
        styles = self.client.get("/static/app.css")
        self.assertIn("[data-local-stack-entry][hidden]", styles.text)

    def test_stage_two_can_keep_the_checking_key_for_loading(self) -> None:
        context = self.app.state.context
        created = _create_workspace(
            context,
            name="One approved Odoo key",
            source_system="Other",
        )
        now = datetime.now(timezone.utc)
        workspace_state = replace(
            created,
            status=WorkspaceStatus.REGISTERED,
            revision=created.revision + 1,
            updated_at=now,
            registered_at=now,
        )
        context.workspace_states.repository.save(
            workspace_state,
            expected_revision=created.revision,
            event_type="WORKSPACE_REGISTERED",
            event_detail="",
            actor=context.actor,
        )
        workspace_id = workspace_state.workspace_id

        tested = self.client.post(
            f"/workspaces/{workspace_id}/target",
            data={
                "csrf_token": self.csrf,
                "revision": str(workspace_state.revision),
                "odoo_connection_mode": "REMOTE",
                "odoo_base_url": "https://edu-ucaps.odoo.com",
                "odoo_database": "edu-ucaps",
                "read_api_key": "approved-check-and-load-key",
                "read_api_key_storage": "vault",
                "keep_api_key_for_loading": "1",
                "action": "test",
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )

        self.assertEqual(tested.status_code, 303, tested.text)
        workspace_state = self.app.state.context.workspace_states.repository.get(
            workspace_id
        )
        read_credential = get_target_credential(
            self.secrets,
            workspace_state,
            TargetCredentialRole.READ,
        )
        write_credential = get_target_credential(
            self.secrets,
            workspace_state,
            TargetCredentialRole.WRITE,
        )
        assert read_credential is not None
        assert write_credential is not None
        self.assertEqual(read_credential.secret, "approved-check-and-load-key")
        self.assertEqual(write_credential.secret, "approved-check-and-load-key")
        self.assertTrue(read_credential.persistent)
        self.assertTrue(write_credential.persistent)
        repository = self.app.state.context.workspace_states.repository
        self.assertTrue(
            repository.has_audit_event(
                workspace_id,
                "ODOO_READ_CREDENTIAL_STORED",
            )
        )
        self.assertTrue(
            repository.has_audit_event(
                workspace_id,
                "ODOO_WRITE_CREDENTIAL_STORED",
            )
        )

        status_page = self.client.get(tested.headers["location"])
        self.assertIn("Checking key: Saved on this computer", status_page.text)
        self.assertIn(
            "Loading key:</strong> Saved on this computer",
            status_page.text,
        )
        self.assertIn(
            "Stage 6 without asking you to enter it again",
            status_page.text,
        )
        self.assertIn("Forget loading key", status_page.text)
        self.assertNotIn("approved-check-and-load-key", status_page.text)

        forgotten = self.client.post(
            f"/workspaces/{workspace_id}/target/write-credential/delete",
            data={"csrf_token": self.csrf},
            headers=POST_HEADERS,
            follow_redirects=False,
        )
        self.assertEqual(forgotten.status_code, 303)
        self.assertIsNone(
            get_target_credential(
                self.secrets,
                workspace_state,
                TargetCredentialRole.WRITE,
            )
        )
        self.assertIsNotNone(
            get_target_credential(
                self.secrets,
                workspace_state,
                TargetCredentialRole.READ,
            )
        )
        forgotten_page = self.client.get(forgotten.headers["location"])
        self.assertIn("The Odoo loading key was forgotten.", forgotten_page.text)

        kept_again = self.client.post(
            f"/workspaces/{workspace_id}/target",
            data={
                "csrf_token": self.csrf,
                "revision": str(workspace_state.revision),
                "odoo_connection_mode": "REMOTE",
                "odoo_base_url": "https://edu-ucaps.odoo.com",
                "odoo_database": "edu-ucaps",
                "keep_api_key_for_loading": "1",
                "action": "test",
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )
        self.assertEqual(kept_again.status_code, 303, kept_again.text)
        current = self.app.state.context.workspace_states.repository.get(workspace_id)
        restored_write = get_target_credential(
            self.secrets,
            current,
            TargetCredentialRole.WRITE,
        )
        assert restored_write is not None
        self.assertEqual(restored_write.secret, "approved-check-and-load-key")
        self.assertTrue(restored_write.persistent)

    def test_summary_reconnects_missing_remote_key_without_losing_schema(
        self,
    ) -> None:
        context = self.app.state.context
        workspace_state, schema = self._registered_remote_schema_workspace()

        def compare_with_reader(_project_id, *, reader, actor):
            self.assertEqual(actor, context.actor)
            return reader(
                PreflightRequirementPlan((), (), (), source_record_count=0)
            )

        with patch.object(
            context.preflight,
            "compare",
            side_effect=compare_with_reader,
        ):
            blocked = self.client.post(
                f"/workspaces/{workspace_state.workspace_id}/summary/compare",
                data={"csrf_token": self.csrf},
                headers=POST_HEADERS,
            )

        self.assertEqual(blocked.status_code, 422)
        self.assertIn("Enter the Odoo read key", blocked.text)
        self.assertIn('name="read_api_key"', blocked.text)
        self.assertIn('name="read_api_key_storage"', blocked.text)
        self.assertIn("Save key and compare", blocked.text)
        self.assertIn("Nothing was changed in Odoo", blocked.text)
        self.assertNotIn("We could not complete that action", blocked.text)
        self.assertEqual(self.schema_calls, [])
        self.assertEqual(self.read_identity_calls, [])

        with patch.object(context.preflight, "compare") as compare:
            reconnected = self.client.post(
                f"/workspaces/{workspace_state.workspace_id}/summary/compare",
                data={
                    "csrf_token": self.csrf,
                    "read_api_key": "replacement-read-key",
                    "read_api_key_storage": "session",
                },
                headers=POST_HEADERS,
                follow_redirects=False,
            )

        self.assertEqual(reconnected.status_code, 303, reconnected.text)
        compare.assert_called_once()
        self.assertEqual(
            self.schema_calls,
            [(workspace_state.workspace_id, "replacement-read-key")],
        )
        self.assertEqual(
            self.read_identity_calls,
            [
                (
                    workspace_state.workspace_id,
                    "replacement-read-key",
                    ("res.partner",),
                )
            ],
        )
        rebound = context.queries.get_odoo_schema_catalog(workspace_state.workspace_id)
        self.assertEqual(rebound.content_hash, schema.content_hash)
        self.assertNotEqual(
            rebound.read_credential_binding_hash,
            schema.read_credential_binding_hash,
        )
        page = self.client.get(reconnected.headers["location"])
        self.assertNotIn("replacement-read-key", page.text)
        self.assertNotIn('id="comparison-recovery"', page.text)

    def test_remote_reference_failure_returns_to_matching_without_key_form(
        self,
    ) -> None:
        context = self.app.state.context
        workspace_state, schema = self._registered_remote_schema_workspace()
        available_key = SimpleNamespace(
            available=True,
            binding_hash=schema.read_credential_binding_hash,
        )

        with (
            patch.object(
                context.preflight,
                "compare",
                side_effect=OdooReadWorkflowError(
                    OdooReadFailureCode.REFERENCE_POLICY_MISMATCH,
                    "Internal wording must not select or leak into recovery",
                    support_reference="res.partner.country_id -> res.country",
                ),
            ),
            patch(
                "impodo.web.presenters.summary.get_target_credential_status",
                return_value=available_key,
            ),
        ):
            blocked = self.client.post(
                f"/workspaces/{workspace_state.workspace_id}/summary/compare",
                data={"csrf_token": self.csrf},
                headers=POST_HEADERS,
            )

        self.assertEqual(blocked.status_code, 422)
        self.assertIn('id="comparison-recovery"', blocked.text)
        self.assertIn("Review the linked field match", blocked.text)
        self.assertIn(
            f'href="/workspaces/{workspace_state.workspace_id}/mapping"',
            blocked.text,
        )
        self.assertNotIn('name="read_api_key"', blocked.text)
        self.assertNotIn("Internal wording", blocked.text)

    def test_remote_compare_recovery_matches_the_classified_failure(self) -> None:
        context = self.app.state.context
        workspace_state, schema = self._registered_remote_schema_workspace()
        available_key = SimpleNamespace(
            available=True,
            binding_hash=schema.read_credential_binding_hash,
        )
        cases = (
            (
                ConnectorAuthenticationError("raw rejected key response"),
                "Replace the Odoo read key",
                'name="read_api_key"',
            ),
            (
                ConnectorTransportError("HTTP 503 raw upstream response"),
                "Try the comparison again",
                f'action="/workspaces/{workspace_state.workspace_id}/summary/compare"',
            ),
            (
                OdooReadWorkflowError(
                    OdooReadFailureCode.SCHEMA_EVIDENCE_STALE,
                    "raw schema mismatch",
                ),
                "Refresh the Odoo data structure",
                f'href="/workspaces/{workspace_state.workspace_id}/schema"',
            ),
            (
                OdooReadWorkflowError(
                    OdooReadFailureCode.PREPARED_EVIDENCE_STALE,
                    "raw prepared mismatch",
                ),
                "Prepare and approve the data again",
                f'href="/workspaces/{workspace_state.workspace_id}/prepare"',
            ),
            (
                SecretStoreError("raw Windows credential-store failure"),
                "Impodo could not save the comparison",
                "COMPARISON_STORAGE_FAILED",
            ),
        )

        for failure, title, expected_action in cases:
            with self.subTest(failure=type(failure).__name__):
                with (
                    patch.object(
                        context.preflight,
                        "compare",
                        side_effect=failure,
                    ),
                    patch(
                        "impodo.web.presenters.summary.get_target_credential_status",
                        return_value=available_key,
                    ),
                ):
                    blocked = self.client.post(
                        f"/workspaces/{workspace_state.workspace_id}/summary/compare",
                        data={"csrf_token": self.csrf},
                        headers=POST_HEADERS,
                    )

                self.assertEqual(blocked.status_code, 422)
                self.assertIn(title, blocked.text)
                self.assertIn(expected_action, blocked.text)
                if not isinstance(failure, ConnectorAuthenticationError):
                    self.assertNotIn('name="read_api_key"', blocked.text)
                self.assertNotIn("raw ", blocked.text)

    def test_remote_connection_failure_shows_safe_red_checks(self) -> None:
        def rejected_connection(_project, _api_key):
            raise ConnectorAuthenticationError(
                "raw remote response and secret must not be displayed"
            )

        self._replace_connection_probes(
            fingerprint_probe=rejected_connection,
        )
        created = self._post(
            "/projects/new",
            {
                "csrf_token": self.csrf,
                "display_name": "Rejected remote connection",
                "source_mode": "FILE",
                "source_system_identity": "Other",
            },
        )
        workspace_id = _created_workspace_id(self.app, created)
        workspace_state = self._complete_setup_before_target(workspace_id)

        tested = self.client.post(
            f"/workspaces/{workspace_id}/target",
            data={
                "csrf_token": self.csrf,
                "revision": str(workspace_state.revision),
                "odoo_connection_mode": "REMOTE",
                "odoo_base_url": "https://odoo.example.com",
                "odoo_database": "migration",
                "api_key": "never-render-this-key",
                "action": "test",
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )

        self.assertEqual(tested.status_code, 303)
        result = self.client.get(tested.headers["location"])
        self.assertIn("connection-state-error", result.text)
        self.assertIn("The Odoo connection is not ready.", result.text)
        self.assertIn("Odoo responded to the read-only check.", result.text)
        self.assertIn(
            "Odoo rejected the access key, database name, or API entitlement.",
            result.text,
        )
        self.assertIn("ODOO_READ_KEY_REJECTED", result.text)
        self.assertIn(">Try again</button>", result.text)
        self.assertNotIn("never-render-this-key", result.text)
        self.assertNotIn("raw remote response", result.text)

    def test_remote_connection_reports_missing_principal_model_access(self) -> None:
        def denied_identity(_project, _api_key, _models):
            raise ConnectorAuthorizationError(
                "internal group and model details must not be rendered"
            )

        self._replace_connection_probes(identity_probe=denied_identity)
        created = self._post(
            "/projects/new",
            {
                "csrf_token": self.csrf,
                "display_name": "Read principal permission",
                "source_mode": "FILE",
                "source_system_identity": "Other",
            },
        )
        workspace_id = _created_workspace_id(self.app, created)
        workspace_state = self._complete_setup_before_target(workspace_id)

        tested = self.client.post(
            f"/workspaces/{workspace_id}/target",
            data={
                "csrf_token": self.csrf,
                "revision": str(workspace_state.revision),
                "odoo_connection_mode": "REMOTE",
                "odoo_base_url": "https://odoo.example.com",
                "odoo_database": "migration",
                "read_api_key": "never-render-this-key",
                "action": "test",
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )

        self.assertEqual(tested.status_code, 303)
        result = self.client.get(tested.headers["location"])
        self.assertIn("ODOO_READ_ACCESS_MISSING", result.text)
        self.assertIn(
            "The authenticated principal lacks the required read access.",
            result.text,
        )
        self.assertNotIn("never-render-this-key", result.text)
        self.assertNotIn("internal group", result.text)

    def test_remote_connection_distinguishes_api_version_and_network_failures(
        self,
    ) -> None:
        def wrong_version(workspace_state, _api_key):
            return replace(
                _browser_schema(workspace_state).fingerprint,
                odoo_version="18.0",
            )

        def missing_api(_project, _api_key):
            raise ConnectorTransportError("Odoo JSON-2 read failed with HTTP 404")

        def unreachable(_project, _api_key):
            raise ConnectorTransportError(
                "Odoo JSON-2 read timed out or was unreachable"
            )

        cases = (
            (
                wrong_version,
                "Impodo requires Odoo 19; this target reported Odoo 18.0.",
                "ODOO_VERSION_UNSUPPORTED",
            ),
            (
                missing_api,
                "The JSON-2 API was not available at this address.",
                "ODOO_API_HTTP_404",
            ),
            (
                unreachable,
                "Impodo could not reach Odoo. Check the address and network connection.",
                "ODOO_TARGET_UNREACHABLE",
            ),
        )

        for index, (tester, message, support_code) in enumerate(cases, start=1):
            with self.subTest(support_code=support_code):
                self._replace_connection_probes(fingerprint_probe=tester)
                created = self._post(
                    "/projects/new",
                    {
                        "csrf_token": self.csrf,
                        "display_name": f"Remote failure {index}",
                        "source_mode": "FILE",
                        "source_system_identity": "Other",
                    },
                )
                workspace_id = _created_workspace_id(self.app, created)
                workspace_state = self._complete_setup_before_target(workspace_id)
                tested = self.client.post(
                    f"/workspaces/{workspace_id}/target",
                    data={
                        "csrf_token": self.csrf,
                        "revision": str(workspace_state.revision),
                        "odoo_connection_mode": "REMOTE",
                        "odoo_base_url": "https://odoo.example.com",
                        "odoo_database": f"migration_{index}",
                        "api_key": f"secret-{index}",
                        "action": "test",
                    },
                    headers=POST_HEADERS,
                    follow_redirects=False,
                )
                self.assertEqual(tested.status_code, 303)
                result = self.client.get(tested.headers["location"])
                self.assertIn("connection-state-error", result.text)
                self.assertIn(message, result.text)
                self.assertIn(support_code, result.text)
                self.assertNotIn(f"secret-{index}", result.text)

    def test_incompatible_project_explains_recovery(
        self,
    ) -> None:
        context = self.app.state.context
        workspace_state = _create_workspace(
            context,
            name="Historical rehearsal",
            source_system="Other",
        )
        repository = context.workspace_states.repository
        workspace_dir = repository.workspace_directory(workspace_state.workspace_id)
        database_path = workspace_dir / "workspace-engine.duckdb"
        with repository._connect(database_path) as connection:
            connection.execute("DROP TABLE schema_version")
            connection.execute(
                "CREATE TABLE schema_version (version INTEGER NOT NULL)"
            )
            connection.execute("INSERT INTO schema_version VALUES (1)")

        opened = self.client.get(f"/workspaces/{workspace_state.workspace_id}")
        self.assertEqual(opened.status_code, 409)
        self.assertIn("uses a saved-data generation or version", opened.text)
        self.assertNotIn("Traceback", opened.text)

    def test_load_receipt_rows_offer_twenty_or_fifty_with_pagination(self) -> None:
        context = self.app.state.context
        workspace_state = _create_workspace(
            context,
            name="Paginated load review",
            source_system="Other",
        )
        workspace_state = context.workspace_states.update_target(
            workspace_state.workspace_id,
            actor=context.actor,
            expected_revision=workspace_state.revision,
            odoo_connection_mode="REMOTE",
            odoo_base_url="https://odoo.example.test",
            odoo_database="migration",
            intended_applications=("Contacts",),
            intended_models=(),
        )
        status = SimpleNamespace(value="COMMITTED")
        rows = tuple(
            SimpleNamespace(
                dataset="contacts",
                source_row=index,
                target_model="res.partner",
                odoo_id=index,
                operation="CREATE",
                status=status,
                safe_error="",
            )
            for index in range(1, 56)
        )
        current_run = SimpleNamespace(
            rows=rows,
            committed_count=55,
            failed_count=0,
            blocked_count=0,
            partially_applied_count=0,
            unknown_count=0,
            run_id="run-1",
        )
        preview = SimpleNamespace(
            snapshot=SimpleNamespace(
                counts={"CREATE": 55, "UPDATE": 0, "UNCHANGED": 0},
                target_database="migration",
                target_odoo_version="19.0",
                semantic_hash="sha256:" + "a" * 64,
                target_hash="sha256:" + "b" * 64,
            ),
            datasets=(),
            current_run=current_run,
            can_load=False,
        )

        with (
            patch.object(
                type(context.execution),
                "current_preview",
                return_value=preview,
            ),
            patch.object(
                type(context.reconciliation),
                "current",
                return_value=None,
            ),
        ):
            page = self.client.get(
                f"/workspaces/{workspace_state.workspace_id}/load"
                "?rows_page=2&rows_per_page=20"
            )

        self.assertEqual(page.status_code, 200)
        self.assertEqual(page.text.count("data-load-row"), 20)
        self.assertIn("Showing 21-40 of 55 records", page.text)
        self.assertIn("Rows per page:", page.text)
        self.assertIn(">20</a>", page.text)
        self.assertIn(">50</a>", page.text)
        self.assertIn("Page 2 of 3", page.text)
        self.assertIn("row 21", page.text)
        self.assertIn("row 40", page.text)
        self.assertNotIn("row 41", page.text)

    def test_load_review_separates_read_only_review_from_confirmation(self) -> None:
        context = self.app.state.context
        workspace_state = _create_workspace(
            context,
            name="Staged load review",
            source_system="Other",
        )
        workspace_state = context.workspace_states.update_target(
            workspace_state.workspace_id,
            actor=context.actor,
            expected_revision=workspace_state.revision,
            odoo_connection_mode="REMOTE",
            odoo_base_url="https://odoo.example.test",
            odoo_database="migration",
            intended_applications=("Contacts",),
            intended_models=(),
        )
        preview = SimpleNamespace(
            snapshot=SimpleNamespace(
                counts={"CREATE": 2, "UPDATE": 1, "UNCHANGED": 4},
                write_count=3,
                target_database="migration",
                target_odoo_version="19.0",
                semantic_hash="sha256:" + "a" * 64,
                target_hash="sha256:" + "b" * 64,
            ),
            datasets=(),
            current_run=None,
            can_load=True,
            scope_error="",
            deferred_create_count=0,
        )

        with (
            patch.object(
                type(context.execution),
                "current_preview",
                return_value=preview,
            ),
            patch.object(
                type(context.reconciliation),
                "current",
                return_value=None,
            ),
        ):
            review = self.client.get(
                f"/workspaces/{workspace_state.workspace_id}/load/review"
            )
            confirm = self.client.get(
                f"/workspaces/{workspace_state.workspace_id}/load/confirm"
            )
            outcome = self.client.get(
                f"/workspaces/{workspace_state.workspace_id}/load/outcome",
                follow_redirects=False,
            )
            preview.can_load = False
            preview.scope_error = "One reviewed field is no longer available."
            blocked_confirm = self.client.get(
                f"/workspaces/{workspace_state.workspace_id}/load/confirm",
                follow_redirects=False,
            )

        self.assertEqual(review.status_code, 200)
        self.assertIn("Check what will change in Odoo", review.text)
        self.assertIn("Nothing is written from this page", review.text)
        self.assertNotIn('name="write_api_key"', review.text)
        self.assertEqual(confirm.status_code, 200)
        self.assertIn("Load 3 records into migration", confirm.text)
        self.assertIn('name="write_api_key"', confirm.text)
        self.assertIn("Advanced settings", confirm.text)
        self.assertIn("Load 3 records into Odoo", confirm.text)
        self.assertEqual(outcome.status_code, 303)
        self.assertEqual(
            outcome.headers["location"],
            f"/workspaces/{workspace_state.workspace_id}/load/review",
        )
        self.assertEqual(blocked_confirm.status_code, 303)
        self.assertEqual(
            blocked_confirm.headers["location"],
            f"/workspaces/{workspace_state.workspace_id}/load/review",
        )

    def test_load_progress_names_target_and_exposes_non_secret_counts(self) -> None:
        context = self.app.state.context
        workspace_state = _create_workspace(
            context,
            name="Visible load progress",
            source_system="Other",
        )
        workspace_state = context.workspace_states.update_target(
            workspace_state.workspace_id,
            actor=context.actor,
            expected_revision=workspace_state.revision,
            odoo_connection_mode="REMOTE",
            odoo_base_url="https://odoo.example.test",
            odoo_database="migration",
            intended_applications=("Contacts",),
            intended_models=(),
        )
        manager = context.load_jobs
        assert manager is not None
        job = manager.enqueue(
            workspace_state.workspace_id,
            workspace_state.name,
            target_database="migration",
            target_server="odoo.example.test",
            target_environment="Test",
            total_rows=3,
            access_context=context.workspace_access.resolve(
                workspace_state.workspace_id,
                actor=context.actor,
                capability=Capability.EXPORT_PLAN_EXECUTE,
            ),
            work=lambda _access, _writing, _verifying: LoadJobResult(
                execution_run_id="run-1",
                verification_complete=False,
            ),
        )
        progress_url = f"/workspaces/{workspace_state.workspace_id}/load/progress/{job.job_id}"

        finished = _wait_for_load(self.client, progress_url)
        page = self.client.get(progress_url)

        self.assertEqual(page.status_code, 200, page.text)
        self.assertIn("data-load-job", page.text)
        self.assertIn(f"Loading {workspace_state.name} into Test Odoo", page.text)
        self.assertIn("odoo.example.test", page.text)
        self.assertIn("migration", page.text)
        self.assertEqual(finished["total_rows"], 3)
        self.assertEqual(finished["created_count"], 0)
        self.assertEqual(finished["updated_count"], 0)
        self.assertEqual(finished["attention_count"], 0)
        self.assertNotIn("write_api_key", page.text)

    def test_confirmed_load_runs_in_background_and_redirects_to_progress(self) -> None:
        context = self.app.state.context
        workspace_state = _create_workspace(
            context,
            name="Background load",
            source_system="Other",
        )
        workspace_state = context.workspace_states.update_target(
            workspace_state.workspace_id,
            actor=context.actor,
            expected_revision=workspace_state.revision,
            odoo_connection_mode="REMOTE",
            odoo_base_url="https://odoo.example.test",
            odoo_database="migration",
            intended_applications=("Contacts",),
            intended_models=(),
        )
        read_credential = store_target_credential(
            context.secret_store,
            workspace_state,
            TargetCredentialRole.READ,
            "read-secret",
            persistent=False,
        )
        store_target_credential(
            context.secret_store,
            workspace_state,
            TargetCredentialRole.WRITE,
            "write-secret",
            persistent=False,
        )
        semantic_hash = "sha256:" + "a" * 64
        target_hash = "sha256:" + "b" * 64
        read_context_hash = "sha256:" + "c" * 64
        snapshot = SimpleNamespace(
            semantic_hash=semantic_hash,
            target_hash=target_hash,
            target_database="migration",
            target_odoo_version="19.0",
            write_count=2,
            readable_models=("res.partner",),
            read_context_hash=read_context_hash,
            read_credential_binding_hash=read_credential.binding_hash,
        )
        preview = SimpleNamespace(
            snapshot=snapshot,
            api_scope=SimpleNamespace(models=()),
            scope_error="",
            current_run=None,
            can_load=True,
        )
        attempts = (
            SimpleNamespace(
                operation="CREATE",
                status=ExecutionRowStatus.COMMITTED,
            ),
            SimpleNamespace(
                operation="UPDATE",
                status=ExecutionRowStatus.COMMITTED,
            ),
        )
        completed_run = SimpleNamespace(
            run_id="11111111-1111-4111-8111-111111111111",
            total_count=2,
            planned_count=0,
            status=ExecutionRunStatus.COMPLETED,
            rows=attempts,
        )

        def execute(_project_id, **kwargs):
            kwargs["progress"](completed_run)
            return completed_run

        read_identity = OdooReadIdentity(
            target_hash=target_hash,
            principal_hash="sha256:" + "d" * 64,
            permission_hash="sha256:" + "e" * 64,
            context_hash=read_context_hash,
            readable_models=("res.partner",),
            observed_at="2026-08-23T00:00:00Z",
        )
        write_identity = OdooWriteIdentity(
            target_hash=target_hash,
            principal_hash="sha256:" + "1" * 64,
            permission_hash="sha256:" + "2" * 64,
            context_hash=read_context_hash,
            readable_models=("res.partner",),
            writable_models=("res.partner",),
            observed_at="2026-08-23T00:00:00Z",
        )

        with (
            patch.object(
                type(context.execution),
                "current_preview",
                return_value=preview,
            ),
            patch.object(type(context.execution), "execute", side_effect=execute),
            patch.object(
                type(context.reconciliation),
                "reconcile",
                return_value=SimpleNamespace(
                    unknown_count=0,
                    fallout_count=0,
                ),
            ),
            patch.object(
                type(context.cutover_plans),
                "assert_application_can_execute",
                return_value=None,
            ),
            patch.object(
                type(context.production_runs),
                "credential_workspace",
                return_value=workspace_state,
            ),
            patch.object(
                type(context.production_runs),
                "assert_execution_authority",
                return_value=None,
            ),
            patch.object(
                type(context.queries),
                "get_odoo_schema_catalog",
                return_value=SimpleNamespace(read_context_hash=read_context_hash),
            ),
            patch.object(context, "read_identity_probe", return_value=read_identity),
            patch.object(context, "write_identity_probe", return_value=write_identity),
            patch.object(context, "write_executor_factory", return_value=object()),
            patch.object(context, "readback_reader_factory", return_value=object()),
        ):
            started = self.client.post(
                f"/workspaces/{workspace_state.workspace_id}/load",
                data={
                    "csrf_token": self.csrf,
                    "snapshot_hash": semantic_hash,
                    "batch_rows": "10",
                },
                headers=POST_HEADERS,
                follow_redirects=False,
            )
            self.assertEqual(started.status_code, 303, started.text)
            self.assertIn("/load/progress/", started.headers["location"])
            finished = _wait_for_load(self.client, started.headers["location"])

        self.assertEqual(finished["status"], "SUCCEEDED", finished)
        self.assertEqual(finished["completed_rows"], 2)
        self.assertEqual(finished["created_count"], 1)
        self.assertEqual(finished["updated_count"], 1)
        self.assertTrue(finished["verification_complete"])
        self.assertEqual(
            finished["redirect_url"],
            f"/workspaces/{workspace_state.workspace_id}/load/outcome",
        )

    def test_local_schema_draft_does_not_call_the_odoo_api(self) -> None:
        context = self.app.state.context
        created = _create_workspace(
            context,
            name="Local draft",
            source_system="CSV",
        )
        now = datetime.now(timezone.utc)
        registered = replace(
            created,
            odoo_connection_mode=OdooConnectionMode.LOCAL,
            odoo_base_url="http://127.0.0.1:8069",
            odoo_database="odoo19_local",
            intended_models=("res.partner",),
            status=WorkspaceStatus.REGISTERED,
            revision=2,
            updated_at=now,
            registered_at=now,
        )
        context.workspace_states.repository.save(
            registered,
            expected_revision=created.revision,
            event_type="WORKSPACE_REGISTERED",
            event_detail="",
            actor=context.actor,
        )
        _replace_run_target_setup(
            context,
            registered.workspace_id,
            connection_mode=OdooConnectionMode.LOCAL,
            base_url="http://127.0.0.1:8069",
            database="odoo19_local",
        )
        context.sources.sources.save_source_selection(
            registered.workspace_id,
            SourceSelection(
                selection_id=str(uuid4()),
                version=1,
                data_version_id=_workspace_data_version_id(
                    context,
                    registered.workspace_id,
                ),
                created_at=now,
                created_by=context.actor.identity.display_name,
                datasets=(),
                content_hash="sha256:" + "a" * 64,
            ),
            actor=context.actor,
        )

        drafted = self._post(
            f"/workspaces/{registered.workspace_id}/schema/local-draft",
            {
                "csrf_token": self.csrf,
                "acknowledge_local_draft": "1",
                "manual_model_label_0": "Contact",
                "manual_fields_0": "name | Name | char | yes | no",
            },
        )
        self.assertEqual(drafted.status_code, 303, drafted.text)
        self.assertEqual(self.schema_calls, [])
        schema = context.schema_workspace.schemas.get_odoo_schema_catalog(
            registered.workspace_id
        )
        self.assertIsNotNone(schema)
        self.assertEqual(schema.origin, SchemaOrigin.LOCAL_MANUAL)
        schema_page = self.client.get(drafted.headers["location"])
        self.assertIn("Needs Odoo check", schema_page.text)
        self.assertIn("name | Name | char", schema_page.text)

    def test_registered_local_schema_uses_selected_config_without_api_key(
        self,
    ) -> None:
        context = self.app.state.context
        created = _create_workspace(
            context,
            name="Keyless local schema",
            source_system="CSV",
        )
        now = datetime.now(timezone.utc)
        registered = replace(
            created,
            odoo_connection_mode=OdooConnectionMode.LOCAL,
            odoo_base_url="http://127.0.0.1:18069",
            odoo_database="odoo19_local",
            intended_applications=("Contacts",),
            status=WorkspaceStatus.REGISTERED,
            revision=2,
            updated_at=now,
            registered_at=now,
        )
        context.workspace_states.repository.save(
            registered,
            expected_revision=created.revision,
            event_type="WORKSPACE_REGISTERED",
            event_detail="",
            actor=context.actor,
        )
        _replace_run_target_setup(
            context,
            registered.workspace_id,
            connection_mode=OdooConnectionMode.LOCAL,
            base_url="http://127.0.0.1:18069",
            database="odoo19_local",
            intended_applications=("Contacts",),
        )
        context.sources.sources.save_source_selection(
            registered.workspace_id,
            SourceSelection(
                selection_id=str(uuid4()),
                version=1,
                data_version_id=_workspace_data_version_id(
                    context,
                    registered.workspace_id,
                ),
                created_at=now,
                created_by=context.actor.identity.display_name,
                datasets=(),
                content_hash="sha256:" + "b" * 64,
            ),
            actor=context.actor,
        )
        unconfigured_page = self.client.get(
            f"/workspaces/{registered.workspace_id}/schema"
        )
        self.assertIn(
            "Find local Odoo first",
            unconfigured_page.text,
        )
        self.assertIn(
            f'action="/workspaces/{registered.workspace_id}/schema/local-config"',
            unconfigured_page.text,
        )
        blocked = self.client.post(
            f"/workspaces/{registered.workspace_id}/schema/models/refresh",
            data={"csrf_token": self.csrf},
            headers=POST_HEADERS,
        )
        self.assertEqual(blocked.status_code, 422)
        self.assertIn("Impodo could not check Odoo", blocked.text)

        workspace = Path(self.temporary.name) / "local-odoo"
        config = workspace / "config" / "odoo.conf"
        config.parent.mkdir(parents=True)
        config.write_text(
            "\n".join(
                (
                    "[options]",
                    "http_interface = 127.0.0.1",
                    "http_port = 18069",
                    "db_host = 127.0.0.1",
                    "db_port = 5544",
                    "db_user = odoo",
                    "db_name = odoo19_local",
                )
            ),
            encoding="utf-8",
        )
        for relative_path in (
            "venv/Scripts/python.exe",
            "odoo/odoo-bin",
        ):
            executable = workspace / relative_path
            executable.parent.mkdir(parents=True, exist_ok=True)
            executable.touch()
        status = context.local_stack.select_config(
            registered.workspace_id,
            config,
        )
        self.assertIsNotNone(status.profile)
        configured_local_stack = context.local_stack

        self.local_odoo_reader.get_model_catalog.return_value = (
            _browser_model_catalog(registered)
        )
        self.local_odoo_reader.get_model_metadata.return_value = (
            _browser_schema(registered)
        )

        page = self.client.get(f"/workspaces/{registered.workspace_id}/schema")
        self.assertIn("Local Odoo is ready to check", page.text)
        self.assertNotIn("Access key", page.text)
        self.assertIn("Show available Odoo data", page.text)
        self.assertNotIn("Verify access and load models", page.text)
        refreshed = self._post(
            f"/workspaces/{registered.workspace_id}/schema/models/refresh",
            {"csrf_token": self.csrf},
        )
        self.assertEqual(refreshed.status_code, 303)
        self.local_odoo_reader.get_model_catalog.assert_called_once()
        self.assertEqual(self.model_catalog_calls, [])
        verified_page = self.client.get(refreshed.headers["location"])
        self.assertIn("The Odoo list is ready", verified_page.text)
        self.assertIn(
            "The connection to local Odoo was checked during this session",
            verified_page.text,
        )
        self.assertIn("Update available choices", verified_page.text)
        self.assertIn("Has Odoo changed?", verified_page.text)
        self.assertIn("Save choices and continue", verified_page.text)
        self.assertNotIn("Save Odoo choices", verified_page.text)
        context.local_stack = LocalStackService()
        cached_page = self.client.get(
            f"/workspaces/{registered.workspace_id}/schema"
        )
        self.assertIn("The Odoo list is ready", cached_page.text)
        self.assertIn(
            "Support details",
            cached_page.text,
        )
        self.assertIn("The saved Odoo list is still available", cached_page.text)
        self.assertIn("res.partner", cached_page.text)
        self.local_odoo_reader.get_model_catalog.assert_called_once()

        context.local_stack = configured_local_stack
        workspace_state = context.workspace_states.repository.get(registered.workspace_id)
        scoped = self._post(
            f"/workspaces/{registered.workspace_id}/schema",
            {
                "csrf_token": self.csrf,
                "revision": str(workspace_state.revision),
                "permitted_models": "res.partner",
            },
        )
        self.assertEqual(scoped.status_code, 303)
        self.assertEqual(
            scoped.headers["location"],
            f"/workspaces/{registered.workspace_id}/schema#odoo-details",
        )
        self.local_odoo_reader.get_model_catalog.assert_called_once()
        self.local_odoo_reader.get_model_metadata.assert_called_once()
        metadata_call = (
            self.local_odoo_reader.get_model_metadata.call_args.args
        )
        self.assertEqual(metadata_call[2], ("res.partner",))
        self.assertEqual(self.schema_calls, [])
        scoped_page = self.client.get(scoped.headers["location"])
        self.assertIn(
            "Odoo data is ready",
            scoped_page.text,
        )
        unchanged_project = context.workspace_states.repository.get(
            registered.workspace_id
        )
        unchanged_scope = self._post(
            f"/workspaces/{registered.workspace_id}/schema",
            {
                "csrf_token": self.csrf,
                "revision": str(unchanged_project.revision),
                "permitted_models": "res.partner",
            },
        )
        self.assertEqual(unchanged_scope.status_code, 303)
        self.local_odoo_reader.get_model_metadata.assert_called_once()
        context.local_stack = LocalStackService()
        cached_schema_page = self.client.get(
            f"/workspaces/{registered.workspace_id}/schema"
        )
        self.assertIn(
            "Odoo details are ready",
            cached_schema_page.text,
        )
        self.assertIn('id="odoo-details"', cached_schema_page.text)
        self.assertIn(
            "The snapshot includes inherited fields and is used without another Odoo call",
            cached_schema_page.text,
        )
        self.assertIn(
            "Check for Odoo changes",
            cached_schema_page.text,
        )
        self.local_odoo_reader.get_model_metadata.assert_called_once()

        context.local_stack = configured_local_stack
        current_schema = context.queries.get_odoo_schema_catalog(
            registered.workspace_id
        )
        assert current_schema is not None
        self.local_odoo_reader.get_model_metadata.return_value = _browser_schema(
            registered
        )
        unchanged_check = self._post(
            f"/workspaces/{registered.workspace_id}/schema/capture",
            {"csrf_token": self.csrf},
        )
        self.assertEqual(unchanged_check.status_code, 303)
        unchanged_page = self.client.get(unchanged_check.headers["location"])
        self.assertIn("Odoo details are unchanged", unchanged_page.text)
        unchanged_schema = context.queries.get_odoo_schema_catalog(
            registered.workspace_id
        )
        assert unchanged_schema is not None
        self.assertEqual(unchanged_schema.content_hash, current_schema.content_hash)
        self.assertIsNone(unchanged_schema.pending_refresh)

        changed_snapshot = _browser_schema(registered)
        changed_partner = changed_snapshot.models["res.partner"]
        self.local_odoo_reader.get_model_metadata.return_value = replace(
            changed_snapshot,
            models={
                "res.partner": replace(
                    changed_partner,
                    fields={
                        **changed_partner.fields,
                        "name": replace(
                            changed_partner.fields["name"],
                            required=False,
                        ),
                    },
                )
            },
        )
        changed_check = self._post(
            f"/workspaces/{registered.workspace_id}/schema/capture",
            {"csrf_token": self.csrf},
        )
        self.assertEqual(changed_check.status_code, 303)
        changed_page = self.client.get(changed_check.headers["location"])
        self.assertIn("Odoo details changed", changed_page.text)
        self.assertIn("Field behavior changed", changed_page.text)
        self.assertIn("Use updated Odoo details", changed_page.text)
        self.assertIn("Needs attention", changed_page.text)
        reviewed_schema = context.queries.get_odoo_schema_catalog(
            registered.workspace_id
        )
        assert reviewed_schema is not None
        pending = reviewed_schema.pending_refresh
        assert pending is not None
        unconfirmed = self._post(
            f"/workspaces/{registered.workspace_id}/schema/capture/confirm",
            {
                "csrf_token": self.csrf,
                "expected_current_content_hash": (
                    pending.expected_current_content_hash
                ),
                "candidate_id": pending.candidate_id,
                "candidate_semantic_hash": pending.semantic_hash,
            },
        )
        self.assertEqual(unconfirmed.status_code, 422)
        self.assertIn("Confirm that Impodo may replace", unconfirmed.text)
        still_pending = context.queries.get_odoo_schema_catalog(
            registered.workspace_id
        )
        assert still_pending is not None
        self.assertEqual(still_pending.pending_refresh, pending)
        confirmed = self._post(
            f"/workspaces/{registered.workspace_id}/schema/capture/confirm",
            {
                "csrf_token": self.csrf,
                "expected_current_content_hash": (
                    pending.expected_current_content_hash
                ),
                "candidate_id": pending.candidate_id,
                "candidate_semantic_hash": pending.semantic_hash,
                "confirm_schema_refresh": "1",
            },
        )
        self.assertEqual(confirmed.status_code, 303)
        confirmed_schema = context.queries.get_odoo_schema_catalog(
            registered.workspace_id
        )
        assert confirmed_schema is not None
        self.assertIsNone(confirmed_schema.pending_refresh)
        self.assertNotEqual(
            confirmed_schema.content_hash,
            current_schema.content_hash,
        )

        self.local_odoo_reader.get_model_metadata.side_effect = ConnectorError(
            "raw local reader failure"
        )
        workspace_state = context.workspace_states.repository.get(registered.workspace_id)
        failed_load = self.client.post(
            f"/workspaces/{registered.workspace_id}/schema",
            data={
                "csrf_token": self.csrf,
                "revision": str(workspace_state.revision),
                "permitted_models": "product.template",
            },
            headers=POST_HEADERS,
        )
        self.assertEqual(failed_load.status_code, 422)
        self.assertIn(
            "Your Odoo choices were saved, but their details could not be loaded",
            failed_load.text,
        )
        self.assertIn("Try loading details again", failed_load.text)
        self.assertGreater(
            failed_load.text.index('id="odoo-load-error"'),
            failed_load.text.index('id="odoo-details"'),
        )
        self.assertLess(
            failed_load.text.index('id="odoo-load-error"'),
            failed_load.text.index("Try loading details again"),
        )
        self.assertNotIn("We could not complete that action", failed_load.text)
        saved_after_failure = context.workspace_states.repository.get(
            registered.workspace_id
        )
        self.assertEqual(
            saved_after_failure.intended_models,
            ("product.template",),
        )
        product_snapshot = _browser_schema(saved_after_failure)
        product_model = replace(
            product_snapshot.models["res.partner"],
            model="product.template",
        )
        self.local_odoo_reader.get_model_metadata.side_effect = None
        self.local_odoo_reader.get_model_metadata.return_value = replace(
            product_snapshot,
            models={"product.template": product_model},
        )
        recovered = self._post(
            f"/workspaces/{registered.workspace_id}/schema/capture",
            {"csrf_token": self.csrf},
        )
        self.assertEqual(recovered.status_code, 303)
        recovered_page = self.client.get(recovered.headers["location"])
        self.assertIn("Odoo data is ready", recovered_page.text)
        self.assertIn("Check for Odoo changes", recovered_page.text)

    def test_odoo_source_setup_skips_file_export_and_opens_schema_first(
        self,
    ) -> None:
        new_page = self.client.get("/projects/new")
        self.assertIn(">Files<", new_page.text)
        self.assertIn(">Data already in Odoo<", new_page.text)

        created = self._post(
            "/projects/new",
            {
                "csrf_token": self.csrf,
                "display_name": "Odoo product cleanup",
                "source_system_identity": "Odoo 19",
                "source_mode": "ODOO",
            },
        )
        self.assertEqual(created.status_code, 303)
        workspace_id = _created_workspace_id(self.app, created)
        data_project_id = self.app.state.context.migration_workspaces.get(
            workspace_id,
            actor=self.app.state.context.actor,
        ).project_id
        self.assertEqual(
            created.headers["location"],
            f"/projects/{data_project_id}",
        )
        project_page = self.client.get(created.headers["location"])
        self.assertIn(
            f'href="/workspaces/{workspace_id}/overview"',
            project_page.text,
        )
        target_page = self.client.get(f"/workspaces/{workspace_id}/target")
        self.assertIn("Connect the Odoo source", target_page.text)
        self.assertIn(
            "Moving records between two Odoo databases is not available yet.",
            target_page.text,
        )
        self.assertIn("It does not discover models or fields", target_page.text)
        files = self.client.get(
            f"/workspaces/{workspace_id}/files",
            follow_redirects=False,
        )
        self.assertEqual(files.status_code, 303)
        self.assertEqual(files.headers["location"], f"/workspaces/{workspace_id}/target")

        unchecked = self._post(
            f"/workspaces/{workspace_id}/target",
            {
                "csrf_token": self.csrf,
                "revision": "1",
                "odoo_connection_mode": "REMOTE",
                "odoo_base_url": "https://odoo.example.test",
                "odoo_database": "odoo_review",
                "read_api_key": "read-secret",
                "action": "save",
            },
        )
        self.assertEqual(unchecked.status_code, 422)
        self.assertIn("Check the Odoo connection before continuing", unchecked.text)

        checked = self._post(
            f"/workspaces/{workspace_id}/target",
            {
                "csrf_token": self.csrf,
                "revision": "2",
                "odoo_connection_mode": "REMOTE",
                "odoo_base_url": "https://odoo.example.test",
                "odoo_database": "odoo_review",
                "action": "test",
            },
        )
        self.assertEqual(checked.status_code, 303)
        self.assertEqual(
            self.read_identity_calls,
            [(workspace_id, "read-secret", ("res.users",))],
        )

        target = self._post(
            f"/workspaces/{workspace_id}/target",
            {
                "csrf_token": self.csrf,
                "revision": "3",
                "odoo_connection_mode": "REMOTE",
                "odoo_base_url": "https://odoo.example.test",
                "odoo_database": "odoo_review",
                "action": "save",
            },
        )
        self.assertEqual(
            target.headers["location"],
            f"/workspaces/{workspace_id}/schema",
        )
        workspace_state = self.app.state.context.workspace_states.repository.get(workspace_id)
        self.assertEqual(workspace_state.status, WorkspaceStatus.REGISTERED)
        self.assertEqual(workspace_state.source_mode, SourceMode.ODOO)
        self.assertEqual(workspace_state.source_system, "Odoo 19")
        self.assertEqual(workspace_state.source_files, ())

        schema_page = self.client.get(f"/workspaces/{workspace_id}/schema")
        self.assertIn("Stage 1 of 6", schema_page.text)
        self.assertIn("Odoo data", schema_page.text)
        self.assertIn("Choose the Odoo source record type", schema_page.text)

        refreshed = self._post(
            f"/workspaces/{workspace_id}/schema/models/refresh",
            {"csrf_token": self.csrf},
        )
        self.assertEqual(refreshed.status_code, 303)
        current = self.app.state.context.workspace_states.repository.get(workspace_id)
        scoped = self._post(
            f"/workspaces/{workspace_id}/schema",
            {
                "csrf_token": self.csrf,
                "revision": str(current.revision),
                "permitted_models": "res.partner",
            },
        )
        self.assertEqual(scoped.status_code, 303)
        self.assertEqual(self.schema_calls, [(workspace_id, "read-secret")])

        source_page = self.client.get(f"/workspaces/{workspace_id}/sources")
        self.assertEqual(source_page.status_code, 200)
        self.assertIn("Define a bounded Odoo capture", source_page.text)
        self.assertIn("Freezing is read-only", source_page.text)
        calls_before_selection = len(self.schema_calls)
        selected = self._post(
            f"/workspaces/{workspace_id}/sources/odoo-selection",
            {
                "csrf_token": self.csrf,
                "dataset_name": "odoo_contacts",
                "model": "res.partner",
                "field_names": "name",
                "include_archived": "",
                "max_rows": "1000",
            },
        )
        self.assertEqual(selected.status_code, 303)
        self.assertEqual(len(self.schema_calls), calls_before_selection)
        selection = (
            self.app.state.context.sources.sources
            .get_current_odoo_capture_selection(workspace_id)
        )
        self.assertIsNotNone(selection)
        self.assertEqual(selection.field_names, ("name",))
        saved_page = self.client.get(selected.headers["location"])
        self.assertIn("Capture plan version 1", saved_page.text)
        self.assertIn("Freeze these Odoo records", saved_page.text)
        self.assertIn("Ready to freeze", saved_page.text)

        context = self.app.state.context
        original_schema_reader = context.schema_reader
        changed_snapshot = _browser_schema(workspace_state)
        changed_partner = changed_snapshot.models["res.partner"]
        changed_snapshot = replace(
            changed_snapshot,
            models={
                "res.partner": replace(
                    changed_partner,
                    fields={
                        **changed_partner.fields,
                        "name": replace(
                            changed_partner.fields["name"],
                            required=False,
                        ),
                    },
                )
            },
        )
        context.schema_reader = lambda _workspace_state, _api_key: changed_snapshot
        change_check = self._post(
            f"/workspaces/{workspace_id}/schema/capture",
            {"csrf_token": self.csrf},
        )
        self.assertEqual(change_check.status_code, 303)
        attention_page = self.client.get(f"/workspaces/{workspace_id}/sources")
        self.assertIn("Odoo data needs attention", attention_page.text)
        self.assertIn("Review Odoo changes", attention_page.text)
        blocked_by_change = self._post(
            f"/workspaces/{workspace_id}/sources/odoo-capture",
            {
                "csrf_token": self.csrf,
                "selection_id": selection.selection_id,
                "selection_hash": selection.content_hash,
                "confirm_capture": "1",
            },
        )
        self.assertEqual(blocked_by_change.status_code, 422)
        self.assertIn("Review the checked Odoo changes", blocked_by_change.text)
        context.schema_reader = original_schema_reader
        cleared_check = self._post(
            f"/workspaces/{workspace_id}/schema/capture",
            {"csrf_token": self.csrf},
        )
        self.assertEqual(cleared_check.status_code, 303)

        stale = self._post(
            f"/workspaces/{workspace_id}/sources/odoo-capture",
            {
                "csrf_token": self.csrf,
                "selection_id": selection.selection_id,
                "selection_hash": "sha256:" + "0" * 64,
                "confirm_capture": "1",
            },
        )
        self.assertEqual(stale.status_code, 422)
        self.assertIn("out of date", stale.text)

        schema = self.app.state.context.queries.get_odoo_schema_catalog(workspace_id)
        self.assertIsNotNone(schema)
        gateway = _BrowserOdooCaptureGateway(workspace_state, schema)
        self.app.state.context.source_capture_factory = (
            lambda selected_workspace_state, _secret: gateway
        )
        started = self._post(
            f"/workspaces/{workspace_id}/sources/odoo-capture",
            {
                "csrf_token": self.csrf,
                "selection_id": selection.selection_id,
                "selection_hash": selection.content_hash,
                "confirm_capture": "1",
            },
        )
        self.assertEqual(started.status_code, 303)
        progress_url = started.headers["location"]
        progress_page = self.client.get(progress_url)
        self.assertIn("data-odoo-capture-job", progress_page.text)
        finished = _wait_for_odoo_capture(self.client, progress_url)
        self.assertEqual(finished["status"], "SUCCEEDED", finished)
        self.assertEqual(finished["completed_rows"], 2)
        self.assertEqual(finished["page_count"], 1)
        calls_after_capture = tuple(gateway.calls)

        frozen_page = self.client.get(finished["redirect_url"])
        self.assertEqual(tuple(gateway.calls), calls_after_capture)
        self.assertIn("Current frozen Odoo source", frozen_page.text)
        self.assertIn("2</dd>", frozen_page.text)
        self.assertIn("Protected history", frozen_page.text)
        self.assertIn("Frozen versions", frozen_page.text)

        mapping_page = self.client.get(f"/workspaces/{workspace_id}/mapping")
        self.assertEqual(mapping_page.status_code, 200)
        self.assertIn("Update only the records selected from Odoo", mapping_page.text)
        self.assertIn("Allow Impodo to update this field", mapping_page.text)
        self.assertIn('value="odoo_pinned_update"', mapping_page.text)
        self.assertNotIn("Which column uniquely identifies each row?", mapping_page.text)
        self.assertEqual(tuple(gateway.calls), calls_after_capture)

        source_selection = (
            self.app.state.context.queries.get_mapping_source_selection(workspace_id)
        )
        self.assertIsNotNone(source_selection)
        assert source_selection is not None
        name_row = re.search(
            r'data-target-field="name".*?name="scalar_value_source_0_(\d+)"',
            mapping_page.text,
            re.DOTALL,
        )
        self.assertIsNotNone(name_row)
        assert name_row is not None
        field_index = name_row.group(1)
        dataset = source_selection.datasets[0]
        source_column = dataset.columns[0]
        mapping_entries = [
            ["csrf_token", self.csrf],
            ["action", "draft"],
            ["expected_parent_version", ""],
            ["expected_working_draft_version", ""],
            ["editable_dataset_id", dataset.dataset_id],
            ["target_model_0", "res.partner"],
            ["mode_0", "odoo_pinned_update"],
            ["visible_scalar_target_0", "name"],
            [f"scalar_value_source_0_{field_index}", "source"],
            [f"scalar_source_0_{field_index}", source_column.stable_key],
            [f"scalar_type_0_{field_index}", "string"],
            [f"scalar_case_0_{field_index}", "uppercase"],
            [f"scalar_compare_0_{field_index}", "1"],
            ["approved_write_field_0", "name"],
        ]
        checked = self.client.post(
            f"/workspaces/{workspace_id}/mapping/save",
            json={"entries": mapping_entries},
            headers={**POST_HEADERS, "X-CSRF-Token": self.csrf},
        )
        self.assertEqual(checked.status_code, 200, checked.text)
        revision = self.app.state.context.queries.get_mapping_revision(workspace_id)
        working = self.app.state.context.queries.get_mapping_working_draft(workspace_id)
        self.assertIsNotNone(revision)
        self.assertIsNotNone(working)
        assert revision is not None
        assert working is not None
        mapping_entries[1] = ["action", "submit"]
        mapping_entries[2] = ["expected_parent_version", str(revision.version)]
        mapping_entries[3] = [
            "expected_working_draft_version",
            str(working.version),
        ]
        submitted = self.client.post(
            f"/workspaces/{workspace_id}/mapping/save",
            json={"entries": mapping_entries},
            headers={**POST_HEADERS, "X-CSRF-Token": self.csrf},
        )
        self.assertEqual(submitted.status_code, 200, submitted.text)
        prepared_page = self.client.get(f"/workspaces/{workspace_id}/prepare")
        self.assertEqual(prepared_page.status_code, 200)
        self.assertIn("Prepare data", prepared_page.text)
        self.assertEqual(tuple(gateway.calls), calls_after_capture)
        normalization = self.app.state.context.preparation.prepare(
            workspace_id,
            actor=self.app.state.context.actor,
        )
        self.assertEqual(
            normalization.eligible_record_count,
            2,
            self.app.state.context.quality.current_summary(workspace_id),
        )
        normalization_page = self.client.get(f"/workspaces/{workspace_id}/normalization")
        self.assertEqual(normalization_page.status_code, 200)
        self.assertIn("Review prepared changes", normalization_page.text)
        self.assertEqual(tuple(gateway.calls), calls_after_capture)
        approved = self._post(
            f"/workspaces/{workspace_id}/normalization/approve",
            {
                "csrf_token": self.csrf,
                "run_id": normalization.run_id,
                "lifecycle_version": str(normalization.lifecycle_version),
            },
        )
        self.assertEqual(approved.status_code, 303)
        self.assertEqual(
            approved.headers["location"],
            f"/workspaces/{workspace_id}/summary",
        )
        approved_page = self.client.get(approved.headers["location"])
        self.assertIn("Compare the approved data with Odoo", approved_page.text)
        self.assertIn("Compare with Odoo", approved_page.text)
        self.assertIn("Final review", approved_page.text)
        self.assertEqual(tuple(gateway.calls), calls_after_capture)

        def pinned_reader(
            selected_workspace_state,
            metadata_requests,
            record_requests,
        ):
            self.readiness_calls.append(
                (
                    selected_workspace_state.workspace_id,
                    metadata_requests,
                    record_requests,
                )
            )
            available = _browser_schema(selected_workspace_state)
            metadata = replace(
                available,
                models={
                    request.model: replace(
                        available.models[request.model],
                        fields={
                            field: available.models[request.model].fields[field]
                            for field in request.fields
                        },
                    )
                    for request in metadata_requests
                },
            )
            requested_fields = record_requests[0].fields
            return metadata, RecordSnapshot(
                fingerprint=metadata.fingerprint,
                records={
                    "res.partner": (
                        TargetRecord(
                            "res.partner",
                            11,
                            {
                                "name": "Alice",
                                "write_date": gateway.now.isoformat(),
                            },
                        ),
                        TargetRecord(
                            "res.partner",
                            12,
                            {
                                "name": "Bob",
                                "write_date": (
                                    gateway.now + timedelta(seconds=1)
                                ).isoformat(),
                            },
                        ),
                    )
                },
                requested_fields={"res.partner": requested_fields},
            )

        self.app.state.context.readiness_reader = pinned_reader
        compared = self._post(
            f"/workspaces/{workspace_id}/summary/compare",
            {"csrf_token": self.csrf},
        )
        self.assertEqual(compared.status_code, 303, compared.text)
        comparison_page = self.client.get(compared.headers["location"])
        self.assertIn("Comparison complete", comparison_page.text)
        self.assertIn("Ready to update", comparison_page.text)
        self.assertIn("2 records ready to update", comparison_page.text)
        self.assertNotIn("New in Odoo", comparison_page.text)
        self.assertNotIn("Create review workbook", comparison_page.text)
        self.assertIn("Load into Odoo", comparison_page.text)
        self.assertIn("Not yet available", comparison_page.text)
        report = self.app.state.context.preflight.current_report(workspace_id)
        self.assertIsNotNone(report)
        assert report is not None
        self.assertEqual(report.create_count, 0)
        self.assertEqual(report.update_count, 2)
        self.assertEqual(report.blocked_count, 0)
        protected_comparison = (
            self.app.state.context.preflight.current_odoo_comparison(
                workspace_id,
                actor=self.app.state.context.actor,
            )
        )
        self.assertIsNotNone(protected_comparison)
        assert protected_comparison is not None
        self.assertEqual(
            tuple(item.odoo_id for item in protected_comparison.rows),
            (11, 12),
        )
        self.assertEqual(
            {item.outcome for item in protected_comparison.rows},
            {OdooComparisonOutcome.UPDATE},
        )
        with self.app.state.context.artifacts.materialize_report(
            workspace_id,
            report.run_id,
            MANIFEST_NAME,
        ) as manifest_path:
            portable_manifest = manifest_path.read_text("utf-8")
        self.assertNotIn('"odoo_id"', portable_manifest)
        self.assertNotIn("Alice", portable_manifest)
        self.assertEqual(len(self.readiness_calls), 1)
        _workspace_id, _metadata_requests, record_requests = self.readiness_calls[0]
        self.assertEqual(len(record_requests), 1)
        self.assertEqual(record_requests[0].fields, ("name", "write_date"))
        self.assertEqual(record_requests[0].domain, (["id", "in", [11, 12]],))

        def conflicting_reader(
            selected_workspace_state,
            metadata_requests,
            record_requests,
        ):
            metadata, records = pinned_reader(
                selected_workspace_state,
                metadata_requests,
                record_requests,
            )
            return metadata, replace(
                records,
                records={
                    "res.partner": (
                        TargetRecord(
                            "res.partner",
                            11,
                            {
                                "name": "Changed elsewhere",
                                "write_date": (
                                    gateway.now + timedelta(minutes=1)
                                ).isoformat(),
                            },
                        ),
                        records.records["res.partner"][1],
                    )
                },
            )

        self.app.state.context.readiness_reader = conflicting_reader
        blocked_compare = self._post(
            f"/workspaces/{workspace_id}/summary/compare",
            {"csrf_token": self.csrf},
        )
        self.assertEqual(blocked_compare.status_code, 303, blocked_compare.text)
        blocked_page = self.client.get(blocked_compare.headers["location"])
        self.assertIn("Refresh the captured Odoo records", blocked_page.text)
        self.assertIn("Needs refresh", blocked_page.text)
        self.assertNotIn("data-preflight-compare", blocked_page.text)
        blocked_report = self.app.state.context.preflight.current_report(workspace_id)
        self.assertIsNotNone(blocked_report)
        assert blocked_report is not None
        self.assertEqual(blocked_report.blocked_count, 1)
        blocked_artifact = self.app.state.context.preflight.current_odoo_comparison(
            workspace_id,
            actor=self.app.state.context.actor,
        )
        self.assertIsNotNone(blocked_artifact)
        assert blocked_artifact is not None
        self.assertEqual(
            blocked_artifact.rows[0].outcome,
            OdooComparisonOutcome.CONCURRENT_FIELD_CHANGE,
        )

        previous_source_hash = source_selection.content_hash
        refresh_gateway = _BrowserOdooCaptureGateway(workspace_state, schema)
        refreshed_publication = (
            self.app.state.context.odoo_capture_publication.publish(
                workspace_id,
                refresh_gateway,
                actor=self.app.state.context.actor,
            )
        )
        self.assertNotEqual(
            refreshed_publication.source_selection.content_hash,
            previous_source_hash,
        )
        self.assertIsNone(
            self.app.state.context.queries.get_mapping_revision(workspace_id)
        )
        self.assertIsNone(
            self.app.state.context.preflight.current_staging(workspace_id)
        )
        self.assertIsNone(
            self.app.state.context.normalization.current_summary(workspace_id)
        )
        self.assertIsNone(
            self.app.state.context.preflight.current_report(workspace_id)
        )
        self.assertIsNone(
            self.app.state.context.preflight.current_odoo_comparison(
                workspace_id,
                actor=self.app.state.context.actor,
            )
        )

    def test_source_review_saves_table_choices_without_an_intermediate_page(
        self,
    ) -> None:
        created = self._post(
            "/projects/new",
            {
                "csrf_token": self.csrf,
                "display_name": "Inline source confirmation",
                "source_mode": "FILE",
                "source_system_identity": "Fictional ERP",
            },
        )
        workspace_id = _created_workspace_id(self.app, created)
        uploaded = self.client.post(
            f"/workspaces/{workspace_id}/files",
            data={"csrf_token": self.csrf, "revision": "1"},
            files={
                "source_file": (
                    "customers.csv",
                    b"code,name\nC001,Example\n",
                    "text/csv",
                )
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )
        self.assertEqual(uploaded.status_code, 303)
        context = self.app.state.context
        workspace_state = context.queries.get(workspace_id)
        registered = self._post(
            f"/workspaces/{workspace_id}/register",
            {
                "csrf_token": self.csrf,
                "revision": str(workspace_state.revision),
            },
        )
        self.assertEqual(registered.status_code, 303)

        inspected = self.client.post(
            f"/workspaces/{workspace_id}/sources/inspect",
            data={"csrf_token": self.csrf},
            headers=POST_HEADERS,
            follow_redirects=False,
        )
        self.assertEqual(inspected.status_code, 303)
        catalog = context.queries.get_source_catalogs(workspace_id)[0]
        confirmed = self.client.post(
            f"/workspaces/{workspace_id}/sources/{catalog.file_id}/configure",
            data={
                "csrf_token": self.csrf,
                "action": "confirm",
                "encoding": "utf-8",
                "delimiter": ",",
                "header_row_0": "1",
                "selected_0": "1",
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )
        self.assertEqual(confirmed.status_code, 303)
        source_page = self.client.get(confirmed.headers["location"])
        self.assertIn("Save the tables for this data version", source_page.text)
        self.assertIn('name="dataset_name_0"', source_page.text)
        self.assertIn(
            f'action="/workspaces/{workspace_id}/datasets/freeze"',
            source_page.text,
        )

        unfinished_page = self.client.get(
            f"/workspaces/{workspace_id}/datasets",
            follow_redirects=False,
        )
        self.assertEqual(unfinished_page.status_code, 303)
        self.assertEqual(
            unfinished_page.headers["location"],
            f"/workspaces/{workspace_id}/sources#table-choices",
        )

        frozen = self.client.post(
            f"/workspaces/{workspace_id}/datasets/freeze",
            data={
                "csrf_token": self.csrf,
                "dataset_name_0": "customers",
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )
        self.assertEqual(frozen.status_code, 303)
        saved_page = self.client.get(frozen.headers["location"])
        self.assertIn("Saved source tables", saved_page.text)
        self.assertIn("Tables ready for the next step", saved_page.text)
        self.assertNotIn('name="dataset_name_0"', saved_page.text)

    def test_complete_project_setup_registration_without_yaml(self) -> None:
        created = self._post(
            "/projects/new",
            {
                "csrf_token": self.csrf,
                "display_name": "Customer migration",
                "source_mode": "FILE",
                "source_system_identity": "Dynamics AX 2012",
            },
        )
        self.assertEqual(created.status_code, 303)
        workspace_id = _created_workspace_id(self.app, created)
        data_project_id = self.app.state.context.migration_workspaces.get(
            workspace_id,
            actor=self.app.state.context.actor,
        ).project_id
        self.assertEqual(
            created.headers["location"],
            f"/projects/{data_project_id}",
        )

        uploaded = self.client.post(
            f"/workspaces/{workspace_id}/files",
            data={"csrf_token": self.csrf, "revision": "1"},
            files=[
                (
                    "source_file",
                    (
                    "customers.csv",
                    b"code,name\nC001,Example\n",
                    "text/csv",
                    ),
                ),
                (
                    "source_file",
                    (
                    "products.xlsx",
                    _workbook_bytes(),
                    "application/vnd.openxmlformats-officedocument."
                    "spreadsheetml.sheet",
                    ),
                ),
            ],
            headers=POST_HEADERS,
            follow_redirects=False,
        )
        self.assertEqual(uploaded.status_code, 303)

        registered = self._post(
            f"/workspaces/{workspace_id}/register",
            {"csrf_token": self.csrf, "revision": "3"},
        )
        self.assertEqual(registered.status_code, 303)
        self.assertEqual(
            registered.headers["location"],
            f"/workspaces/{workspace_id}/sources",
        )
        summary = self.client.get(f"/workspaces/{workspace_id}/overview")
        self.assertIn("Data version overview", summary.text)
        self.assertIn("Ready for source check", summary.text)
        self.assertIn("Check source data", summary.text)
        self.assertIn(
            '<div class="overview-stage-number" aria-hidden="true">1</div>',
            summary.text,
        )
        self.assertIn(
            '<div class="overview-stage-number" aria-hidden="true">6</div>',
            summary.text,
        )
        self.assertIn("Source data", summary.text)
        self.assertIn("Load into Odoo", summary.text)
        self.assertIn(
            f'href="/workspaces/{workspace_id}/sources"',
            summary.text,
        )
        workspace_state = self.app.state.context.workspace_states.repository.get(workspace_id)
        self.assertEqual(workspace_state.status, WorkspaceStatus.REGISTERED)
        self.assertIsNone(workspace_state.odoo_connection_mode)
        self.assertEqual(workspace_state.mapping_version, None)
        manifest = (
            self.app.state.context.workspace_states.repository.workspace_directory(workspace_id)
            / "audit"
            / f"workspace-registration-r{workspace_state.revision}.json"
        )
        self.assertTrue(manifest.is_file())
        manifest_text = manifest.read_text()
        self.assertIn('"contract_version":6', manifest_text)
        self.assertIn('"odoo_connection_mode":null', manifest_text)
        for removed_field in (
            "business_unit",
            "data_manager",
            "description",
            "export_date",
            "export_status",
            "functional_owner",
            "support_access",
        ):
            self.assertNotIn(f'"{removed_field}"', manifest_text)

        source_discovery = self.client.get(registered.headers["location"])
        self.assertEqual(source_discovery.status_code, 200)
        self.assertIn("Stage 1 of 6", source_discovery.text)
        self.assertIn("Source data", source_discovery.text)
        self.assertIn('aria-current="step"', source_discovery.text)
        self.assertIn('aria-current="page"', source_discovery.text)
        self.assertIn("Check source files", source_discovery.text)
        self.assertIn("Your files have not been checked yet", source_discovery.text)
        self.assertIn("data-source-review-page", source_discovery.text)
        self.assertIn("data-source-review-form", source_discovery.text)
        inspected = self.client.post(
            f"/workspaces/{workspace_id}/sources/inspect",
            data={"csrf_token": self.csrf},
            headers=POST_HEADERS,
            follow_redirects=False,
        )
        self.assertEqual(inspected.status_code, 303)
        self.assertEqual(
            inspected.headers["location"],
            f"/workspaces/{workspace_id}/sources#source-files",
        )
        inspection_page = self.client.get(inspected.headers["location"])
        self.assertIn("Checked 2 source file", inspection_page.text)
        self.assertIn("customers.csv", inspection_page.text)
        self.assertIn("C001", inspection_page.text)
        self.assertIn("products.xlsx", inspection_page.text)
        self.assertIn("ProductTable", inspection_page.text)
        self.assertIn('class="source-table-summary"', inspection_page.text)
        self.assertIn('class="source-table-title"', inspection_page.text)
        self.assertIn("covers the same data", inspection_page.text)
        self.assertNotIn("Use separate Excel tables instead", inspection_page.text)
        self.assertIn("Likely content", inspection_page.text)
        self.assertIn("data-source-review-card", inspection_page.text)
        catalogs = self.app.state.context.sources.sources.get_source_catalogs(workspace_id)
        self.assertEqual(len(catalogs), 2)
        catalogs_by_name = {catalog.display_name: catalog for catalog in catalogs}
        customer_catalog = catalogs_by_name["customers.csv"]
        product_catalog = catalogs_by_name["products.xlsx"]
        source_files_by_name = {
            source.display_name: source for source in workspace_state.source_files
        }
        self.assertEqual(
            customer_catalog.source_sha256,
            source_files_by_name["customers.csv"].sha256,
        )
        self.assertIn("Data in customers.csv", inspection_page.text)
        self.assertNotIn(
            f"<h3>{customer_catalog.tables[0].name}</h3>",
            inspection_page.text,
        )

        configured = self.client.post(
            f"/workspaces/{workspace_id}/sources/{customer_catalog.file_id}/configure",
            data={
                "csrf_token": self.csrf,
                "action": "confirm",
                "encoding": "utf-8",
                "delimiter": ",",
                "header_row_0": "1",
                "selected_0": "1",
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )
        self.assertEqual(configured.status_code, 303)
        self.assertEqual(
            configured.headers["location"],
            f"/workspaces/{workspace_id}/sources#source-{customer_catalog.file_id}",
        )
        configured_page = self.client.get(configured.headers["location"])
        self.assertIn("Confirmed customers.csv", configured_page.text)

        workbook_configured = self.client.post(
            f"/workspaces/{workspace_id}/sources/{product_catalog.file_id}/configure",
            data={
                "csrf_token": self.csrf,
                "action": "confirm",
                "encoding": "",
                "delimiter": "",
                "header_row_0": "1",
                "selected_0": "1",
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )
        self.assertEqual(workbook_configured.status_code, 303)
        self.assertEqual(
            workbook_configured.headers["location"],
            f"/workspaces/{workspace_id}/sources#source-{product_catalog.file_id}",
        )
        configured_page = self.client.get(workbook_configured.headers["location"])
        self.assertIn("Confirmed products.xlsx", configured_page.text)
        self.assertIn("Save the tables for this data version", configured_page.text)
        self.assertIn(
            f'action="/workspaces/{workspace_id}/datasets/freeze"',
            configured_page.text,
        )
        self.assertIn('name="dataset_name_0"', configured_page.text)
        self.assertIn('name="dataset_name_1"', configured_page.text)
        self.assertIn(
            "source files cannot be added or removed from this data version",
            configured_page.text,
        )

        datasets = self.client.get(
            f"/workspaces/{workspace_id}/datasets",
            follow_redirects=False,
        )
        self.assertEqual(datasets.status_code, 303)
        self.assertEqual(
            datasets.headers["location"],
            f"/workspaces/{workspace_id}/sources#table-choices",
        )
        dataset_names: dict[str, str] = {}
        for configuration in (
            self.app.state.context.queries.get_source_configurations(workspace_id)
        ):
            for _table_key in configuration.selected_table_keys:
                dataset_names[f"dataset_name_{len(dataset_names)}"] = (
                    "customers"
                    if configuration.file_id == customer_catalog.file_id
                    else "products"
                )
        frozen = self.client.post(
            f"/workspaces/{workspace_id}/datasets/freeze",
            data={
                "csrf_token": self.csrf,
                **dataset_names,
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )
        self.assertEqual(frozen.status_code, 303)
        self.assertEqual(
            frozen.headers["location"],
            f"/workspaces/{workspace_id}/datasets#tables-ready",
        )
        saved_datasets = self.client.get(frozen.headers["location"])
        self.assertIn("Saved source tables", saved_datasets.text)
        self.assertIn('id="tables-ready"', saved_datasets.text)
        self.assertIn("Tables ready for the next step", saved_datasets.text)
        self.assertIn("Choose Odoo data", saved_datasets.text)
        self.assertNotIn('name="dataset_name_0"', saved_datasets.text)
        self.assertNotIn("Save tables for this data version", saved_datasets.text)

        odoo_data = self.client.get(
            f"/workspaces/{workspace_id}/schema",
            follow_redirects=False,
        )
        self.assertEqual(odoo_data.status_code, 303)
        self.assertEqual(
            odoo_data.headers["location"],
            f"/workspaces/{workspace_id}/target",
        )
        target_page = self.client.get(odoo_data.headers["location"])
        self.assertIn("Connect the Odoo destination", target_page.text)
        self.assertIn("It does not discover models or fields", target_page.text)

        workspace_state = self.app.state.context.queries.get(workspace_id)
        tested_target = self.client.post(
            f"/workspaces/{workspace_id}/target",
            data={
                "csrf_token": self.csrf,
                "revision": str(workspace_state.revision),
                "odoo_connection_mode": "REMOTE",
                "odoo_base_url": "https://odoo.example.test",
                "odoo_database": "odoo19_target",
                "read_api_key": "super-secret-token",
                "action": "test",
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )
        self.assertEqual(tested_target.status_code, 303)
        self.assertEqual(
            self.connection_calls,
            [(workspace_id, "super-secret-token", OdooConnectionMode.REMOTE)],
        )
        workspace_state = self.app.state.context.queries.get(workspace_id)
        saved_target = self._post(
            f"/workspaces/{workspace_id}/target",
            {
                "csrf_token": self.csrf,
                "revision": str(workspace_state.revision),
                "odoo_connection_mode": "REMOTE",
                "odoo_base_url": "https://odoo.example.test",
                "odoo_database": "odoo19_target",
                "action": "save",
            },
        )
        self.assertEqual(
            saved_target.headers["location"],
            f"/workspaces/{workspace_id}/schema",
        )
        workspace_state = self.app.state.context.queries.get(workspace_id)
        self.assertEqual(workspace_state.status, WorkspaceStatus.REGISTERED)
        self.assertEqual(workspace_state.odoo_connection_mode, OdooConnectionMode.REMOTE)
        self.assertNotIn(
            b"super-secret-token",
            (
                self.app.state.context.workspace_states.repository.workspace_directory(workspace_id)
                / "workspace-engine.duckdb"
            ).read_bytes(),
        )

        derived_page = self.client.get(
            f"/workspaces/{workspace_id}/derived-entities"
        )
        self.assertIn("Separate combined information", derived_page.text)
        self.assertIn("Stage 1 of 6", derived_page.text)
        self.assertIn("Optional source organization", derived_page.text)
        self.assertIn("You are viewing Source data", derived_page.text)
        self.assertIn("Current data-version work:", derived_page.text)
        self.assertIn("Stage 2", derived_page.text)
        self.assertIn("Odoo data", derived_page.text)
        self.assertIn(
            "Saved rules are repeated consistently for every row",
            derived_page.text,
        )
        self.assertIn("Create two related tables", derived_page.text)
        self.assertIn("Which field groups rows together?", derived_page.text)
        self.assertIn(
            "Which field identifies each row within its group?",
            derived_page.text,
        )
        self.assertNotIn("BOMId", derived_page.text)
        self.assertNotIn("dataAreaId", derived_page.text)
        self.assertNotIn("LineNum", derived_page.text)
        self.assertIn("Show available Odoo record types", derived_page.text)
        self.assertNotIn(
            (
                f'action="/workspaces/{workspace_id}/derived-entities/'
                'lookup/preview#lookup-preview"'
            ),
            derived_page.text,
        )
        selection = (
            self.app.state.context.sources.sources.get_source_selection(workspace_id)
        )
        self.assertIsNotNone(selection)
        assert selection is not None
        datasets_by_name = {item.name: item for item in selection.datasets}
        customer_dataset = datasets_by_name["customers"]
        product_dataset = datasets_by_name["products"]
        source_choices = _source_value_choices(
            self.app.state.context,
            workspace_id,
            customer_dataset.dataset_id,
            customer_dataset.columns[0].stable_key,
        )
        self.assertEqual(source_choices, ({"value": "C001", "count": 1},))
        product_name = product_dataset.columns[1]
        product_code = product_dataset.columns[0]
        related_preview = self.client.post(
            f"/workspaces/{workspace_id}/derived-entities/related/preview",
            data={
                "csrf_token": self.csrf,
                "expected_parent_version": "",
                "source_dataset_id": product_dataset.dataset_id,
                "parent_dataset_name": "product_groups",
                "child_dataset_name": "product_rows",
                "parent_key_column_key": product_name.stable_key,
                "scope_column_key": "",
                "child_key_column_key": product_code.stable_key,
                "blank_policy": "block",
            },
            headers=POST_HEADERS,
        )
        self.assertEqual(related_preview.status_code, 200)
        self.assertIn("Review before creating", related_preview.text)
        self.assertIn("Create these separate tables", related_preview.text)
        saved_related = self.client.post(
            f"/workspaces/{workspace_id}/derived-entities/related/save",
            data={
                "csrf_token": self.csrf,
                "expected_parent_version": "",
                "source_dataset_id": product_dataset.dataset_id,
                "parent_dataset_name": "product_groups",
                "child_dataset_name": "product_rows",
                "parent_key_column_key": product_name.stable_key,
                "scope_column_key": "",
                "child_key_column_key": product_code.stable_key,
                "blank_policy": "block",
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )
        self.assertEqual(saved_related.status_code, 303)
        related_page = self.client.get(saved_related.headers["location"])
        self.assertIn(
            "Created the separate tables product_groups and product_rows",
            related_page.text,
        )
        related_plan = (
            self.app.state.context.derived_entities.derived_entities.get_derived_entity_plan(workspace_id)
        )
        self.assertIsNotNone(related_plan)
        removed_related = self.client.post(
            (
                f"/workspaces/{workspace_id}/derived-entities/"
                f"{related_plan.rules[0].rule_id}/delete"
            ),
            data={
                "csrf_token": self.csrf,
                "expected_parent_version": "1",
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )
        self.assertEqual(removed_related.status_code, 303)
        derived_rule_data = {
            "csrf_token": self.csrf,
            "expected_parent_version": "2",
            "source_binding": (
                f"{product_dataset.dataset_id}|{product_name.stable_key}"
            ),
            "output_dataset_name": "product_names",
            "target_model": "res.partner",
            "target_name_field": "name",
            "external_id_namespace": "dynamics_ax_2012",
            "parent_separator": "",
            "blank_policy": "block",
        }
        blocked_without_models = self.client.post(
            f"/workspaces/{workspace_id}/derived-entities/lookup/preview",
            data=derived_rule_data,
            headers=POST_HEADERS,
        )
        self.assertEqual(blocked_without_models.status_code, 422)
        self.assertIn(
            "Show the available Odoo record types before choosing one",
            blocked_without_models.text,
        )

        workspace_state = self.app.state.context.workspace_states.repository.get(workspace_id)
        refreshed_lookup_models = self._post(
            f"/workspaces/{workspace_id}/derived-entities/models/refresh",
            {"csrf_token": self.csrf},
        )
        self.assertEqual(refreshed_lookup_models.status_code, 303)
        self.assertEqual(
            refreshed_lookup_models.headers["location"],
            f"/workspaces/{workspace_id}/derived-entities#lookup-extraction",
        )
        self.assertEqual(
            self.model_catalog_calls,
            [(workspace_id, "super-secret-token")],
        )
        lookup_model_page = self.client.get(
            refreshed_lookup_models.headers["location"]
        )
        self.assertIn("Odoo record types are ready", lookup_model_page.text)
        self.assertIn(
            (
                f'action="/workspaces/{workspace_id}/derived-entities/'
                'lookup/preview#lookup-preview"'
            ),
            lookup_model_page.text,
        )
        self.assertIn('value="res.partner" label="Contact"', lookup_model_page.text)
        self.assertIn("Start typing an Odoo record type", lookup_model_page.text)
        self.assertNotIn('placeholder="product_categories"', lookup_model_page.text)
        self.assertNotIn("Article and Service", lookup_model_page.text)

        rejected_lookup_model = self.client.post(
            f"/workspaces/{workspace_id}/derived-entities/lookup/preview",
            data={**derived_rule_data, "target_model": "x.not.available"},
            headers=POST_HEADERS,
        )
        self.assertEqual(rejected_lookup_model.status_code, 422)
        self.assertIn(
            "Choose an existing Odoo record type from the loaded list",
            rejected_lookup_model.text,
        )
        rejected_lookup_save = self.client.post(
            f"/workspaces/{workspace_id}/derived-entities/save",
            data={**derived_rule_data, "target_model": "x.not.available"},
            headers=POST_HEADERS,
        )
        self.assertEqual(rejected_lookup_save.status_code, 422)
        self.assertIn(
            "Choose an existing Odoo record type from the loaded list",
            rejected_lookup_save.text,
        )
        lookup_preview = self.client.post(
            f"/workspaces/{workspace_id}/derived-entities/lookup/preview",
            data=derived_rule_data,
            headers=POST_HEADERS,
        )
        self.assertEqual(lookup_preview.status_code, 200)
        self.assertIn("Review before creating", lookup_preview.text)
        self.assertIn('id="lookup-preview"', lookup_preview.text)
        self.assertIn("Create this related table", lookup_preview.text)
        saved_derived = self.client.post(
            f"/workspaces/{workspace_id}/derived-entities/save",
            data=derived_rule_data,
            headers=POST_HEADERS,
            follow_redirects=False,
        )
        self.assertEqual(saved_derived.status_code, 303)
        saved_plan = (
            self.app.state.context.derived_entities.derived_entities.get_derived_entity_plan(
                workspace_id
            )
        )
        self.assertIsNotNone(saved_plan)
        saved_rule = next(
            rule
            for rule in saved_plan.rules
            if getattr(rule, "output_dataset_name", None) == "product_names"
        )
        self.assertEqual(
            saved_derived.headers["location"],
            (
                f"/workspaces/{workspace_id}/derived-entities"
                f"#lookup-rule-{saved_rule.rule_id}"
            ),
        )
        derived_preview = self.client.get(saved_derived.headers["location"])
        self.assertIn(
            f'id="lookup-rule-{saved_rule.rule_id}"',
            derived_preview.text,
        )
        self.assertIn("Created the related table product_names", derived_preview.text)
        self.assertIn("Example product", derived_preview.text)
        self.assertIn("impodo_dynamics_ax_2012.res_partner_", derived_preview.text)
        self.assertIn(
            "available when you match data beside the original rows",
            derived_preview.text,
        )
        self.assertNotIn("entity:P001", derived_preview.text)

        refreshed_models = self._post(
            f"/workspaces/{workspace_id}/schema/models/refresh",
            {"csrf_token": self.csrf},
        )
        self.assertEqual(refreshed_models.status_code, 303)
        self.assertEqual(
            self.model_catalog_calls,
            [
                (workspace_id, "super-secret-token"),
                (workspace_id, "super-secret-token"),
            ],
        )
        model_page = self.client.get(refreshed_models.headers["location"])
        self.assertIn("Stage 2 of 6", model_page.text)
        self.assertIn("Odoo data", model_page.text)
        self.assertIn('aria-current="step"', model_page.text)
        self.assertIn('aria-current="page"', model_page.text)
        self.assertIn("Choose the Odoo data you need", model_page.text)
        self.assertNotIn("Odoo areas included:", model_page.text)
        self.assertIn("Contact", model_page.text)
        self.assertIn("res.partner", model_page.text)
        self.assertIn(
            "No Odoo area was selected, so all available business records are shown.",
            model_page.text,
        )
        self.assertIn(
            f'action="/workspaces/{workspace_id}/schema"',
            model_page.text,
        )
        self.assertIn('aria-live="polite"', model_page.text)
        self.assertIn(
            'data-model-search-text="product product.template product stock"',
            model_page.text,
        )

        model_picker_script = self.client.get("/static/app.js")
        self.assertIn("const hasQuery = Boolean(query);", model_picker_script.text)
        self.assertIn(
            "matches && (hasQuery || browseAll || choice.inFocus || selected)",
            model_picker_script.text,
        )
        self.assertIn(
            "Saving choices and loading Odoo data...",
            model_picker_script.text,
        )
        model_picker_styles = self.client.get("/static/app.css")
        self.assertIn("label.model-choice[hidden]", model_picker_styles.text)

        rejected_scope = self.client.post(
            f"/workspaces/{workspace_id}/schema",
            data={
                "csrf_token": self.csrf,
                "revision": str(workspace_state.revision),
                "permitted_models": "x.not.available",
            },
            headers=POST_HEADERS,
        )
        self.assertEqual(rejected_scope.status_code, 422)
        self.assertIn(
            "not in the refreshed Odoo model catalogue",
            rejected_scope.text,
        )

        scope = self.client.post(
            f"/workspaces/{workspace_id}/schema",
            data={
                "csrf_token": self.csrf,
                "revision": str(workspace_state.revision),
                "permitted_models": "res.partner",
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )
        self.assertEqual(scope.status_code, 303)
        workspace_state = self.app.state.context.workspace_states.repository.get(workspace_id)
        self.assertEqual(workspace_state.intended_models, ("res.partner",))
        self.assertEqual(self.schema_calls, [(workspace_id, "super-secret-token")])
        workspace_state = self.app.state.context.workspace_states.repository.get(workspace_id)
        read_credential = get_target_credential(
            self.secrets,
            workspace_state,
            TargetCredentialRole.READ,
        )
        assert read_credential is not None
        model_catalog = (
            self.app.state.context.schema_workspace.schemas.get_odoo_model_catalog(
                workspace_id
            )
        )
        schema_catalog = (
            self.app.state.context.schema_workspace.schemas.get_odoo_schema_catalog(
                workspace_id
            )
        )
        assert model_catalog is not None
        assert schema_catalog is not None
        self.assertEqual(
            model_catalog.read_credential_binding_hash,
            read_credential.binding_hash,
        )
        self.assertEqual(
            schema_catalog.read_credential_binding_hash,
            read_credential.binding_hash,
        )
        self.assertEqual(
            model_catalog.read_principal_hash,
            "sha256:" + "1" * 64,
        )
        self.assertEqual(
            schema_catalog.read_principal_hash,
            model_catalog.read_principal_hash,
        )
        self.assertEqual(
            schema_catalog.read_context_hash,
            model_catalog.read_context_hash,
        )
        self.assertNotEqual(
            schema_catalog.read_permission_hash,
            model_catalog.read_permission_hash,
        )
        schema_page = self.client.get(scope.headers["location"])
        self.assertIn("Tell Impodo how to find existing records", schema_page.text)
        self.assertIn("How should Impodo find an existing Contact?", schema_page.text)
        self.assertNotIn("Reference (ref)", schema_page.text)
        self.assertIn("Search fields", schema_page.text)
        self.assertIn("Show fields that Odoo controls", schema_page.text)
        self.assertIn("Impodo found no single safe recommendation", schema_page.text)
        self.assertIn("Reference", schema_page.text)
        self.assertIn("Support options for combined matching", schema_page.text)
        governed = self.client.post(
            f"/workspaces/{workspace_id}/schema/govern",
            data={
                "csrf_token": self.csrf,
                "primary_key_field_0": "ref",
                "primary_scope_field_0": "",
                "key_fields_0": "",
                "scope_fields_0": "",
                "key_description_0": "Unique contact reference",
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )
        self.assertEqual(governed.status_code, 303)
        mapping_page = self.client.get(governed.headers["location"])
        self.assertIn("Stage 3 of 6", mapping_page.text)
        self.assertIn("Match data", mapping_page.text)
        self.assertIn('aria-current="step"', mapping_page.text)
        self.assertIn('aria-current="page"', mapping_page.text)
        self.assertIn("Match your data to Odoo", mapping_page.text)
        self.assertIn('<details class="technical-evidence">', mapping_page.text)
        self.assertNotIn("Evidence binding", mapping_page.text)
        self.assertIn("Which column uniquely identifies each row?", mapping_page.text)
        self.assertIn("How should Impodo find the same record in Odoo?", mapping_page.text)
        self.assertIn('class="identity-pair"', mapping_page.text)
        self.assertIn("Use existing Odoo records only", mapping_page.text)
        self.assertIn("Source value, or backup when blank", mapping_page.text)
        self.assertNotIn("Let Odoo choose", mapping_page.text)
        self.assertIn("Choose what goes into each Odoo field", mapping_page.text)
        self.assertNotIn("Scalar target fields", mapping_page.text)
        self.assertIn("Find a field", mapping_page.text)
        self.assertIn("For example: Sales Price or Barcode", mapping_page.text)
        self.assertNotIn("list_price", mapping_page.text)
        self.assertIn(
            "Choose where each Odoo field gets its value",
            mapping_page.text,
        )
        self.assertIn("Connect values to existing Odoo lists", mapping_page.text)
        self.assertIn("such as a product category", mapping_page.text)
        self.assertIn("Find an Odoo list or linked field", mapping_page.text)
        self.assertIn("data-relation-pagination", mapping_page.text)
        self.assertIn("data-scalar-pagination", mapping_page.text)
        self.assertIn("data-scalar-table-scroll-top", mapping_page.text)
        self.assertIn(
            'aria-label="Scroll scalar target fields horizontally"',
            mapping_page.text,
        )
        self.assertIn("data-scalar-table-scroll", mapping_page.text)
        self.assertIn("Preview", mapping_page.text)
        self.assertIn("Prepare and check values", mapping_page.text)
        self.assertIn("Must be exactly", mapping_page.text)
        self.assertIn("The first characters", mapping_page.text)
        self.assertIn("Add cleanup step", mapping_page.text)
        self.assertIn("Remove separators between numbers", mapping_page.text)
        self.assertIn("data-text-step-storage", mapping_page.text)
        self.assertIn("Advanced: custom pattern", mapping_page.text)
        self.assertIn("Advanced: formula or custom calculation", mapping_page.text)
        self.assertIn("Safe formulas only", mapping_page.text)

        mapping_script = self.client.get("/static/app.js")
        self.assertIn("updateScalarTableScroll", mapping_script.text)
        self.assertIn(
            "new ResizeObserver(updateScalarTableScroll)",
            mapping_script.text,
        )
        self.assertIn('window.addEventListener("beforeunload"', mapping_script.text)
        self.assertIn("rememberMappingPosition", mapping_script.text)
        self.assertIn("restoreMappingPosition", mapping_script.text)
        self.assertIn("rememberMappingInteraction", mapping_script.text)
        self.assertIn("visibleMappingRow", mapping_script.text)
        self.assertIn(
            'mappingForm.addEventListener("pointerdown"',
            mapping_script.text,
        )
        self.assertIn("window.sessionStorage", mapping_script.text)
        self.assertIn("preventScroll: true", mapping_script.text)
        self.assertIn("rememberNormalizationPosition", mapping_script.text)
        self.assertIn("restoreNormalizationPosition", mapping_script.text)
        self.assertIn("data-normalization-reject-form", mapping_script.text)
        self.assertIn("normalizationApproveDialog", mapping_script.text)
        self.assertIn("normalizationRejectDialog", mapping_script.text)
        self.assertIn("rememberSourceReviewPosition", mapping_script.text)
        self.assertIn("restoreSourceReviewPosition", mapping_script.text)
        self.assertIn("data-source-review-form", mapping_script.text)
        self.assertIn("scheduleScalarCatalogSearch", mapping_script.text)
        self.assertIn("catalogRequestUrl", mapping_script.text)
        self.assertIn("relationRequestUrl", mapping_script.text)
        self.assertIn("new AbortController()", mapping_script.text)
        self.assertIn("new DOMParser()", mapping_script.text)
        self.assertIn("window.history.replaceState", mapping_script.text)
        self.assertIn("restoreScalarRow(row)", mapping_script.text)
        self.assertIn("restoreRelationRow(row)", mapping_script.text)
        self.assertIn(
            "is ordinary text here.",
            mapping_script.text,
        )
        self.assertIn("internationalPhoneTextSteps", mapping_script.text)
        self.assertIn('characters: " .-/"', mapping_script.text)
        self.assertIn("scheduleRelationCatalogSearch", mapping_script.text)
        self.assertIn("relationDraftRows", mapping_script.text)
        self.assertNotIn("pendingRedirect", mapping_script.text)
        self.assertNotIn(
            "mappingForm.requestSubmit(saveProgress)",
            mapping_script.text,
        )
        self.assertIn(
            "Searching Odoo fields",
            mapping_script.text,
        )
        self.assertIn(
            'mappingForm.getAttribute("action")',
            mapping_script.text,
        )
        self.assertNotIn("fetch(mappingForm.action", mapping_script.text)
        self.assertIn(
            "updateMappingVersionFields(payload)",
            mapping_script.text,
        )
        self.assertIn("workingVersionUpdated", mapping_script.text)
        self.assertIn('if (action === "save_progress")', mapping_script.text)
        self.assertIn("navigateToMappingResult", mapping_script.text)
        self.assertIn(
            "Your unsaved changes are still on this page",
            mapping_script.text,
        )
        self.assertIn(
            "Your checked matches are unchanged",
            mapping_script.text,
        )
        self.assertIn("hydrateSourceOptions", mapping_script.text)
        self.assertIn("option.defaultSelected = selected", mapping_script.text)
        mapping_styles = self.client.get("/static/app.css")
        self.assertIn(".scalar-table-scroll-top", mapping_styles.text)
        self.assertIn("overflow-x: scroll", mapping_styles.text)
        self.assertIn(".mapping-save-state.unsaved", mapping_styles.text)
        self.assertIn(".source-table-summary", mapping_styles.text)
        self.assertIn(".source-table-title", mapping_styles.text)

        selection = (
            self.app.state.context.sources.sources.get_source_selection(workspace_id)
        )
        schema_governance = (
            self.app.state.context.schema_workspace.schemas.get_schema_governance(workspace_id)
        )
        self.assertIsNotNone(selection)
        self.assertIsNotNone(schema_governance)
        assert selection is not None
        physical_datasets_by_name = {item.name: item for item in selection.datasets}
        customer = physical_datasets_by_name["customers"]
        product = physical_datasets_by_name["products"]
        customer_code, customer_name = customer.columns
        product_code, product_name = product.columns
        mapping_selection = (
            self.app.state.context.sources.sources.get_mapping_source_selection(
                workspace_id
            )
        )
        self.assertIsNotNone(mapping_selection)
        mapping_datasets = {
            item.name: item for item in mapping_selection.datasets
        }
        mapped_customer = mapping_datasets["customers"]
        product_names = mapping_datasets["product_names"]
        mapped_product = mapping_datasets["products"]
        product_name_key, product_name_value = product_names.columns
        dataset_indices = {
            item.dataset_id: index
            for index, item in enumerate(mapping_selection.datasets)
        }
        customer_index = dataset_indices[mapped_customer.dataset_id]
        product_names_index = dataset_indices[product_names.dataset_id]
        product_index = dataset_indices[mapped_product.dataset_id]
        self.assertEqual(mapped_customer.dataset_id, customer.dataset_id)
        self.assertEqual(mapped_product.dataset_id, product.dataset_id)
        business_key_id = schema_governance.business_keys[0].key_id
        mapping_filter_query = (
            "scalar_page=1&field_query=nam&"
            f"mapping_dataset={customer_index}&relation_page=1"
        )
        saved_progress = self.client.post(
            f"/workspaces/{workspace_id}/mapping/save?{mapping_filter_query}",
            data={
                "csrf_token": self.csrf,
                "action": "save_progress",
                "expected_parent_version": "",
                "expected_working_draft_version": "",
                f"target_model_{customer_index}": "res.partner",
                f"mode_{customer_index}": "upsert",
                f"scalar_value_source_{customer_index}_1": "source",
                f"scalar_type_{customer_index}_1": "string",
                f"scalar_case_{customer_index}_1": "preserve",
                f"scalar_formula_{customer_index}_1": (
                    'coalesce(value, "Unnamed contact")'
                ),
                f"scalar_compare_{customer_index}_1": "1",
                f"scalar_null_{customer_index}_1": "distinct",
                f"target_model_{product_names_index}": "res.partner",
                f"mode_{product_names_index}": "upsert",
                f"target_model_{product_index}": "res.partner",
                f"mode_{product_index}": "upsert",
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )
        self.assertEqual(saved_progress.status_code, 303)
        self.assertEqual(
            saved_progress.headers["location"],
            (
                f"/workspaces/{workspace_id}/mapping?{mapping_filter_query}"
                f"#mapping-dataset-{customer_index}"
            ),
        )
        saved_progress_page = self.client.get(
            saved_progress.headers["location"]
        )
        self.assertIn(
            f'id="mapping-dataset-{customer_index}"',
            saved_progress_page.text,
        )
        self.assertIn("data-mapping-dataset", saved_progress_page.text)
        self.assertIn(
            "Saved your matching progress. Check the matches when ready.",
            saved_progress_page.text,
        )
        saved_draft = (
            self.app.state.context.mapping_workspace.mappings.get_mapping_working_draft(
                workspace_id
            )
        )
        current_mapping_selection = (
            self.app.state.context.queries.get_mapping_source_selection(workspace_id)
        )
        self.assertIsNotNone(saved_draft)
        self.assertIsNotNone(current_mapping_selection)
        self.assertEqual(
            saved_draft.definition.source_selection_hash,
            current_mapping_selection.content_hash,
        )
        current_governance = (
            self.app.state.context.queries.get_schema_governance(workspace_id)
        )
        self.assertIsNotNone(current_governance)
        self.assertEqual(
            saved_draft.definition.schema_hash,
            current_governance.content_hash,
        )
        self.assertIn("Saved changes need checking", saved_progress_page.text)
        self.assertIn("Your saved work is loaded", saved_progress_page.text)
        working_draft = (
            self.app.state.context.mapping_workspace.mappings.get_mapping_working_draft(
                workspace_id
            )
        )
        self.assertEqual(working_draft.version, 1)
        working_by_dataset = {
            item.dataset_id: item
            for item in working_draft.definition.datasets
        }
        self.assertEqual(
            working_by_dataset[customer.dataset_id].fields[0].source_column_key,
            "",
        )
        self.assertIsNone(
            self.app.state.context.mapping_workspace.mappings.get_mapping_revision(workspace_id)
        )
        mapping_data = {
                "csrf_token": self.csrf,
                f"target_model_{customer_index}": "res.partner",
                f"mode_{customer_index}": "upsert",
                f"source_identity_{customer_index}": customer_code.stable_key,
                f"business_key_{customer_index}": business_key_id,
                f"identity_source_{customer_index}_0": customer_code.stable_key,
                f"scalar_value_source_{customer_index}_1": "source_with_fallback",
                f"scalar_source_{customer_index}_1": customer_name.stable_key,
                f"scalar_literal_{customer_index}_1": "Unnamed contact",
                f"scalar_type_{customer_index}_1": "string",
                f"scalar_trim_{customer_index}_1": "1",
                f"scalar_collapse_{customer_index}_1": "1",
                f"scalar_empty_null_{customer_index}_1": "1",
                f"scalar_case_{customer_index}_1": "preserve",
                f"scalar_formula_{customer_index}_1": (
                    'coalesce(value, "Unnamed contact")'
                ),
                f"scalar_compare_{customer_index}_1": "1",
                f"scalar_null_{customer_index}_1": "distinct",
                f"target_model_{product_names_index}": "res.partner",
                f"mode_{product_names_index}": "upsert",
                f"source_identity_{product_names_index}": product_name_key.stable_key,
                f"business_key_{product_names_index}": business_key_id,
                f"identity_source_{product_names_index}_0": product_name_value.stable_key,
                f"scalar_value_source_{product_names_index}_1": "source",
                f"scalar_source_{product_names_index}_1": product_name_value.stable_key,
                f"scalar_type_{product_names_index}_1": "string",
                f"scalar_compare_{product_names_index}_1": "1",
                f"scalar_null_{product_names_index}_1": "distinct",
                f"target_model_{product_index}": "res.partner",
                f"mode_{product_index}": "upsert",
                f"source_identity_{product_index}": product_code.stable_key,
                f"business_key_{product_index}": business_key_id,
                f"identity_source_{product_index}_0": product_code.stable_key,
                f"scalar_value_source_{product_index}_1": "constant",
                f"scalar_literal_{product_index}_1": "Imported product",
                f"scalar_type_{product_index}_1": "string",
                f"scalar_case_{product_index}_1": "sentence",
                f"scalar_text_steps_{product_index}_1": (
                    '[{"kind":"find_replace","search_value":"Imported",'
                    '"replacement_value":"imported","search_mode":"literal",'
                    '"replace_all":true,"characters":""}]'
                ),
                f"scalar_exact_length_{product_index}_1": "16",
                f"scalar_segment_location_{product_index}_1": "first",
                f"scalar_segment_length_{product_index}_1": "1",
                f"scalar_character_class_{product_index}_1": "uppercase",
                f"scalar_pattern_{product_index}_1": "[A-Z][a-z ]{15}",
                f"scalar_compare_{product_index}_1": "1",
                f"scalar_null_{product_index}_1": "distinct",
        }
        checked = self.client.post(
            f"/workspaces/{workspace_id}/mapping/save",
            data={
                **mapping_data,
                "action": "draft",
                "expected_parent_version": "",
                "expected_working_draft_version": "1",
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )
        self.assertEqual(checked.status_code, 303)
        checked_page = self.client.get(checked.headers["location"])
        self.assertIn("Matches checked and ready to confirm", checked_page.text)
        self.assertIn("Ready to confirm", checked_page.text)
        checked_draft = (
            self.app.state.context.mapping_workspace.mappings.get_mapping_working_draft(
                workspace_id
            )
        )
        self.assertEqual(checked_draft.version, 2)
        checked_revision = (
            self.app.state.context.mapping_workspace.mappings.get_mapping_revision(
                workspace_id
            )
        )
        self.assertEqual(checked_revision.version, 1)

        submitted = self.client.post(
            f"/workspaces/{workspace_id}/mapping/save",
            data={
                **mapping_data,
                "action": "submit",
                "expected_parent_version": "1",
                "expected_working_draft_version": "2",
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )
        self.assertEqual(submitted.status_code, 303)
        self.assertEqual(
            submitted.headers["location"],
            f"/workspaces/{workspace_id}/prepare",
        )
        submitted_page = self.client.get(submitted.headers["location"])
        self.assertIn("Field matches confirmed", submitted_page.text)
        self.assertIn("Prepare all source rows", submitted_page.text)
        confirmed_mapping_page = self.client.get(
            f"/workspaces/{workspace_id}/mapping"
        )
        self.assertIn("Field matches are confirmed", confirmed_mapping_page.text)
        self.assertIn("Continue to Prepare data", confirmed_mapping_page.text)
        self.assertRegex(
            confirmed_mapping_page.text,
            rf'class="button secondary" href="/workspaces/{workspace_id}/prepare"',
        )
        revision = (
            self.app.state.context.mapping_workspace.mappings.get_mapping_revision(workspace_id)
        )
        self.assertEqual(revision.version, 1)
        self.assertEqual(
            [
                item.version
                for item in self.app.state.context.mapping_workspace.mappings.list_mapping_revisions(
                    workspace_id
                )
            ],
            [1],
        )
        revision_by_dataset = {
            item.dataset_id: item for item in revision.definition.datasets
        }
        self.assertEqual(
            revision_by_dataset[
                customer.dataset_id
            ].fields[0].value_source.value,
            "source_with_fallback",
        )
        self.assertTrue(
            revision_by_dataset[customer.dataset_id].fields[0].transform.trim
        )
        self.assertEqual(
            revision_by_dataset[customer.dataset_id].fields[0].transform.formula,
            'coalesce(value, "Unnamed contact")',
        )
        self.assertEqual(
            revision_by_dataset[
                product.dataset_id
            ].fields[0].value_source.value,
            "constant",
        )
        self.assertEqual(
            revision_by_dataset[product.dataset_id].fields[0].literal_value,
            "Imported product",
        )
        product_field = revision_by_dataset[product.dataset_id].fields[0]
        self.assertEqual(product_field.transform.case_mode, "sentence")
        self.assertEqual(
            product_field.transform.text_steps[0].search_value,
            "Imported",
        )
        self.assertEqual(product_field.validation.exact_length, 16)
        self.assertEqual(product_field.validation.segment_location, "first")
        self.assertEqual(product_field.validation.character_class, "uppercase")

        impact_link = (
            f"/workspaces/{workspace_id}/mapping/transformation-impact"
        )
        self.assertIn("Review rule effects (optional)", confirmed_mapping_page.text)
        impact_page = self.client.get(impact_link)
        self.assertEqual(impact_page.status_code, 200)
        self.assertIn("Review rule effects", impact_page.text)
        self.assertIn("Stage 3 of 6", impact_page.text)
        self.assertIn("Optional rule review", impact_page.text)
        self.assertIn('aria-current="step"', impact_page.text)
        self.assertIn('aria-current="page"', impact_page.text)
        self.assertIn("This review is optional", impact_page.text)
        self.assertIn("Prepare preview", impact_page.text)
        self.assertNotIn("What each rule did", impact_page.text)
        self.assertIn("your checked field rules", impact_page.text)
        self.assertNotIn("data-impact-row", impact_page.text)
        prepared = self.client.post(
            f"{impact_link}/prepare",
            data={"csrf_token": self.csrf},
            headers=POST_HEADERS,
            follow_redirects=False,
        )
        self.assertEqual(prepared.status_code, 303)
        impact_page = self.client.get(prepared.headers["location"])
        self.assertIn("Original", impact_page.text)
        self.assertIn("Prepared", impact_page.text)
        self.assertIn("All Odoo fields", impact_page.text)
        self.assertIn("Download matching rows (.csv)", impact_page.text)
        self.assertIn("Download all affected rows (.csv)", impact_page.text)
        self.assertIn("Your registered Excel or CSV source remains unchanged", impact_page.text)
        self.assertIn("Showing 1", impact_page.text)
        self.assertNotIn("data-impact-row", impact_page.text)
        impact_csv = self.client.post(
            f"{impact_link}.csv",
            data={"csrf_token": self.csrf},
            headers=POST_HEADERS,
        )
        self.assertEqual(impact_csv.status_code, 200)
        self.assertIn("text/csv", impact_csv.headers["content-type"])
        self.assertIn("Raw source", impact_csv.text)
        self.assertIn("Proposed value", impact_csv.text)
        mapping_script = self.client.get("/static/app.js")
        self.assertNotIn("data-impact-export", mapping_script.text)

        prepare_page = self.client.get(f"/workspaces/{workspace_id}/prepare")
        self.assertEqual(prepare_page.status_code, 200)
        self.assertIn("Stage 4 of 6", prepare_page.text)
        self.assertIn("Prepare data", prepare_page.text)
        self.assertIn("Prepare all source rows", prepare_page.text)
        self.assertIn(
            "Impodo prepares from the accepted Data version selected by this workspace",
            prepare_page.text,
        )
        self.assertIn(
            f'action="/workspaces/{workspace_id}/summary/check"',
            prepare_page.text,
        )
        self.assertIn('aria-current="step"', prepare_page.text)
        self.assertIn('aria-current="page"', prepare_page.text)

        summary = self.client.get(f"/workspaces/{workspace_id}/summary")
        self.assertIn("Prepare and review data", summary.text)
        self.assertIn("Uses Impodo’s stored local copy", summary.text)
        checked = self.client.post(
            f"/workspaces/{workspace_id}/summary/check",
            data={"csrf_token": self.csrf},
            headers=POST_HEADERS,
            follow_redirects=False,
        )
        self.assertEqual(checked.status_code, 303)
        self.assertIn("/preparation/", checked.headers["location"])
        progress_page = self.client.get(checked.headers["location"])
        self.assertIn("Stage 4 of 6", progress_page.text)
        self.assertIn("Prepare data", progress_page.text)
        self.assertIn(
            "Impodo is preparing from its stored local copy",
            progress_page.text,
        )
        self.assertIn('aria-current="step"', progress_page.text)
        self.assertIn('aria-current="page"', progress_page.text)
        completed_job = _wait_for_preparation(
            self.client,
            checked.headers["location"],
        )
        self.assertEqual(completed_job["status"], "SUCCEEDED", completed_job)
        manager = self.app.state.context.preparation_jobs
        assert manager is not None
        worker_deadline = time.monotonic() + 2.0
        while (
            manager.worker_alive(str(completed_job["job_id"]))
            and time.monotonic() < worker_deadline
        ):
            time.sleep(0.01)
        self.assertFalse(manager.worker_alive(str(completed_job["job_id"])))
        review_page = self.client.get(str(completed_job["redirect_url"]))
        self.assertIn("Review what Impodo prepared", review_page.text)
        self.assertIn("Nothing is sent to Odoo", review_page.text)
        self.assertIn("data-normalization-review", review_page.text)
        self.assertIn("Approve all prepared data", review_page.text)
        self.assertIn("Send back to fix", review_page.text)
        self.assertNotIn("Accept this change", review_page.text)
        self.assertIn("data-normalization-approve-dialog", review_page.text)
        self.assertIn("data-normalization-reject-dialog", review_page.text)
        self.assertIn("data-normalization-table-scroll", review_page.text)
        self.assertEqual(len(self.readiness_calls), 0)

        context = self.app.state.context
        quality_summary = context.quality.current_summary(workspace_id)
        assert quality_summary is not None
        quality_page = context.queries.get_quality_review_page(
            workspace_id,
            quality_summary.run_id,
            status="",
            dataset="",
            page=1,
            page_size=20,
        )
        self.assertLessEqual(len(quality_page.items), 20)
        prepared_summary_page = self.client.get(
            f"/workspaces/{workspace_id}/summary"
        )
        self.assertIn(
            f"Records 1-{min(20, quality_page.matching_count)} "
            f"of {quality_page.matching_count}",
            prepared_summary_page.text,
        )
        if quality_page.matching_count > 10:
            self.assertIn("Records per page:", prepared_summary_page.text)

        normalization_service = self.app.state.context.normalization
        review = normalization_service.current_review(workspace_id)
        assert review is not None
        normalization, evaluation, dry_run = review
        decision_group = next(
            group for group in evaluation.groups if group.requires_decision
        )
        rejected = self.client.post(
            f"/workspaces/{workspace_id}/normalization/groups/"
            f"{decision_group.group_id}/reject?status=pending&page=2",
            data={
                "csrf_token": self.csrf,
                "run_id": normalization.run_id,
                "lifecycle_version": str(normalization.lifecycle_version),
                "reason": "The prepared value needs another review.",
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )
        self.assertEqual(rejected.status_code, 303)
        blocked_page = self.client.get(rejected.headers["location"])
        self.assertIn("Fix the change that was sent back", blocked_page.text)
        self.assertIn("The prepared value needs another review", blocked_page.text)
        self.assertIn("Reopen review", blocked_page.text)
        self.assertNotIn("Accept this change", blocked_page.text)
        normalization = normalization_service.current_summary(workspace_id)
        assert normalization is not None
        reopened = self.client.post(
            f"/workspaces/{workspace_id}/normalization/reopen",
            data={
                "csrf_token": self.csrf,
                "run_id": normalization.run_id,
                "lifecycle_version": str(normalization.lifecycle_version),
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )
        self.assertEqual(reopened.status_code, 303)
        self.assertEqual(
            reopened.headers["location"],
            f"/workspaces/{workspace_id}/normalization?status=pending#review-groups",
        )
        normalization = normalization_service.current_summary(workspace_id)
        assert normalization is not None
        approved = self.client.post(
            f"/workspaces/{workspace_id}/normalization/approve",
            data={
                "csrf_token": self.csrf,
                "run_id": normalization.run_id,
                "lifecycle_version": str(normalization.lifecycle_version),
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )
        self.assertEqual(approved.status_code, 303)
        normalization = normalization_service.current_summary(workspace_id)
        assert normalization is not None
        self.assertTrue(normalization.frozen)
        frozen_review = normalization_service.current_review(workspace_id)
        assert frozen_review is not None
        self.assertEqual(
            frozen_review[2].approved_groups,
            frozen_review[2].summary.required_group_keys,
        )
        workspace = self.app.state.context.migration_workspaces.get(
            workspace_id,
            actor=self.app.state.context.actor,
        )
        package = (
            self.app.state.context.data_version_source_projection.packages.repository
            .get_source_package(workspace.data_version_id)
        )
        assert package is not None
        source_artifact = (
            Path(self.temporary.name)
            / "artifacts"
            / "dv"
            / workspace.data_version_id
            / "inbox"
            / package.files[0].storage_key
        )
        source_artifact.unlink()
        compared = self.client.post(
            f"/workspaces/{workspace_id}/summary/compare",
            data={"csrf_token": self.csrf},
            headers=POST_HEADERS,
            follow_redirects=False,
        )
        self.assertEqual(compared.status_code, 303, compared.text)
        readiness_page = self.client.get(compared.headers["location"])
        self.assertIn("Included in preparation", readiness_page.text)
        self.assertIn("Set aside", readiness_page.text)
        self.assertIn("Needs correction", readiness_page.text)
        self.assertIn("New in Odoo", readiness_page.text)
        self.assertIn("Different from Odoo", readiness_page.text)
        self.assertIn("Already matches", readiness_page.text)
        self.assertIn("Needs attention", readiness_page.text)
        self.assertIn('id="quality-rows"', readiness_page.text)
        self.assertIn("Ready", readiness_page.text)
        self.assertIn("Needs a decision", readiness_page.text)
        self.assertIn("Needs correction", readiness_page.text)
        self.assertIn("Rows", readiness_page.text)
        self.assertIn("Support details", readiness_page.text)
        self.assertIn("Create review workbook", readiness_page.text)
        self.assertIn("Keep these rules for another data version", readiness_page.text)
        self.assertIn(
            "Your submitted preparation, matching, relationship, and checking "
            "rules are already saved in this workspace.",
            readiness_page.text,
        )
        self.assertIn("Save these rules as a Recipe", readiness_page.text)
        self.assertIn(
            f'href="/projects/{workspace.project_id}#project-recipes-title"',
            readiness_page.text,
        )
        self.assertIn(
            "The review workbook documents this data check. "
            "It does not save the rules for reuse.",
            readiness_page.text,
        )
        self.assertIn("Odoo remains unchanged", readiness_page.text)
        self.assertIn("prepared rows safely saved", readiness_page.text)
        self.assertIn("data-staging-summary", readiness_page.text)
        self.assertIn("<summary>Support details</summary>", readiness_page.text)
        self.assertIn("data-preflight-compare", readiness_page.text)
        self.assertIn(
            "Comparing with Odoo... Keep this page open.",
            readiness_page.text,
        )

        report = self.app.state.context.preflight.current_report(workspace_id)
        assert report is not None
        self.assertEqual(
            report.create_count
            + report.update_count
            + report.unchanged_count
            + report.ambiguous_count
            + report.blocked_count,
            report.total_count,
        )
        staging = self.app.state.context.preflight.current_staging(workspace_id)
        assert staging is not None
        self.assertEqual(report.staging_run_id, staging.run_id)
        self.assertEqual(report.staging_content_hash, staging.content_hash)
        restored_staging = (
            self.app.state.context.preflight.staging.get_canonical_staging_run(
                workspace_id,
                staging.run_id,
            )
        )
        self.assertIsNotNone(restored_staging)
        self.assertEqual(
            restored_staging.content_hash,
            staging.content_hash,
        )
        restart_app = create_local_app(
            self.temporary.name,
            secret_store=self.secrets,
            readiness_reader=lambda *_args: self.fail(
                "Restart retrieval must not contact Odoo"
            ),
        )
        restarted_report = restart_app.state.context.preflight.current_report(
            workspace_id
        )
        self.assertIsNotNone(restarted_report)
        assert restarted_report is not None
        self.assertEqual(restarted_report.run_id, report.run_id)

        sample_row = self.app.state.context.preflight.readiness_rows(
            workspace_id,
            report.run_id,
        ).items[0]
        self.assertIn(
            sample_row.source_trace_id,
            {item.row_id for item in restored_staging.rows},
        )

        database_path = (
            self.app.state.context.workspace_states.repository.workspace_directory(workspace_id)
            / "workspace-engine.duckdb"
        )
        staging_repository = self.app.state.context.preflight.staging
        with staging_repository._connect(database_path) as connection:
            stored_row = connection.execute(
                """
                SELECT ordinal, row_json
                  FROM canonical_staging_row
                 WHERE run_id = ?
                 ORDER BY ordinal
                 LIMIT 1
                """,
                [staging.run_id],
            ).fetchone()
            assert stored_row is not None
            tampered_payload = json.loads(str(stored_row[1]))
            tampered_payload["target_model"] = "x.tampered"
            connection.execute(
                """
                UPDATE canonical_staging_row
                   SET row_json = ?
                 WHERE run_id = ? AND ordinal = ?
                """,
                [
                    canonical_json_text(tampered_payload),
                    staging.run_id,
                    int(stored_row[0]),
                ],
            )
        try:
            rejected_tamper = self.client.post(
                f"/workspaces/{workspace_id}/summary/compare",
                data={"csrf_token": self.csrf},
                headers=POST_HEADERS,
                follow_redirects=False,
            )
        finally:
            with staging_repository._connect(database_path) as connection:
                connection.execute(
                    """
                    UPDATE canonical_staging_row
                       SET row_json = ?
                     WHERE run_id = ? AND ordinal = ?
                    """,
                    [str(stored_row[1]), staging.run_id, int(stored_row[0])],
                )
        self.assertEqual(rejected_tamper.status_code, 422)
        self.assertEqual(len(self.readiness_calls), 1)

        compared_again = self.client.post(
            f"/workspaces/{workspace_id}/summary/compare",
            data={"csrf_token": self.csrf},
            headers=POST_HEADERS,
            follow_redirects=False,
        )
        self.assertEqual(compared_again.status_code, 303, compared_again.text)
        repeated_report = self.app.state.context.preflight.current_report(workspace_id)
        assert repeated_report is not None
        self.assertNotEqual(repeated_report.run_id, report.run_id)
        self.assertEqual(repeated_report.staging_run_id, report.staging_run_id)
        self.assertEqual(repeated_report.quality_run_id, report.quality_run_id)
        self.assertEqual(
            repeated_report.normalization_run_id,
            report.normalization_run_id,
        )
        with staging_repository._connect(database_path) as connection:
            run_count = connection.execute(
                "SELECT COUNT(*) FROM readiness_run"
            ).fetchone()
            current_run = connection.execute(
                "SELECT run_id FROM preflight_current WHERE singleton_id = 1"
            ).fetchone()
            superseded = connection.execute(
                """
                SELECT detail
                  FROM preflight_transition
                 WHERE run_id = ? AND event_type = 'SUPERSEDED'
                """,
                [report.run_id],
            ).fetchone()
        self.assertEqual(run_count, (2,))
        self.assertEqual(current_run, (repeated_report.run_id,))
        self.assertEqual(superseded, (repeated_report.run_id,))
        current_normalization = normalization_service.current_summary(workspace_id)
        assert current_normalization is not None
        self.assertEqual(
            current_normalization.run_id,
            report.normalization_run_id,
        )

        load_preview = self.app.state.context.execution.current_preview(workspace_id)
        assert load_preview is not None
        final_review = self.client.get(f"/workspaces/{workspace_id}/summary")
        self.assertIn("Final review complete", final_review.text)
        self.assertNotIn("Preview the Odoo load", final_review.text)
        self.assertLess(
            final_review.text.index("Checked rows"),
            final_review.text.index("Final review complete"),
        )
        load_landing = self.client.get(
            f"/workspaces/{workspace_id}/load",
            follow_redirects=False,
        )
        self.assertEqual(load_landing.status_code, 303)
        self.assertTrue(load_landing.headers["location"].endswith("/load/review"))
        load_page = self.client.get(load_landing.headers["location"])
        self.assertEqual(load_page.status_code, 200)
        self.assertIn("Check what will change in Odoo", load_page.text)
        self.assertIn("Nothing is written from this page", load_page.text)
        self.assertIn("Reviewed captured Odoo fields only", load_page.text)
        self.assertIn("Continue to confirmation", load_page.text)
        self.assertNotIn('name="write_api_key"', load_page.text)
        outcome_before_load = self.client.get(
            f"/workspaces/{workspace_id}/load/outcome",
            follow_redirects=False,
        )
        self.assertEqual(outcome_before_load.status_code, 303)
        self.assertTrue(
            outcome_before_load.headers["location"].endswith("/load/review")
        )
        confirmation_page = self.client.get(
            f"/workspaces/{workspace_id}/load/confirm"
        )
        self.assertEqual(confirmation_page.status_code, 200)
        self.assertIn("Confirm the Odoo load", confirmation_page.text)
        self.assertIn("bounded batches in dependency order", confirmation_page.text)
        self.assertIn("will not retry blindly", confirmation_page.text)
        self.assertIn("Advanced settings", confirmation_page.text)
        self.assertIn('name="write_api_key"', confirmation_page.text)
        self.assertIn('name="remember_write_api_key"', confirmation_page.text)
        self.assertIn("checking-only key is not reused here", confirmation_page.text)

        class FakeWriteExecutor:
            target_hash = load_preview.snapshot.target_hash
            scope_hash = load_preview.api_scope.semantic_hash

            def __init__(self):
                self.created = []
                self.updated = []
                self.records = {}
                self.next_id = 100

            def find_ids(self, model, domain):
                del model, domain
                return (42,)

            def create_rows(self, model, values):
                rows = tuple(dict(item) for item in values)
                self.created.append((model, rows))
                identifiers = tuple(
                    range(self.next_id, self.next_id + len(rows))
                )
                self.next_id += len(rows)
                for identifier, row in zip(identifiers, rows, strict=True):
                    self.records[(model, identifier)] = dict(row)
                return identifiers

            def load_create_rows(self, model, values, external_ids):
                del external_ids
                return self.create_rows(model, values)

            def update_row(self, model, record_id, values):
                self.updated.append((model, record_id, dict(values)))
                self.records.setdefault((model, record_id), {}).update(values)

        class FakeReadbackReader:
            target_hash = load_preview.snapshot.target_hash
            scope_hash = load_preview.api_scope.semantic_hash
            imports_external_ids = False

            def __init__(self, writer):
                self.writer = writer

            def read_ids(self, model, identifiers, fields):
                return tuple(
                    ReadbackRecord(
                        identifier,
                        {
                            field: self.writer.records[(model, identifier)][field]
                            for field in fields
                        },
                    )
                    for identifier in identifiers
                    if (model, identifier) in self.writer.records
                )

            def find_records(self, model, domain, fields):
                del domain
                matches = [
                    ReadbackRecord(
                        identifier,
                        {field: values[field] for field in fields},
                    )
                    for (stored_model, identifier), values in self.writer.records.items()
                    if stored_model == model and all(field in values for field in fields)
                ]
                if matches:
                    return tuple(matches[:2])
                return (ReadbackRecord(42, {}),) if not fields else ()

            def find_records_many(self, model, lookups):
                return tuple(
                    self.find_records(model, lookup.domain, lookup.fields)
                    for lookup in lookups
                )

            def read_external_ids(self, external_ids):
                del external_ids
                return ()

        fake_writer = FakeWriteExecutor()
        write_factory_keys = []
        readback_factory_keys = []

        def write_factory(_project, api_key, _scope):
            write_factory_keys.append(api_key)
            return fake_writer

        def readback_factory(_project, api_key, _scope):
            readback_factory_keys.append(api_key)
            return FakeReadbackReader(fake_writer)

        def write_identity_probe(_project, _api_key, scope):
            return OdooWriteIdentity(
                target_hash=load_preview.snapshot.target_hash,
                principal_hash="sha256:" + "5" * 64,
                permission_hash="sha256:" + "6" * 64,
                context_hash=load_preview.snapshot.read_context_hash,
                readable_models=tuple(item.model for item in scope.models),
                writable_models=tuple(
                    item.model for item in scope.models if item.write_fields
                ),
                observed_at="2026-08-21T00:00:00Z",
            )

        self.app.state.context.write_executor_factory = write_factory
        self.app.state.context.readback_reader_factory = readback_factory
        self.app.state.context.write_identity_probe = write_identity_probe
        missing_write_key = self.client.post(
            f"/workspaces/{workspace_id}/load",
            data={
                "csrf_token": self.csrf,
                "snapshot_hash": load_preview.snapshot.semantic_hash,
                "batch_rows": "10",
            },
            headers=POST_HEADERS,
        )
        self.assertEqual(missing_write_key.status_code, 422)
        self.assertIn("Enter an Odoo API key approved for loading", missing_write_key.text)
        self.assertIn("Confirm the Odoo load", missing_write_key.text)
        self.assertEqual(write_factory_keys, [])
        self.assertEqual(readback_factory_keys, [])

        loaded = self.client.post(
            f"/workspaces/{workspace_id}/load",
            data={
                "csrf_token": self.csrf,
                "snapshot_hash": load_preview.snapshot.semantic_hash,
                "write_api_key": "load-secret",
                "batch_rows": "10",
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )
        self.assertEqual(loaded.status_code, 303, loaded.text)
        progress_url = loaded.headers["location"]
        self.assertIn("/load/progress/", progress_url)
        progress_page = self.client.get(progress_url)
        self.assertEqual(progress_page.status_code, 200, progress_page.text)
        self.assertIn("data-load-job", progress_page.text)
        self.assertIn("odoo.example.test", progress_page.text)
        self.assertIn("migration", progress_page.text)
        finished_load = _wait_for_load(self.client, progress_url)
        self.assertEqual(finished_load["status"], "SUCCEEDED", finished_load)
        self.assertEqual(
            finished_load["completed_rows"],
            repeated_report.create_count + repeated_report.update_count,
        )
        self.assertEqual(finished_load["created_count"], repeated_report.create_count)
        self.assertEqual(finished_load["updated_count"], repeated_report.update_count)
        self.assertEqual(finished_load["attention_count"], 0)
        self.assertTrue(finished_load["verification_complete"])
        self.assertNotIn("load-secret", json.dumps(finished_load))
        self.assertEqual(write_factory_keys, ["load-secret"])
        self.assertEqual(readback_factory_keys, ["load-secret"])
        outcome_page = self.client.get(finished_load["redirect_url"])
        self.assertIn("Odoo read-back complete", outcome_page.text)
        self.assertIn("Verified in Odoo", outcome_page.text)
        self.assertIn("Odoo now matches every field", outcome_page.text)
        self.assertIn("Verify result", outcome_page.text)
        self.assertEqual(
            outcome_page.text.count("data-load-row"),
            repeated_report.create_count + repeated_report.update_count,
        )
        self.assertNotIn("load-secret", outcome_page.text)

        report = repeated_report
        sample_row = self.app.state.context.preflight.readiness_rows(
            workspace_id,
            report.run_id,
        ).items[0]
        paged_rows = tuple(
            replace(
                sample_row,
                source_trace_id=f"sha256:{index:064x}",
                source_row=index,
                status="blocked" if index <= 120 else "ready",
                identity=f"ROW-{index:04d}",
            )
            for index in range(1, 202)
        )
        with staging_repository._connect(database_path) as connection:
            connection.execute(
                "DELETE FROM preflight_decision WHERE run_id = ?",
                [report.run_id],
            )
            connection.executemany(
                """
                INSERT INTO preflight_decision (
                    run_id, ordinal, source_trace_id, dataset,
                    source_row, status, decision_json
                ) VALUES (?, ?, ?, ?, ?, ?, ?)
                """,
                [
                    [
                        report.run_id,
                        index,
                        item.source_trace_id,
                        item.dataset,
                        item.source_row,
                        item.status,
                        canonical_json_text(asdict(item)),
                    ]
                    for index, item in enumerate(paged_rows)
                ],
            )

        with self.subTest("persisted readiness paging"):
            first_page = self.client.get(
                f"/workspaces/{workspace_id}/summary"
            )
            self.assertEqual(
                first_page.text.count("data-readiness-row"),
                20,
            )
            self.assertIn("Rows 1-20 of 201", first_page.text)
            self.assertIn("Page 1 of 11", first_page.text)
            self.assertIn("Rows per page:", first_page.text)
            for size in (10, 20, 50, 100):
                self.assertIn(f">{size}</a>", first_page.text)
            self.assertIn("ROW-0020", first_page.text)
            self.assertNotIn("ROW-0021", first_page.text)
            next_match = re.search(
                r'href="([^"]+)" data-readiness-next',
                first_page.text,
            )
            assert next_match is not None
            next_query = parse_qs(
                urlsplit(unescape(next_match.group(1))).query
            )
            self.assertEqual(next_query["page"], ["2"])

            second_page = self.client.get(
                f"/workspaces/{workspace_id}/summary?page=2"
            )
            self.assertEqual(
                second_page.text.count("data-readiness-row"),
                20,
            )
            self.assertIn("Rows 21-40 of 201", second_page.text)
            self.assertIn("ROW-0021", second_page.text)
            self.assertIn("ROW-0040", second_page.text)
            self.assertNotIn("ROW-0001", second_page.text)

            clamped_page = self.client.get(
                f"/workspaces/{workspace_id}/summary?page=999"
            )
            self.assertEqual(
                clamped_page.text.count("data-readiness-row"),
                1,
            )
            self.assertIn("Rows 201-201 of 201", clamped_page.text)
            self.assertIn("Page 11 of 11", clamped_page.text)
            self.assertIn("ROW-0201", clamped_page.text)

            filtered_page = self.client.get(
                f"/workspaces/{workspace_id}/summary",
                params={
                    "status": "blocked",
                    "dataset": sample_row.dataset,
                    "page": "3",
                    "page_size": "50",
                },
            )
            self.assertEqual(
                filtered_page.text.count("data-readiness-row"),
                20,
            )
            self.assertIn("Rows 101-120 of 120", filtered_page.text)
            self.assertIn("Page 3 of 3", filtered_page.text)
            self.assertNotIn("data-readiness-next", filtered_page.text)
            previous_match = re.search(
                r'href="([^"]+)" data-readiness-previous',
                filtered_page.text,
            )
            assert previous_match is not None
            previous_query = parse_qs(
                urlsplit(unescape(previous_match.group(1))).query
            )
            self.assertEqual(previous_query["status"], ["blocked"])
            self.assertEqual(
                previous_query["dataset"],
                [sample_row.dataset],
            )
            self.assertEqual(previous_query["page"], ["2"])
            self.assertEqual(previous_query["page_size"], ["50"])

        self.assertEqual(len(self.readiness_calls), 2)
        readiness_requests = self.readiness_calls[-1][2]
        self.assertTrue(readiness_requests)
        self.assertEqual(
            {item.model for item in readiness_requests},
            {"res.partner"},
        )
        self.assertTrue(all(item.domain for item in readiness_requests))
        evidence = self.client.get(
            f"/workspaces/{workspace_id}/summary/manifest"
        )
        self.assertEqual(evidence.status_code, 200)
        self.assertIn(
            "application/json",
            evidence.headers["content-type"],
        )
        with patch("impodo.web.routers.preflight.write_review_workbook") as builder:
            builder.side_effect = lambda _manifest, workbook, **_options: Path(
                workbook
            ).write_bytes(b"review package")
            packaged = self.client.post(
                f"/workspaces/{workspace_id}/summary/package",
                data={"csrf_token": self.csrf},
                headers=POST_HEADERS,
                follow_redirects=False,
            )
            review_evidence = builder.call_args.kwargs["review_evidence"]
            self.assertIsNotNone(review_evidence)
            current_report = self.app.state.context.preflight.current_report(
                workspace_id
            )
            self.assertEqual(
                len(review_evidence.records),
                current_report.total_count,
            )
            self.assertEqual(
                review_evidence.frozen_input_hash,
                current_report.frozen_input_hash,
            )
        self.assertEqual(packaged.status_code, 303)
        packaged_page = self.client.get(packaged.headers["location"])
        self.assertIn("Download review workbook", packaged_page.text)
        self.assertIn("Recreate review workbook", packaged_page.text)
        workbook = self.client.get(
            f"/workspaces/{workspace_id}/summary/workbook"
        )
        self.assertEqual(workbook.status_code, 200)
        self.assertIn(
            "spreadsheetml.sheet",
            workbook.headers["content-type"],
        )

        workspace_state = self.app.state.context.workspace_states.repository.get(workspace_id)

        def company_schema_reader(workspace_state, api_key):
            self.schema_calls.append((workspace_state.workspace_id, api_key))
            snapshot = _browser_schema(workspace_state)
            company = replace(
                snapshot.models["res.partner"],
                model="res.company",
            )
            return replace(snapshot, models={"res.company": company})

        self.app.state.context.schema_reader = company_schema_reader
        changed_scope = self.client.post(
            f"/workspaces/{workspace_id}/schema",
            data={
                "csrf_token": self.csrf,
                "revision": str(workspace_state.revision),
                "permitted_models": "res.company",
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )
        self.assertEqual(changed_scope.status_code, 303)
        workspace_state = self.app.state.context.workspace_states.repository.get(workspace_id)
        self.assertEqual(workspace_state.intended_models, ("res.company",))
        self.assertIsNone(workspace_state.mapping_version)
        self.assertEqual(workspace_state.approval_status.value, "INVALIDATED")
        refreshed_schema = (
            self.app.state.context.schema_workspace.schemas.get_odoo_schema_catalog(
                workspace_id
            )
        )
        self.assertIsNotNone(refreshed_schema)
        self.assertEqual(
            tuple(model.name for model in refreshed_schema.models),
            ("res.company",),
        )
        self.assertIsNone(
            self.app.state.context.schema_workspace.schemas.get_schema_governance(workspace_id)
        )
        self.assertIsNone(
            self.app.state.context.mapping_workspace.mappings.get_mapping_revision(workspace_id)
        )
        self.assertIsNone(
            self.app.state.context.preflight.current_staging(workspace_id)
        )
        self.assertIsNotNone(
            self.app.state.context.preflight.staging.get_canonical_staging_run(
                workspace_id,
                staging.run_id,
            )
        )
        self.assertIsNotNone(
            self.app.state.context.mapping_workspace.mappings.get_mapping_working_draft(
                workspace_id
            )
        )
        stale_mapping = self.client.get(
            f"/workspaces/{workspace_id}/mapping"
        )
        self.assertIn(
            "Your source tables or Odoo choices changed, so older matching work was not loaded",
            stale_mapping.text,
        )

    def test_saved_key_is_not_reused_after_target_change(self) -> None:
        created = self._post(
            "/projects/new",
            {
                "csrf_token": self.csrf,
                "display_name": "Credential binding",
                "source_mode": "FILE",
                "source_system_identity": "Other",
            },
        )
        workspace_id = _created_workspace_id(self.app, created)
        workspace_state = self._complete_setup_before_target(workspace_id)
        local = self._post(
            f"/workspaces/{workspace_id}/target",
            {
                "csrf_token": self.csrf,
                "revision": str(workspace_state.revision),
                "odoo_connection_mode": "LOCAL",
                "odoo_base_url": "http://127.0.0.1:8069",
                "odoo_database": "odoo19_local",
                "api_key": "local-only-key",
                "action": "test",
            },
        )
        self.assertEqual(local.status_code, 303)

        remote = self.client.post(
            f"/workspaces/{workspace_id}/target",
            data={
                "csrf_token": self.csrf,
                "revision": str(workspace_state.revision + 1),
                "odoo_connection_mode": "REMOTE",
                "odoo_base_url": "https://odoo.example.com",
                "odoo_database": "odoo_review",
                "action": "test",
            },
            headers=POST_HEADERS,
        )
        self.assertEqual(remote.status_code, 200)
        self.assertIn(
            "Enter an Odoo access key for this remote target.",
            remote.text,
        )
        self.assertEqual(len(self.connection_calls), 1)
        self.assertEqual(self.connection_calls[0][1], "local-only-key")
        self.assertEqual(self.secrets.values, {})

    def test_first_identity_mapping_save_persists_without_validation(self) -> None:
        workspace_id, dataset, business_key = self._mapping_ready_workspace(
            scalar_field_count=30
        )
        source_identity = dataset.columns[0]
        context = self.app.state.context

        saved = self.client.post(
            f"/workspaces/{workspace_id}/mapping/save",
            json={
                "entries": [
                    ["csrf_token", self.csrf],
                    ["action", "save_progress"],
                    ["expected_parent_version", ""],
                    ["expected_working_draft_version", ""],
                    ["editable_dataset_id", dataset.dataset_id],
                    ["target_model_0", "res.partner"],
                    ["mode_0", "upsert"],
                    ["on_existing_0", "block"],
                    ["source_identity_0", source_identity.stable_key],
                    ["business_key_0", business_key.key_id],
                    [
                        "identity_source_0_0",
                        source_identity.stable_key,
                    ],
                ]
            },
            headers={
                **POST_HEADERS,
                "X-CSRF-Token": self.csrf,
            },
        )

        self.assertEqual(saved.status_code, 200)
        self.assertEqual(
            saved.json()["message"],
            "Progress saved. Check matches when ready.",
        )
        working = context.mapping_workspace.mappings.get_mapping_working_draft(workspace_id)
        self.assertIsNotNone(working)
        self.assertEqual(working.version, 1)
        self.assertEqual(
            working.definition.datasets[0].target_identity[0].source_column_keys,
            (source_identity.stable_key,),
        )
        self.assertIsNone(context.mapping_workspace.mappings.get_mapping_revision(workspace_id))

    def test_selection_choices_load_and_save_from_the_mapping_dialog(self) -> None:
        workspace_id, dataset, business_key = self._mapping_ready_workspace(
            scalar_field_count=1,
            selection_field=True,
        )
        source_identity, source_value = dataset.columns

        page = self.client.get(f"/workspaces/{workspace_id}/mapping")
        self.assertEqual(page.status_code, 200)
        self.assertIn("data-value-match-dialog", page.text)
        self.assertIn("Match values", page.text)
        self.assertIn("Choice field", page.text)
        self.assertIn("2 choice(s) captured from Odoo", page.text)
        self.assertIn("Match source choices to Odoo", page.text)
        self.assertIn("How must source choices be covered?", page.text)
        self.assertIn("French (France) — fr_FR", page.text)
        self.assertNotIn("datalist", page.text)
        mapping_script = self.client.get("/static/app.js")
        self.assertIn(
            'event.target.closest?.("[data-open-value-match]")',
            mapping_script.text,
        )
        self.assertIn(
            'mappingForm.addEventListener("change", (event) => {',
            mapping_script.text,
        )
        self.assertNotIn(
            'for (const trigger of mappingForm.querySelectorAll(\n'
            '      "[data-open-value-match]"',
            mapping_script.text,
        )
        with patch(
            "impodo.web.routers.mapping._source_value_choices",
            return_value=(
                {"value": "French", "count": 12},
                {"value": "German", "count": 4},
            ),
        ):
            choices = self.client.post(
                f"/workspaces/{workspace_id}/mapping/value-choices",
                data={
                    "csrf_token": self.csrf,
                    "kind": "scalar",
                    "dataset_id": dataset.dataset_id,
                    "source_column_key": source_value.stable_key,
                    "target_model": "res.partner",
                    "target_field": "field_0000",
                    "business_key_id": "",
                },
                headers=POST_HEADERS,
            )

        self.assertEqual(choices.status_code, 200)
        self.assertEqual(
            choices.json()["target_choices"],
            [
                {"value": "fr_FR", "label": "French (France)"},
                {"value": "de_DE", "label": "German"},
            ],
        )
        entries = [
            ["csrf_token", self.csrf],
            ["action", "save_progress"],
            ["expected_parent_version", ""],
            ["expected_working_draft_version", ""],
            ["editable_dataset_id", dataset.dataset_id],
            ["target_model_0", "res.partner"],
            ["mode_0", "upsert"],
            ["on_existing_0", "block"],
            ["source_identity_0", source_identity.stable_key],
            ["business_key_0", business_key.key_id],
            ["identity_source_0_0", source_identity.stable_key],
            ["visible_scalar_target_0", "field_0000"],
            ["scalar_value_source_0_1", "source"],
            ["scalar_source_0_1", source_value.stable_key],
            ["scalar_type_0_1", "string"],
            ["scalar_case_0_1", "preserve"],
            ["scalar_compare_0_1", "1"],
            ["scalar_null_0_1", "distinct"],
            [
                "scalar_value_matches_0_1",
                '[{"source_value":"French","target_value":"fr_FR"}]',
            ],
        ]
        saved = self.client.post(
            f"/workspaces/{workspace_id}/mapping/save",
            json={"entries": entries},
            headers={**POST_HEADERS, "X-CSRF-Token": self.csrf},
        )

        self.assertEqual(saved.status_code, 200)
        working = self.app.state.context.mapping_workspace.mappings.get_mapping_working_draft(
            workspace_id
        )
        self.assertEqual(
            working.definition.datasets[0].fields[0].value_mappings,
            (ValueMapping("French", "fr_FR"),),
        )
        self.assertEqual(
            working.definition.datasets[0].fields[0].categorical_policy,
            CategoricalCoveragePolicy.EXPLICIT_VALUE_MATCH,
        )

    def test_relationship_choices_request_the_inline_read_key_recovery(
        self,
    ) -> None:
        workspace_id, dataset, business_key = self._mapping_ready_workspace(
            scalar_field_count=0,
            relationship_field_count=1,
        )
        source_value = dataset.columns[1]
        with (
            patch(
                "impodo.web.routers.mapping._source_value_choices",
                return_value=({"value": "FRA", "count": 3},),
            ),
            patch(
                "impodo.web.routers.mapping._relationship_value_choices",
                side_effect=OdooReadCredentialMissingError(
                    "Enter the read-only Odoo key to load current countries"
                ),
            ),
        ):
            response = self.client.post(
                f"/workspaces/{workspace_id}/mapping/value-choices",
                data={
                    "csrf_token": self.csrf,
                    "kind": "relationship",
                    "dataset_id": dataset.dataset_id,
                    "source_column_key": source_value.stable_key,
                    "target_model": "res.partner",
                    "target_field": "relation_0000",
                    "business_key_id": business_key.key_id,
                },
                headers=POST_HEADERS,
            )

        self.assertEqual(response.status_code, 422, response.text)
        self.assertTrue(
            response.json()["read_credential_required"],
            response.text,
        )
        self.assertEqual(
            response.json()["read_credential_failure_code"],
            OdooReadFailureCode.READ_KEY_MISSING.value,
        )
        self.assertNotIn("Odoo API key", response.text)

        script = self.client.get("/static/app.js")
        self.assertIn("payload.read_credential_required === true", script.text)
        self.assertIn('"impodo:read-credential-saved"', script.text)

    def test_relationship_choices_are_read_once_without_exposing_odoo_ids(
        self,
    ) -> None:
        workspace_id, dataset, business_key = self._mapping_ready_workspace(
            scalar_field_count=0,
            relationship_field_count=1,
        )
        source_value = dataset.columns[1]
        context = self.app.state.context
        calls = []

        def readiness_reader(workspace_state, metadata_requests, record_requests):
            calls.append((metadata_requests, record_requests))
            metadata = _browser_schema(workspace_state)
            return metadata, RecordSnapshot(
                fingerprint=metadata.fingerprint,
                records={
                    "res.partner": (
                        TargetRecord("res.partner", 10, {"ref": "FR"}),
                        TargetRecord("res.partner", 11, {"ref": "DE"}),
                        TargetRecord("res.partner", 12, {"ref": "BE"}),
                        TargetRecord("res.partner", 13, {"ref": "BE"}),
                    )
                },
                requested_fields={"res.partner": ("ref",)},
            )

        context.readiness_reader = readiness_reader
        with patch(
            "impodo.web.routers.mapping._source_value_choices",
            return_value=({"value": "FRA", "count": 3},),
        ):
            response = self.client.post(
                f"/workspaces/{workspace_id}/mapping/value-choices",
                data={
                    "csrf_token": self.csrf,
                    "kind": "relationship",
                    "dataset_id": dataset.dataset_id,
                    "source_column_key": source_value.stable_key,
                    "target_model": "res.partner",
                    "target_field": "relation_0000",
                    "business_key_id": business_key.key_id,
                },
                headers=POST_HEADERS,
            )

        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(len(calls), 1)
        self.assertEqual(len(calls[0][1]), 1)
        self.assertEqual(
            response.json()["target_choices"],
            [
                {"value": "DE", "label": "DE"},
                {"value": "FR", "label": "FR"},
            ],
        )
        self.assertEqual(response.json()["ambiguous_values"], ["BE"])
        self.assertNotIn("odoo_id", response.text)
        source_identity = dataset.columns[0]
        saved = self.client.post(
            f"/workspaces/{workspace_id}/mapping/save",
            json={
                "entries": [
                    ["csrf_token", self.csrf],
                    ["action", "save_progress"],
                    ["expected_parent_version", ""],
                    ["expected_working_draft_version", ""],
                    ["editable_dataset_id", dataset.dataset_id],
                    ["target_model_0", "res.partner"],
                    ["mode_0", "upsert"],
                    ["on_existing_0", "block"],
                    ["source_identity_0", source_identity.stable_key],
                    ["business_key_0", business_key.key_id],
                    ["identity_source_0_0", source_identity.stable_key],
                    ["visible_relation_target_0", "relation_0000"],
                    ["relation_source_0_0", source_value.stable_key],
                    ["relation_origin_0_0", "target_catalog"],
                    ["relation_key_0_0", business_key.key_id],
                    ["relation_operation_0_0", "replace"],
                    ["relation_compare_0_0", "1"],
                    ["relation_missing_0_0", "error"],
                    ["relation_ambiguous_0_0", "error"],
                    ["relation_null_0_0", "distinct"],
                    ["relation_separator_0_0", ";"],
                    [
                        "relation_value_matches_0_0",
                        '[{"source_value":"FRA","target_value":"FR"}]',
                    ],
                ]
            },
            headers={**POST_HEADERS, "X-CSRF-Token": self.csrf},
        )

        self.assertEqual(saved.status_code, 200)
        working = context.mapping_workspace.mappings.get_mapping_working_draft(workspace_id)
        self.assertEqual(
            working.definition.datasets[0]
            .relationships[0]
            .resolver.value_mappings,
            (ValueMapping("FRA", "FR"),),
        )
        self.assertEqual(
            working.definition.datasets[0].relationships[0].categorical_policy,
            CategoricalCoveragePolicy.EXPLICIT_KEY_MATCH,
        )

    def test_matching_rule_without_description_uses_odoo_field_label(self) -> None:
        workspace_id, _dataset, _business_key = self._mapping_ready_workspace(
            scalar_field_count=0,
            relationship_field_count=1,
            business_key_description="",
        )

        page = self.client.get(f"/workspaces/{workspace_id}/mapping")

        self.assertEqual(page.status_code, 200)
        self.assertIn(">\n                        Reference\n", page.text)
        self.assertIn(">\n                            Reference\n", page.text)
        self.assertNotIn("Confirmed matching rule", page.text)

    def test_country_matching_uses_reviewed_code_without_schema_recapture(
        self,
    ) -> None:
        workspace_id, dataset, business_key = self._mapping_ready_workspace(
            scalar_field_count=0,
            relationship_field_count=1,
            relationship_model="res.country",
        )
        source_identity, source_country = dataset.columns
        context = self.app.state.context
        calls = []

        def readiness_reader(workspace_state, metadata_requests, record_requests):
            calls.append((metadata_requests, record_requests))
            available = _browser_schema(workspace_state)
            metadata = replace(
                available,
                models={
                    "res.country": ModelMetadata(
                        model="res.country",
                        description="Country",
                        fields={
                            "code": FieldMetadata(
                                name="code",
                                type="char",
                                label="Country Code",
                                required=True,
                            ),
                            "name": FieldMetadata(
                                name="name",
                                type="char",
                                label="Country Name",
                                required=True,
                            ),
                        },
                    )
                },
            )
            return metadata, RecordSnapshot(
                fingerprint=metadata.fingerprint,
                records={
                    "res.country": (
                        TargetRecord(
                            "res.country",
                            1,
                            {"code": "FR", "name": "France"},
                        ),
                        TargetRecord(
                            "res.country",
                            2,
                            {"code": "BE", "name": "Belgium"},
                        ),
                    )
                },
                requested_fields={"res.country": ("code", "name")},
            )

        context.readiness_reader = readiness_reader
        page = self.client.get(f"/workspaces/{workspace_id}/mapping")

        self.assertEqual(page.status_code, 200)
        schema = context.schema_workspace.schemas.get_odoo_schema_catalog(workspace_id)
        self.assertNotEqual(
            schema.content_hash,
            target_identity_hash(
                connection_mode="LOCAL",
                base_url="http://127.0.0.1:8069",
                database="odoo19_local",
            ),
        )
        self.assertIn("Country code — recommended", page.text)
        self.assertIn(
            'value="odoo-standard:res.country:code" selected',
            page.text,
        )
        self.assertIn("data-refresh-value-match", page.text)
        mapping_script = self.client.get("/static/app.js")
        self.assertIn("Using saved Odoo values checked", mapping_script.text)
        self.assertIn(
            'body.set("refresh", refresh ? "1" : "0")',
            mapping_script.text,
        )
        with patch(
            "impodo.web.routers.mapping._source_value_choices",
            return_value=({"value": "FRA", "count": 3},),
        ):
            request_data = {
                "csrf_token": self.csrf,
                "kind": "relationship",
                "dataset_id": dataset.dataset_id,
                "source_column_key": source_country.stable_key,
                "target_model": "res.partner",
                "target_field": "relation_0000",
                "business_key_id": "odoo-standard:res.country:code",
            }
            choices = self.client.post(
                f"/workspaces/{workspace_id}/mapping/value-choices",
                data=request_data,
                headers=POST_HEADERS,
            )
            reused = self.client.post(
                f"/workspaces/{workspace_id}/mapping/value-choices",
                data=request_data,
                headers=POST_HEADERS,
            )
            refreshed = self.client.post(
                f"/workspaces/{workspace_id}/mapping/value-choices",
                data={**request_data, "refresh": "1"},
                headers=POST_HEADERS,
            )

        self.assertEqual(choices.status_code, 200, choices.text)
        self.assertEqual(
            choices.json()["target_choices"],
            [
                {"value": "BE", "label": "Belgium (BE)"},
                {"value": "FR", "label": "France (FR)"},
            ],
        )
        self.assertFalse(choices.json()["target_choices_reused"])
        self.assertTrue(choices.json()["target_checked_at"])
        self.assertEqual(reused.status_code, 200)
        self.assertTrue(reused.json()["target_choices_reused"])
        self.assertEqual(refreshed.status_code, 200)
        self.assertFalse(refreshed.json()["target_choices_reused"])
        self.assertEqual(len(calls), 2)
        self.assertEqual(calls[0][0][0].model, "res.country")
        self.assertEqual(calls[0][0][0].fields, ("code", "name"))
        self.assertEqual(calls[0][1][0].model, "res.country")
        self.assertEqual(calls[0][1][0].fields, ("code", "name"))
        self.assertEqual(calls[0][1][0].limit, 2001)
        self.assertNotIn("odoo_id", choices.text)

        saved = self.client.post(
            f"/workspaces/{workspace_id}/mapping/save",
            json={
                "entries": [
                    ["csrf_token", self.csrf],
                    ["action", "save_progress"],
                    ["expected_parent_version", ""],
                    ["expected_working_draft_version", ""],
                    ["editable_dataset_id", dataset.dataset_id],
                    ["target_model_0", "res.partner"],
                    ["mode_0", "upsert"],
                    ["on_existing_0", "block"],
                    ["source_identity_0", source_identity.stable_key],
                    ["business_key_0", business_key.key_id],
                    ["identity_source_0_0", source_identity.stable_key],
                    ["visible_relation_target_0", "relation_0000"],
                    ["relation_source_0_0", source_country.stable_key],
                    ["relation_origin_0_0", "target_catalog"],
                    [
                        "relation_key_0_0",
                        "odoo-standard:res.country:code",
                    ],
                    ["relation_operation_0_0", "replace"],
                    ["relation_compare_0_0", "1"],
                    ["relation_missing_0_0", "error"],
                    ["relation_ambiguous_0_0", "error"],
                    ["relation_null_0_0", "distinct"],
                    ["relation_separator_0_0", ";"],
                    [
                        "relation_value_matches_0_0",
                        '[{"source_value":"FRA","target_value":"FR"}]',
                    ],
                ]
            },
            headers={**POST_HEADERS, "X-CSRF-Token": self.csrf},
        )

        self.assertEqual(saved.status_code, 200)
        working = context.mapping_workspace.mappings.get_mapping_working_draft(workspace_id)
        resolver = working.definition.datasets[0].relationships[0].resolver
        self.assertEqual(
            resolver.key_mappings,
            (ReferenceKeyMapping(source_country.stable_key, "code"),),
        )
        self.assertEqual(
            resolver.value_mappings,
            (ValueMapping("FRA", "FR"),),
        )

    def test_reviewed_reference_matching_is_not_country_specific(self) -> None:
        for related_model, label in (
            ("res.lang", "Language code — recommended"),
            ("res.currency", "Currency code — recommended"),
        ):
            with self.subTest(related_model=related_model):
                workspace_id, _dataset, _business_key = self._mapping_ready_workspace(
                    scalar_field_count=0,
                    relationship_field_count=1,
                    relationship_model=related_model,
                )

                page = self.client.get(f"/workspaces/{workspace_id}/mapping")

                self.assertEqual(page.status_code, 200)
                self.assertIn(label, page.text)
                self.assertNotIn("No matching rule available", page.text)

    def test_uncaptured_many2one_offers_an_explicit_name_matching_rule(
        self,
    ) -> None:
        workspace_id, _dataset, _business_key = self._mapping_ready_workspace(
            scalar_field_count=0,
            relationship_field_count=1,
            relationship_model="res.company",
        )

        page = self.client.get(f"/workspaces/{workspace_id}/mapping")

        self.assertEqual(page.status_code, 200)
        self.assertIn("Odoo record name", page.text)
        self.assertIn("confirm", page.text)
        self.assertIn("Match values", page.text)
        self.assertNotIn('name="relation_key_0_0" disabled', page.text)
        self.assertNotIn("No matching rule available", page.text)

    def test_product_uom_choices_are_fetched_as_bounded_supporting_data(
        self,
    ) -> None:
        workspace_id, dataset, business_key = self._mapping_ready_workspace(
            scalar_field_count=0,
            relationship_field_count=1,
            relationship_model="uom.uom",
            target_model="product.template",
            relationship_field_names=("uom_id",),
        )
        source_identity, source_uom = dataset.columns
        context = self.app.state.context
        calls = []

        def readiness_reader(workspace_state, metadata_requests, record_requests):
            calls.append((metadata_requests, record_requests))
            available = _browser_schema(workspace_state)
            metadata = replace(
                available,
                models={
                    "uom.uom": ModelMetadata(
                        model="uom.uom",
                        description="Unit of Measure",
                        fields={
                            "name": FieldMetadata(
                                name="name",
                                type="char",
                                label="Unit of Measure",
                                required=True,
                            ),
                        },
                    )
                },
            )
            return metadata, RecordSnapshot(
                fingerprint=metadata.fingerprint,
                records={
                    "uom.uom": (
                        TargetRecord("uom.uom", 41, {"name": "Units"}),
                        TargetRecord("uom.uom", 42, {"name": "kg"}),
                        TargetRecord("uom.uom", 43, {"name": "Hours"}),
                    )
                },
                requested_fields={"uom.uom": ("name",)},
            )

        context.readiness_reader = readiness_reader
        page = self.client.get(f"/workspaces/{workspace_id}/mapping")
        candidate = re.search(
            r'<option value="([^"]+)"[^>]*>\s*Odoo record name',
            page.text,
        )
        self.assertIsNotNone(candidate, page.text)
        candidate_key_id = candidate.group(1)
        self.assertNotIn(f'value="{candidate_key_id}" selected', page.text)

        with patch(
            "impodo.web.routers.mapping._source_value_choices",
            return_value=({"value": "Unit", "count": 3},),
        ):
            choices = self.client.post(
                f"/workspaces/{workspace_id}/mapping/value-choices",
                data={
                    "csrf_token": self.csrf,
                    "kind": "relationship",
                    "dataset_id": dataset.dataset_id,
                    "source_column_key": source_uom.stable_key,
                    "target_model": "product.template",
                    "target_field": "uom_id",
                    "business_key_id": candidate_key_id,
                },
                headers=POST_HEADERS,
            )

        self.assertEqual(choices.status_code, 200, choices.text)
        self.assertEqual(
            choices.json()["target_choices"],
            [
                {"value": "Hours", "label": "Hours"},
                {"value": "kg", "label": "kg"},
                {"value": "Units", "label": "Units"},
            ],
        )
        self.assertEqual(len(calls), 1)
        self.assertEqual(calls[0][0][0].model, "uom.uom")
        self.assertEqual(calls[0][0][0].fields, ("name",))
        self.assertEqual(calls[0][1][0].fields, ("name",))
        self.assertEqual(calls[0][1][0].limit, 2001)
        self.assertNotIn("odoo_id", choices.text)
        self.assertNotIn("41", choices.text)

        schema = context.queries.get_odoo_schema_catalog(workspace_id)
        lookup = context.supporting_lookups.current(
            workspace_id,
            relation_model="uom.uom",
            key_fields=("name",),
            scope_fields=(),
            display_field="name",
            target_hash=schema.connection_target_hash,
            read_credential_binding_hash=schema.read_credential_binding_hash,
            read_principal_hash=schema.read_principal_hash,
            read_context_hash=schema.read_context_hash,
            actor=context.actor,
        )
        self.assertIsNotNone(lookup)
        context.supporting_lookups.capture(
            workspace_id,
            relation_model=lookup.relation_model,
            key_fields=lookup.key_fields,
            scope_fields=lookup.scope_fields,
            display_field=lookup.display_field,
            field_contracts=lookup.field_contracts,
            target_hash=lookup.target_hash,
            read_credential_binding_hash=lookup.read_credential_binding_hash,
            read_principal_hash=lookup.read_principal_hash,
            read_permission_hash="sha256:" + "9" * 64,
            read_context_hash=lookup.read_context_hash,
            captured_at=lookup.captured_at,
            choices=lookup.choices,
            ambiguous_values=lookup.ambiguous_values,
            actor=context.actor,
        )

        saved = self.client.post(
            f"/workspaces/{workspace_id}/mapping/save",
            json={
                "entries": [
                    ["csrf_token", self.csrf],
                    ["action", "save_progress"],
                    ["expected_parent_version", ""],
                    ["expected_working_draft_version", ""],
                    ["editable_dataset_id", dataset.dataset_id],
                    ["target_model_0", "product.template"],
                    ["mode_0", "upsert"],
                    ["on_existing_0", "block"],
                    ["source_identity_0", source_identity.stable_key],
                    ["business_key_0", business_key.key_id],
                    ["identity_source_0_0", source_identity.stable_key],
                    ["visible_relation_target_0", "uom_id"],
                    ["relation_source_0_0", source_uom.stable_key],
                    ["relation_origin_0_0", "target_catalog"],
                    ["relation_key_0_0", candidate_key_id],
                    ["relation_operation_0_0", "replace"],
                    ["relation_compare_0_0", "1"],
                    ["relation_missing_0_0", "error"],
                    ["relation_ambiguous_0_0", "error"],
                    ["relation_null_0_0", "distinct"],
                    ["relation_separator_0_0", ";"],
                    [
                        "relation_value_matches_0_0",
                        '[{"source_value":"Unit","target_value":"Units"}]',
                    ],
                ]
            },
            headers={**POST_HEADERS, "X-CSRF-Token": self.csrf},
        )

        self.assertEqual(saved.status_code, 200, saved.text)
        working = context.mapping_workspace.mappings.get_mapping_working_draft(
            workspace_id
        )
        schema = context.queries.get_odoo_schema_catalog(workspace_id)
        governance = context.queries.get_schema_governance(workspace_id)
        selection = context.queries.get_mapping_source_selection(workspace_id)
        supporting = context.mapping_workspace._current_supporting_references(
            workspace_id,
            working.definition,
            schema,
        )
        validation = context.mapping_workspace.validator.validate(
            working.definition,
            selection,
            schema,
            governance,
            supporting,
        )
        codes = {item.code for item in validation.issues}
        self.assertEqual(len(supporting), 1)
        self.assertEqual(supporting[0].relation_model, "uom.uom")
        self.assertEqual(supporting[0].field_contracts[0].name, "name")
        self.assertNotIn("MAPPING_TARGET_MODEL_UNKNOWN", codes)
        self.assertNotIn("MAPPING_BUSINESS_KEY_NOT_GOVERNED", codes)

    def test_relationship_catalog_is_searchable_and_progressively_disclosed(
        self,
    ) -> None:
        workspace_id, dataset, _business_key = self._mapping_ready_workspace(
            scalar_field_count=1,
            relationship_field_count=51,
        )
        source_identity, source_value = dataset.columns
        context = self.app.state.context
        initial = context.mapping_workspace.save_working_draft(
            workspace_id,
            datasets=(
                DatasetMapping(
                    dataset_id=dataset.dataset_id,
                    target_model="res.partner",
                    mode=MappingTargetMode.UPSERT,
                    source_identity_column_keys=(source_identity.stable_key,),
                    target_identity=(
                        IdentityComponentMapping(
                            source_column_keys=(source_identity.stable_key,),
                            target_fields=("ref",),
                        ),
                    ),
                    relationships=(
                        RelationshipMapping(
                            target_field="relation_0050",
                            kind="many2one",
                            source_column_keys=(source_value.stable_key,),
                            resolver=RelationshipResolver(
                                origin=ResolverOrigin.TARGET_CATALOG,
                                model="res.partner",
                                key_mappings=(
                                    ReferenceKeyMapping(
                                        source_column_key=source_value.stable_key,
                                        target_field="ref",
                                    ),
                                ),
                            ),
                        ),
                    ),
                ),
            ),
            expected_version=None,
            actor=context.actor,
        )
        self.assertEqual(initial.version, 1)

        page = self.client.get(f"/workspaces/{workspace_id}/mapping")
        self.assertEqual(page.status_code, 200)
        self.assertEqual(page.text.count("data-relation-field-row"), 3)
        self.assertIn("Showing 3 of 51 linked fields", page.text)
        self.assertIn('data-relation-page-size="3"', page.text)
        self.assertLess(
            page.text.index("relation_0050"),
            page.text.index("relation_0000"),
        )
        self.assertNotIn("relation_0002</code>", page.text)

        expanded = self.client.get(
            f"/workspaces/{workspace_id}/mapping?relation_page_size=20"
        )
        self.assertEqual(expanded.text.count("data-relation-field-row"), 20)
        self.assertIn('data-relation-page-size="20"', expanded.text)

        searched = self.client.get(
            f"/workspaces/{workspace_id}/mapping?relation_query=relation_0049"
        )
        self.assertEqual(searched.status_code, 200)
        self.assertEqual(searched.text.count("data-relation-field-row"), 1)
        self.assertIn("Linked Field 0049", searched.text)
        self.assertIn("Showing 1 of 1 linked fields", searched.text)

        fragment = self.client.get(
            f"/workspaces/{workspace_id}/mapping/field-catalog"
            "?catalog=relation&relation_query=relation_0049"
        )
        self.assertEqual(fragment.status_code, 200)
        self.assertEqual(fragment.text.count("data-relation-field-row"), 1)
        self.assertIn("Linked Field 0049", fragment.text)
        self.assertNotIn("data-scalar-field-catalog", fragment.text)
        self.assertNotIn("<main", fragment.text)
        self.assertIn("projection;dur=", fragment.headers["server-timing"])

        searched_by_model = self.client.get(
            f"/workspaces/{workspace_id}/mapping?relation_query=res.partner"
        )
        self.assertEqual(
            searched_by_model.text.count("data-relation-field-row"),
            3,
        )
        self.assertIn("Showing 3 of 51 linked fields", searched_by_model.text)

        last_page = self.client.get(
            f"/workspaces/{workspace_id}/mapping?relation_page=17"
        )
        self.assertEqual(last_page.text.count("data-relation-field-row"), 3)
        self.assertIn("relation_0049", last_page.text)

        rejected_size = self.client.get(
            f"/workspaces/{workspace_id}/mapping?relation_page_size=100"
        )
        self.assertEqual(
            rejected_size.text.count("data-relation-field-row"),
            3,
        )
        self.assertIn('data-relation-page-size="3"', rejected_size.text)

    def test_readonly_field_matches_are_hidden_and_recovered_as_one_decision(
        self,
    ) -> None:
        workspace_id, dataset, _business_key = self._mapping_ready_workspace(
            scalar_field_count=4,
            readonly_scalar_indexes=(1, 2),
        )
        source_identity, source_value = dataset.columns
        context = self.app.state.context
        mapping = DatasetMapping(
            dataset_id=dataset.dataset_id,
            target_model="res.partner",
            mode=MappingTargetMode.UPSERT,
            source_identity_column_keys=(source_identity.stable_key,),
            target_identity=(
                IdentityComponentMapping(
                    source_column_keys=(source_identity.stable_key,),
                    target_fields=("ref",),
                ),
            ),
            fields=(
                ScalarFieldMapping(
                    target_field="field_0001",
                    source_column_key=source_value.stable_key,
                ),
                ScalarFieldMapping(
                    target_field="field_0002",
                    source_column_key=source_value.stable_key,
                    compare=False,
                    validate_only=True,
                ),
            ),
        )
        _revision, validation = context.mapping_workspace.check_definition(
            workspace_id,
            datasets=(mapping,),
            expected_parent_version=None,
            expected_working_draft_version=None,
            actor=context.actor,
        )
        self.assertEqual(validation.status, MappingValidationStatus.INVALID)
        self.assertEqual(
            [
                item.target_field
                for item in validation.issues
                if item.code == "MAPPING_TARGET_FIELD_READONLY"
            ],
            ["field_0001"],
        )

        page = self.client.get(f"/workspaces/{workspace_id}/mapping")

        self.assertEqual(page.status_code, 200)
        self.assertIn("Odoo manages 1 selected field", page.text)
        self.assertIn("Remove this field match", page.text)
        self.assertIn("1 decision", page.text)
        self.assertNotIn('data-target-field="field_0001"', page.text)
        self.assertIn('data-target-field="field_0002"', page.text)
        self.assertIn('data-target-field="field_0003"', page.text)
        self.assertIn('name="scalar_value_source_0_3"', page.text)
        self.assertEqual(
            page.text.count("MAPPING_TARGET_FIELD_READONLY"),
            1,
        )

        recovered = self.client.post(
            f"/workspaces/{workspace_id}/mapping/save",
            data={
                "csrf_token": self.csrf,
                "action": "remove_readonly",
                "expected_parent_version": "1",
                "expected_working_draft_version": "1",
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )

        self.assertEqual(recovered.status_code, 303)
        recovered_page = self.client.get(recovered.headers["location"])
        self.assertIn(
            "Removed 1 Odoo-managed field match. Check matches again when ready.",
            recovered_page.text,
        )
        self.assertNotIn("Odoo manages 1 selected field", recovered_page.text)
        working = context.mapping_workspace.mappings.get_mapping_working_draft(
            workspace_id
        )
        self.assertEqual(working.version, 2)
        self.assertEqual(
            [item.target_field for item in working.definition.datasets[0].fields],
            ["field_0002"],
        )

    def test_verified_required_default_is_reviewed_and_confirmed_as_a_group(
        self,
    ) -> None:
        workspace_id, dataset, business_key = self._mapping_ready_workspace(
            scalar_field_count=1,
            required_scalar_indexes=(0,),
            verified_default_scalar_indexes=(0,),
            target_model="sale.order",
        )
        source_identity, _source_value = dataset.columns
        context = self.app.state.context
        _revision, validation = context.mapping_workspace.check_definition(
            workspace_id,
            datasets=(
                DatasetMapping(
                    dataset_id=dataset.dataset_id,
                    target_model="sale.order",
                    mode=MappingTargetMode.UPSERT,
                    source_identity_column_keys=(source_identity.stable_key,),
                    target_identity=(
                        IdentityComponentMapping(
                            source_column_keys=(source_identity.stable_key,),
                            target_fields=business_key.key_fields,
                        ),
                    ),
                ),
            ),
            expected_parent_version=None,
            expected_working_draft_version=None,
            actor=context.actor,
        )
        self.assertEqual(validation.status, MappingValidationStatus.INVALID)

        page = self.client.get(
            f"/workspaces/{workspace_id}/mapping?field_query=not-visible"
        )

        self.assertEqual(page.status_code, 200)
        self.assertIn('id="next-step-blockers"', page.text)
        self.assertGreater(
            page.text.index('id="next-step-blockers"'),
            page.text.index('id="semantic-validation"'),
        )
        self.assertLess(
            page.text.index('id="next-step-blockers"'),
            page.text.index('class="actions mapping-actions"'),
        )
        self.assertIn("You cannot continue yet — 1 reason", page.text)
        self.assertIn("Review 1 Odoo default", page.text)
        self.assertIn("Field 0000:", page.text)
        self.assertIn(
            'value="confirm_defaults"',
            page.text,
        )
        self.assertIn("Odoo default 0000", page.text)
        self.assertIn("Use 1 Odoo default", page.text)
        self.assertRegex(
            page.text,
            r'<button class="button primary"[^>]*disabled>'
            r"Confirm field matches</button>",
        )
        self.assertNotIn('data-target-field="field_0000"', page.text)

        decision = self.client.post(
            f"/workspaces/{workspace_id}/mapping/save",
            data={
                "csrf_token": self.csrf,
                "action": "confirm_defaults",
                "expected_working_draft_version": "1",
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )

        self.assertEqual(decision.status_code, 303)
        working = context.mapping_workspace.mappings.get_mapping_working_draft(
            workspace_id
        )
        self.assertIsNotNone(working)
        assert working is not None
        self.assertEqual(
            working.definition.datasets[0]
            .target_field_dispositions[0]
            .handling,
            TargetFieldHandling.ODOO_DEFAULT,
        )
        decision_page = self.client.get(decision.headers["location"])
        self.assertIn("Odoo will choose this value", decision_page.text)
        self.assertIn(
            'name="target_field_disposition_0"', decision_page.text
        )
        self.assertIn("Saved changes have not been checked yet", decision_page.text)
        self.assertIn("Keep working from the last check", decision_page.text)
        self.assertIn(
            "large_contacts: Field 0000 needs attention",
            decision_page.text,
        )
        self.assertIn("Confirmed 1 Odoo default.", decision_page.text)

        mapping_data = {
            "csrf_token": self.csrf,
            "editable_dataset_id": dataset.dataset_id,
            "target_model_0": "sale.order",
            "mode_0": "upsert",
            "on_existing_0": "block",
            "source_identity_0": source_identity.stable_key,
            "business_key_0": business_key.key_id,
            "identity_source_0_0": source_identity.stable_key,
            "visible_scalar_target_0": "field_0000",
            "target_field_disposition_0": "field_0000:odoo_default",
        }
        checked = self.client.post(
            f"/workspaces/{workspace_id}/mapping/save",
            data={
                **mapping_data,
                "action": "draft",
                "expected_parent_version": "1",
                "expected_working_draft_version": "2",
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )

        self.assertEqual(checked.status_code, 303)
        checked_page = self.client.get(checked.headers["location"])
        self.assertNotIn('id="next-step-blockers"', checked_page.text)
        self.assertIn("Confirm field matches", checked_page.text)
        current_revision = context.mapping_workspace.mappings.get_mapping_revision(
            workspace_id
        )
        current_working = (
            context.mapping_workspace.mappings.get_mapping_working_draft(workspace_id)
        )
        current_validation = (
            context.mapping_workspace.mappings.get_mapping_validation(
                workspace_id,
                current_revision.version,
            )
        )
        self.assertEqual(current_validation.status, MappingValidationStatus.VALID)
        self.assertEqual(current_validation.issues, ())

        submitted = self.client.post(
            f"/workspaces/{workspace_id}/mapping/save",
            json={
                "entries": [
                    *mapping_data.items(),
                    ["action", "submit"],
                    ["expected_parent_version", str(current_revision.version)],
                    [
                        "expected_working_draft_version",
                        str(current_working.version),
                    ],
                ]
            },
            headers={**POST_HEADERS, "X-CSRF-Token": self.csrf},
        )

        self.assertEqual(submitted.status_code, 200)
        self.assertEqual(
            submitted.json()["redirect_url"],
            f"/workspaces/{workspace_id}/prepare",
        )
        submitted_page = self.client.get(submitted.json()["redirect_url"])
        self.assertIn("Field matches confirmed", submitted_page.text)
        self.assertIn("Prepare all source rows", submitted_page.text)

    def test_verified_required_defaults_are_confirmed_in_one_action(
        self,
    ) -> None:
        workspace_id, dataset, business_key = self._mapping_ready_workspace(
            scalar_field_count=2,
            required_scalar_indexes=(0, 1),
            verified_default_scalar_indexes=(0, 1),
        )
        source_identity, _source_value = dataset.columns
        context = self.app.state.context
        _revision, validation = context.mapping_workspace.check_definition(
            workspace_id,
            datasets=(
                DatasetMapping(
                    dataset_id=dataset.dataset_id,
                    target_model="res.partner",
                    mode=MappingTargetMode.UPSERT,
                    source_identity_column_keys=(source_identity.stable_key,),
                    target_identity=(
                        IdentityComponentMapping(
                            source_column_keys=(source_identity.stable_key,),
                            target_fields=business_key.key_fields,
                        ),
                    ),
                ),
            ),
            expected_parent_version=None,
            expected_working_draft_version=None,
            actor=context.actor,
        )
        self.assertEqual(validation.status, MappingValidationStatus.INVALID)

        page = self.client.get(f"/workspaces/{workspace_id}/mapping")
        self.assertIn("Review 2 Odoo defaults", page.text)
        self.assertIn("Odoo default 0000", page.text)
        self.assertIn("Odoo default 0001", page.text)
        self.assertEqual(page.text.count('value="confirm_defaults"'), 1)

        decision = self.client.post(
            f"/workspaces/{workspace_id}/mapping/save",
            data={
                "csrf_token": self.csrf,
                "action": "confirm_defaults",
                "expected_working_draft_version": "1",
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )

        self.assertEqual(decision.status_code, 303)
        working = context.mapping_workspace.mappings.get_mapping_working_draft(
            workspace_id
        )
        self.assertEqual(
            [
                item.target_field
                for item in (
                    working.definition.datasets[0].target_field_dispositions
                )
            ],
            ["field_0000", "field_0001"],
        )

    def test_missing_required_defaults_are_checked_and_confirmed_in_one_action(
        self,
    ) -> None:
        workspace_id, dataset, business_key = self._mapping_ready_workspace(
            scalar_field_count=2,
            required_scalar_indexes=(0, 1),
        )
        source_identity, _source_value = dataset.columns
        context = self.app.state.context
        revision, validation = context.mapping_workspace.check_definition(
            workspace_id,
            datasets=(
                DatasetMapping(
                    dataset_id=dataset.dataset_id,
                    target_model="res.partner",
                    mode=MappingTargetMode.UPSERT,
                    source_identity_column_keys=(source_identity.stable_key,),
                    target_identity=(
                        IdentityComponentMapping(
                            source_column_keys=(source_identity.stable_key,),
                            target_fields=business_key.key_fields,
                        ),
                    ),
                ),
            ),
            expected_parent_version=None,
            expected_working_draft_version=None,
            actor=context.actor,
        )
        self.assertEqual(validation.status, MappingValidationStatus.INVALID)
        schema_before = context.queries.get_odoo_schema_catalog(workspace_id)
        governance_before = context.queries.get_schema_governance(workspace_id)
        calls = []

        def readiness_reader(workspace_state, metadata_requests, record_requests):
            calls.append((metadata_requests, record_requests))
            models = {model.name: model for model in schema_before.models}
            returned_models = {}
            defaults = {}
            for request in metadata_requests:
                model = models[request.model]
                fields = {field.name: field for field in model.fields}
                returned_models[request.model] = ModelMetadata(
                    model=request.model,
                    description=model.label,
                    fields={
                        field_name: FieldMetadata(
                            name=field_name,
                            type=fields[field_name].type,
                            label=fields[field_name].label,
                            required=fields[field_name].required,
                            readonly=fields[field_name].readonly,
                            relation=fields[field_name].relation,
                            relation_field=fields[field_name].relation_field,
                            selection=fields[field_name].selection,
                            stored=fields[field_name].stored,
                            computed=fields[field_name].computed,
                            has_inverse=fields[field_name].has_inverse,
                            related=fields[field_name].related,
                            translated=fields[field_name].translated,
                            company_dependent=fields[field_name].company_dependent,
                            searchable=fields[field_name].searchable,
                            sortable=fields[field_name].sortable,
                            exportable=fields[field_name].exportable,
                            digits=fields[field_name].digits,
                            currency_field=fields[field_name].currency_field,
                        )
                        for field_name in request.fields
                    },
                )
                defaults[request.model] = {
                    field_name: f"Current Odoo default {field_name}"
                    for field_name in request.fields
                }
            metadata = MetadataSnapshot(
                fingerprint=_browser_schema(workspace_state).fingerprint,
                models=returned_models,
                create_defaults=defaults,
            )
            return metadata, RecordSnapshot(
                fingerprint=metadata.fingerprint,
                records={},
                requested_fields={},
            )

        context.readiness_reader = readiness_reader
        page = self.client.get(f"/workspaces/{workspace_id}/mapping")

        self.assertIn("Let Odoo decide for 2 required fields", page.text)
        self.assertIn("read-only for this exact Odoo target", page.text)
        self.assertEqual(page.text.count('value="refresh_defaults"'), 1)

        decision = self.client.post(
            f"/workspaces/{workspace_id}/mapping/save",
            data={
                "csrf_token": self.csrf,
                "action": "refresh_defaults",
                "expected_working_draft_version": "1",
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )

        self.assertEqual(decision.status_code, 303, decision.text)
        self.assertEqual(len(calls), 1)
        self.assertEqual(len(calls[0][0]), 1)
        self.assertEqual(
            calls[0][0][0].fields,
            ("field_0000", "field_0001"),
        )
        self.assertEqual(calls[0][1], ())
        schema_after = context.queries.get_odoo_schema_catalog(workspace_id)
        self.assertEqual(schema_after.content_hash, schema_before.content_hash)
        self.assertEqual(
            context.queries.get_schema_governance(workspace_id),
            governance_before,
        )
        recovered_fields = {
            field.name: field for field in schema_after.models[0].fields
        }
        self.assertEqual(
            recovered_fields["field_0000"].create_default_value,
            "Current Odoo default field_0000",
        )
        self.assertEqual(
            recovered_fields["field_0001"].create_default_value,
            "Current Odoo default field_0001",
        )
        working = context.mapping_workspace.mappings.get_mapping_working_draft(
            workspace_id
        )
        self.assertEqual(
            [
                item.target_field
                for item in working.definition.datasets[0].target_field_dispositions
            ],
            ["field_0000", "field_0001"],
        )
        decision_page = self.client.get(decision.headers["location"])
        self.assertIn("Odoo will decide 2 required fields", decision_page.text)
        self.assertIn("Odoo will choose this value", decision_page.text)

        _checked_revision, checked_validation = (
            context.mapping_workspace.check_definition(
                workspace_id,
                datasets=working.definition.datasets,
                expected_parent_version=revision.version,
                expected_working_draft_version=working.version,
                actor=context.actor,
            )
        )
        self.assertEqual(checked_validation.status, MappingValidationStatus.VALID)
        self.assertEqual(checked_validation.issues, ())

    def test_required_managed_relationship_can_be_left_to_odoo(self) -> None:
        workspace_id, dataset, business_key = self._mapping_ready_workspace(
            scalar_field_count=0,
            relationship_field_count=1,
            relationship_field_type="one2many",
            required_relationship_indexes=(0,),
        )
        source_identity = dataset.columns[0]
        context = self.app.state.context
        _revision, validation = context.mapping_workspace.check_definition(
            workspace_id,
            datasets=(
                DatasetMapping(
                    dataset_id=dataset.dataset_id,
                    target_model="res.partner",
                    mode=MappingTargetMode.UPSERT,
                    source_identity_column_keys=(source_identity.stable_key,),
                    target_identity=(
                        IdentityComponentMapping(
                            source_column_keys=(source_identity.stable_key,),
                            target_fields=business_key.key_fields,
                        ),
                    ),
                ),
            ),
            expected_parent_version=None,
            expected_working_draft_version=None,
            actor=context.actor,
        )
        self.assertEqual(validation.status, MappingValidationStatus.INVALID)

        page = self.client.get(f"/workspaces/{workspace_id}/mapping")
        self.assertIn("Linked Field 0000 needs attention", page.text)
        self.assertIn(
            'value="set_disposition:0:relation_0000:odoo_managed"',
            page.text,
        )

        decision = self.client.post(
            f"/workspaces/{workspace_id}/mapping/save",
            data={
                "csrf_token": self.csrf,
                "action": "set_disposition:0:relation_0000:odoo_managed",
                "expected_working_draft_version": "1",
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )

        self.assertEqual(decision.status_code, 303)
        working = context.mapping_workspace.mappings.get_mapping_working_draft(
            workspace_id
        )
        self.assertEqual(
            working.definition.datasets[0]
            .target_field_dispositions[0]
            .handling,
            TargetFieldHandling.ODOO_MANAGED,
        )
        decision_page = self.client.get(decision.headers["location"])
        self.assertIn("Odoo manages this field", decision_page.text)

    def test_readonly_relationship_fields_are_hidden_without_shifting_indexes(
        self,
    ) -> None:
        workspace_id, _dataset, _business_key = self._mapping_ready_workspace(
            scalar_field_count=1,
            relationship_field_count=3,
            readonly_relationship_indexes=(1,),
        )

        page = self.client.get(f"/workspaces/{workspace_id}/mapping")

        self.assertEqual(page.status_code, 200)
        self.assertIn('data-target-field="relation_0000"', page.text)
        self.assertNotIn('data-target-field="relation_0001"', page.text)
        self.assertIn('data-target-field="relation_0002"', page.text)
        self.assertIn('name="relation_source_0_2"', page.text)

    def test_large_mapping_catalog_is_paged_and_saved_sparsely(self) -> None:
        workspace_id, dataset, business_key = self._mapping_ready_workspace(
            scalar_field_count=1500
        )
        source_identity, source_value = dataset.columns
        context = self.app.state.context
        initial = context.mapping_workspace.save_working_draft(
            workspace_id,
            datasets=(
                DatasetMapping(
                    dataset_id=dataset.dataset_id,
                    target_model="res.partner",
                    mode=MappingTargetMode.UPSERT,
                    source_identity_column_keys=(source_identity.stable_key,),
                    target_identity=(
                        IdentityComponentMapping(
                            source_column_keys=(source_identity.stable_key,),
                            target_fields=("ref",),
                        ),
                    ),
                    fields=tuple(
                        ScalarFieldMapping(
                            target_field=f"field_{index:04d}",
                            source_column_key=source_value.stable_key,
                            value_source=ScalarValueSource.SOURCE,
                        )
                        for index in range(1500)
                    ),
                ),
            ),
            expected_version=None,
            actor=context.actor,
        )
        self.assertEqual(initial.version, 1)

        page = self.client.get(f"/workspaces/{workspace_id}/mapping")
        self.assertEqual(page.status_code, 200)
        self.assertEqual(page.text.count("data-scalar-mapping-row"), 3)
        self.assertIn("Showing 3 of 1500 fields", page.text)
        self.assertIn("Page 1 of 500", page.text)
        self.assertIn("field_0000", page.text)
        self.assertNotIn("field_0003</code>", page.text)
        self.assertIn("template data-source-column-options", page.text)
        self.assertLess(page.text.count(source_value.stable_key), 100)

        last_page = self.client.get(
            f"/workspaces/{workspace_id}/mapping?scalar_page=500"
        )
        self.assertEqual(last_page.text.count("data-scalar-mapping-row"), 3)
        self.assertIn("field_1499", last_page.text)
        self.assertIn("scalar_page=500", last_page.text)

        expanded = self.client.get(
            f"/workspaces/{workspace_id}/mapping?scalar_page_size=50"
        )
        self.assertEqual(expanded.text.count("data-scalar-mapping-row"), 50)
        self.assertIn("field_0049", expanded.text)
        self.assertNotIn("field_0050</code>", expanded.text)
        self.assertIn('data-scalar-page-size="50"', expanded.text)

        rejected_size = self.client.get(
            f"/workspaces/{workspace_id}/mapping?scalar_page_size=100"
        )
        self.assertEqual(
            rejected_size.text.count("data-scalar-mapping-row"),
            3,
        )
        self.assertIn('data-scalar-page-size="3"', rejected_size.text)

        searched = self.client.get(
            f"/workspaces/{workspace_id}/mapping?field_query=field_1499"
        )
        self.assertEqual(searched.status_code, 200)
        self.assertEqual(searched.text.count("data-scalar-mapping-row"), 1)
        self.assertIn("field_1499", searched.text)
        self.assertNotIn("field_0000</code>", searched.text)

        fragment = self.client.get(
            f"/workspaces/{workspace_id}/mapping/field-catalog"
            "?field_query=field_1499"
        )
        self.assertEqual(fragment.status_code, 200)
        self.assertEqual(fragment.text.count("data-scalar-mapping-row"), 1)
        self.assertIn("field_1499", fragment.text)
        self.assertNotIn("data-relation-field-catalog", fragment.text)
        self.assertNotIn("<main", fragment.text)
        self.assertEqual(fragment.headers["cache-control"], "no-store")
        self.assertIn("projection;dur=", fragment.headers["server-timing"])

        entries = [
            ["csrf_token", self.csrf],
            ["action", "save_progress"],
            ["expected_parent_version", ""],
            ["expected_working_draft_version", "1"],
            ["editable_dataset_id", dataset.dataset_id],
            ["target_model_0", "res.partner"],
            ["mode_0", "upsert"],
            ["on_existing_0", "block"],
            ["source_identity_0", source_identity.stable_key],
            ["business_key_0", business_key.key_id],
            ["identity_source_0_0", source_identity.stable_key],
            ["visible_scalar_target_0", "field_0000"],
            ["scalar_value_source_0_1", "constant"],
            ["scalar_literal_0_1", "Updated safely"],
            ["scalar_type_0_1", "string"],
            ["scalar_case_0_1", "preserve"],
            ["scalar_compare_0_1", "1"],
            ["scalar_null_0_1", "distinct"],
        ]
        saved = self.client.post(
            f"/workspaces/{workspace_id}/mapping/save",
            json={"entries": entries},
            headers={
                **POST_HEADERS,
                "X-CSRF-Token": self.csrf,
            },
        )
        self.assertEqual(saved.status_code, 200)
        self.assertIn("redirect_url", saved.json())
        self.assertEqual(saved.json()["expected_working_draft_version"], 2)
        self.assertIn("saved_at", saved.json())
        self.assertEqual(
            saved.json()["redirect_url"],
            f"/workspaces/{workspace_id}/mapping#mapping-dataset-0",
        )
        saved_again_entries = [list(entry) for entry in entries]
        for entry in saved_again_entries:
            if entry[0] == "expected_working_draft_version":
                entry[1] = "2"
            elif entry[0] == "scalar_literal_0_1":
                entry[1] = "Updated safely again"
        saved_again = self.client.post(
            f"/workspaces/{workspace_id}/mapping/save",
            json={"entries": saved_again_entries},
            headers={
                **POST_HEADERS,
                "X-CSRF-Token": self.csrf,
            },
        )
        self.assertEqual(saved_again.status_code, 200)
        self.assertEqual(
            saved_again.json()["expected_working_draft_version"],
            3,
        )
        working = context.mapping_workspace.mappings.get_mapping_working_draft(workspace_id)
        self.assertEqual(working.version, 3)
        self.assertEqual(len(working.definition.datasets[0].fields), 1500)
        updated = {
            item.target_field: item
            for item in working.definition.datasets[0].fields
        }
        self.assertEqual(
            updated["field_0000"].literal_value,
            "Updated safely again",
        )
        self.assertEqual(
            updated["field_1499"].source_column_key,
            source_value.stable_key,
        )

        denied = self.client.post(
            f"/workspaces/{workspace_id}/mapping/save",
            json={"entries": entries},
            headers={
                **POST_HEADERS,
                "X-CSRF-Token": "incorrect",
            },
        )
        self.assertEqual(denied.status_code, 403)
        self.assertEqual(
            context.mapping_workspace.mappings.get_mapping_working_draft(workspace_id).version,
            3,
        )

        invalid_entries = [
            list(item)
            for item in entries
            if item[0] not in {"source_identity_0", "identity_source_0_0"}
        ]
        for item in invalid_entries:
            if item[0] == "action":
                item[1] = "submit"
            elif item[0] == "expected_working_draft_version":
                item[1] = "3"
        invalid = self.client.post(
            f"/workspaces/{workspace_id}/mapping/save",
            json={"entries": invalid_entries},
            headers={
                **POST_HEADERS,
                "X-CSRF-Token": self.csrf,
            },
        )
        self.assertEqual(invalid.status_code, 422)
        self.assertEqual(invalid.json()["expected_working_draft_version"], 3)
        self.assertIsNone(invalid.json()["expected_parent_version"])
        self.assertEqual(
            context.mapping_workspace.mappings.get_mapping_working_draft(
                workspace_id
            ).version,
            3,
        )
        self.assertIsNone(
            context.mapping_workspace.mappings.get_mapping_revision(workspace_id)
        )

        retry_entries = [list(item) for item in entries]
        for item in retry_entries:
            if item[0] == "expected_working_draft_version":
                item[1] = "3"
            elif item[0] == "expected_parent_version":
                item[1] = ""
        retried = self.client.post(
            f"/workspaces/{workspace_id}/mapping/save",
            json={"entries": retry_entries},
            headers={
                **POST_HEADERS,
                "X-CSRF-Token": self.csrf,
            },
        )
        self.assertEqual(retried.status_code, 200)
        self.assertEqual(
            context.mapping_workspace.mappings.get_mapping_working_draft(workspace_id).version,
            4,
        )

        oversized = b'{"entries":[["csrf_token","' + (
            b"x" * (5 * 1024 * 1024)
        ) + b'"]]}'
        rejected = self.client.post(
            f"/workspaces/{workspace_id}/mapping/save",
            content=oversized,
            headers={
                **POST_HEADERS,
                "Content-Type": "application/json",
                "X-CSRF-Token": self.csrf,
            },
        )
        self.assertEqual(rejected.status_code, 413)
        self.assertEqual(
            context.mapping_workspace.mappings.get_mapping_working_draft(workspace_id).version,
            4,
        )

        excessive_form = "&".join(
            [f"field_{index}=x" for index in range(25_001)]
        )
        recovered = self.client.post(
            f"/workspaces/{workspace_id}/mapping/save",
            content=excessive_form,
            headers={
                **POST_HEADERS,
                "Content-Type": "application/x-www-form-urlencoded",
            },
            follow_redirects=False,
        )
        self.assertEqual(recovered.status_code, 303)
        self.assertIn("save_error=request_rejected", recovered.headers["location"])
        recovery_page = self.client.get(recovered.headers["location"])
        self.assertIn("No mapping change was saved", recovery_page.text)

    def test_saved_prepared_data_has_plain_recovery_ui(self) -> None:
        workspace_id, _dataset, _business_key = self._mapping_ready_workspace(
            scalar_field_count=0,
        )
        context = self.app.state.context
        staging = MagicMock(
            total_rows=12,
            mapping_version=3,
            run_id="f0cd6d32-80d9-4e31-9bcb-d316d83cf0b8",
            content_hash="sha256:" + "7" * 64,
            datasets=(),
        )

        with (
            patch.object(
                context.preflight,
                "current_staging",
                return_value=staging,
            ),
            patch.object(
                context.preflight,
                "current_report",
                return_value=None,
            ),
        ):
            page = self.client.get(f"/workspaces/{workspace_id}/summary")

        self.assertEqual(page.status_code, 200)
        self.assertIn("Your prepared data is safe", page.text)
        self.assertIn("12 prepared rows", page.text)
        self.assertIn("Prepare data for review", page.text)
        self.assertIn("Prepared data is stored locally", page.text)
        self.assertIn("<details", page.text)
        self.assertIn("<summary>Support details</summary>", page.text)
        self.assertNotIn("<details open", page.text)
        self.assertNotIn("canonical_staging", page.text)

    def test_prepare_rejects_bad_source_hash_before_publication_and_redirects(
        self,
    ) -> None:
        workspace_id, dataset, business_key = self._mapping_ready_workspace(
            scalar_field_count=0,
        )
        context = self.app.state.context
        source_identity = dataset.columns[0]
        mapping = DatasetMapping(
            dataset_id=dataset.dataset_id,
            target_model="res.partner",
            mode=MappingTargetMode.UPSERT,
            source_identity_column_keys=(source_identity.stable_key,),
            target_identity=(
                IdentityComponentMapping(
                    source_column_keys=(source_identity.stable_key,),
                    target_fields=business_key.key_fields,
                ),
            ),
        )
        revision, validation = context.mapping_workspace.check_definition(
            workspace_id,
            datasets=(mapping,),
            expected_parent_version=None,
            expected_working_draft_version=None,
            actor=context.actor,
        )
        submission = context.mapping_workspace.submit_current(
            workspace_id,
            datasets=(mapping,),
            expected_version=revision.version,
            expected_working_draft_version=1,
            actor=context.actor,
        )
        self.assertNotEqual(validation.status, MappingValidationStatus.INVALID)
        self.assertIsNotNone(submission)
        selection = context.sources.sources.get_source_selection(workspace_id)
        assert selection is not None
        corrupted = json.loads(selection.to_json())
        corrupted["datasets"][0]["source"]["source_sha256"] = (
            "sha256:not-a-digest"
        )
        database_path = (
            context.workspace_states.repository.workspace_directory(workspace_id) / "workspace-engine.duckdb"
        )
        with context.workspace_states.repository._connect(database_path) as connection:
            connection.execute(
                "UPDATE source_selection SET selection_json = ? "
                "WHERE singleton_id = 1",
                [json.dumps(corrupted)],
            )

        failed = self.client.post(
            f"/workspaces/{workspace_id}/summary/check",
            data={"csrf_token": self.csrf},
            headers=POST_HEADERS,
            follow_redirects=False,
        )

        self.assertEqual(failed.status_code, 303)
        self.assertIn(
            f"/workspaces/{workspace_id}/preparation/",
            failed.headers["location"],
        )
        completed_job = _wait_for_preparation(
            self.client,
            failed.headers["location"],
        )
        self.assertEqual(completed_job["status"], "FAILED")
        self.assertTrue(completed_job["failure_code"])
        self.assertIsNone(context.preflight.current_staging(workspace_id))
        self.assertIsNone(context.quality.current_summary(workspace_id))
        self.assertIsNone(
            context.normalization.current_summary(workspace_id)
        )
        self.assertEqual(self.readiness_calls, [])

        recovery = self.client.get(failed.headers["location"])
        self.assertEqual(recovery.status_code, 200)
        self.assertIn("Stored source selection is invalid", recovery.text)
        self.assertIn("data-preparation-failure-code", recovery.text)
        self.assertIn(str(completed_job["failure_code"]), recovery.text)
        retried = self.client.post(
            f"{failed.headers['location']}/retry",
            data={"csrf_token": self.csrf},
            headers=POST_HEADERS,
            follow_redirects=False,
        )
        self.assertEqual(retried.status_code, 303)
        self.assertNotEqual(retried.headers["location"], failed.headers["location"])
        retried_job = _wait_for_preparation(
            self.client,
            retried.headers["location"],
        )
        self.assertEqual(retried_job["status"], "FAILED")

    def test_data_manager_can_save_an_optional_named_total(self) -> None:
        workspace_id, dataset, business_key = self._mapping_ready_workspace(
            scalar_field_count=1,
            numeric_field=True,
        )
        source_identity, source_value = dataset.columns

        page = self.client.get(f"/workspaces/{workspace_id}/mapping")
        self.assertEqual(page.status_code, 200)
        self.assertIn("Check a known total (optional)", page.text)
        self.assertIn("Allowed difference", page.text)
        self.assertNotIn("control_totals_json", page.text)

        saved = self.client.post(
            f"/workspaces/{workspace_id}/mapping/save",
            data={
                "csrf_token": self.csrf,
                "action": "save_progress",
                "expected_parent_version": "",
                "expected_working_draft_version": "",
                "target_model_0": "res.partner",
                "mode_0": "upsert",
                "source_identity_0": source_identity.stable_key,
                "business_key_0": business_key.key_id,
                "identity_source_0_0": source_identity.stable_key,
                "scalar_value_source_0_1": "source",
                "scalar_source_0_1": source_value.stable_key,
                "scalar_type_0_1": "decimal",
                "scalar_compare_0_1": "1",
                "scalar_null_0_1": "distinct",
                "control_name_0_0": "Opening balance",
                "control_target_0_0": "field_0000",
                "control_expected_0_0": "1234.50",
                "control_unit_0_0": "EUR",
                "control_tolerance_0_0": "0.01",
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )

        self.assertEqual(saved.status_code, 303)
        working = (
            self.app.state.context.mapping_workspace.mappings.get_mapping_working_draft(
                workspace_id
            )
        )
        control = working.definition.datasets[0].effective_control_totals[0]
        self.assertEqual(control.name, "Opening balance")
        self.assertEqual(control.target_field, "field_0000")
        self.assertEqual(control.expected_total, "1234.50")
        self.assertEqual(control.unit, "EUR")
        self.assertEqual(control.tolerance, "0.01")

    def test_active_preparation_blocks_mapping_autosave_before_database_read(
        self,
    ) -> None:
        workspace_id, _dataset, _business_key = self._mapping_ready_workspace(
            scalar_field_count=1,
        )
        context = self.app.state.context
        manager = context.preparation_jobs
        assert manager is not None
        active = SimpleNamespace(job_id=str(uuid4()))

        with patch.object(manager, "active", return_value=active):
            mapping_page = self.client.get(
                f"/workspaces/{workspace_id}/mapping",
                follow_redirects=False,
            )
            autosave = self.client.post(
                f"/workspaces/{workspace_id}/mapping/save",
                json={"entries": []},
                headers={**POST_HEADERS, "X-CSRF-Token": self.csrf},
            )

        progress_url = f"/workspaces/{workspace_id}/preparation/{active.job_id}"
        self.assertEqual(mapping_page.status_code, 303)
        self.assertEqual(mapping_page.headers["location"], progress_url)
        self.assertEqual(autosave.status_code, 409)
        self.assertEqual(autosave.json()["redirect_url"], progress_url)
        self.assertIn("Wait for it to finish", autosave.json()["detail"])

    def test_data_manager_can_save_a_guided_business_data_check(self) -> None:
        workspace_id, dataset, business_key = self._mapping_ready_workspace(
            scalar_field_count=2,
        )
        source_identity, source_value = dataset.columns
        context = self.app.state.context
        _revision, validation = (
            context.mapping_workspace.check_definition(
                workspace_id,
                datasets=(
                    DatasetMapping(
                        dataset_id=dataset.dataset_id,
                        target_model="res.partner",
                        mode=MappingTargetMode.UPSERT,
                        source_identity_column_keys=(
                            source_identity.stable_key,
                        ),
                        target_identity=(
                            IdentityComponentMapping(
                                source_column_keys=(
                                    source_identity.stable_key,
                                ),
                                target_fields=business_key.key_fields,
                            ),
                        ),
                        fields=(
                            ScalarFieldMapping(
                                target_field="field_0000",
                                source_column_key=source_value.stable_key,
                                value_source=ScalarValueSource.SOURCE,
                            ),
                            ScalarFieldMapping(
                                target_field="field_0001",
                                source_column_key=source_value.stable_key,
                                value_source=ScalarValueSource.SOURCE,
                            ),
                        ),
                    ),
                ),
                expected_parent_version=None,
                expected_working_draft_version=None,
                actor=context.actor,
            )
        )
        self.assertNotEqual(
            validation.status,
            MappingValidationStatus.INVALID,
        )

        page = self.client.get(f"/workspaces/{workspace_id}/mapping")
        self.assertEqual(page.status_code, 200)
        self.assertIn("Data checks", page.text)
        self.assertIn("Recommended checks are already on", page.text)
        self.assertIn("Add business check 1", page.text)
        self.assertNotIn("ruleset_json", page.text)

        saved = self.client.post(
            f"/workspaces/{workspace_id}/mapping/quality",
            data={
                "csrf_token": self.csrf,
                "quality_dataset_id": dataset.dataset_id,
                "quality_name_0": "Opening before closing",
                "quality_family_0": "ORDERED_COMPARISON",
                "quality_field_a_0": "field_0000",
                "quality_field_b_0": "field_0001",
                "quality_equals_0": "",
                "quality_outcome_0": "QUARANTINE",
                "quality_owner_0": "FUNCTIONAL_OWNER",
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )

        self.assertEqual(saved.status_code, 303)
        ruleset = context.quality.quality.get_current_quality_ruleset(workspace_id)
        self.assertIsNotNone(ruleset)
        self.assertEqual(len(ruleset.manager_rules), 1)
        rule = ruleset.manager_rules[0]
        self.assertEqual(rule.name, "Opening before closing")
        self.assertEqual(rule.family, QualityRuleFamily.ORDERED_COMPARISON)
        self.assertEqual(rule.outcome, QualityOutcomePolicy.QUARANTINE)
        self.assertEqual(rule.owner_role, QualityOwnerRole.FUNCTIONAL_OWNER)
        restored = self.client.get(saved.headers["location"])
        self.assertIn("Opening before closing", restored.text)
        self.assertIn("Functional owner", restored.text)

    def test_failed_named_total_has_plain_review_ui_and_blocks_package(self) -> None:
        workspace_id, _dataset, _business_key = self._mapping_ready_workspace(
            scalar_field_count=0,
        )
        context = self.app.state.context
        total = CanonicalControlTotal(
            control_id="sha256:" + "d" * 64,
            name="Opening balance",
            dataset="contacts",
            target_field="credit_limit",
            expected_total="1000",
            actual_total="900",
            tolerance="0",
            unit="EUR",
            included_rows=12,
            empty_rows=0,
        )
        staging = MagicMock(
            total_rows=12,
            mapping_version=3,
            run_id="f0cd6d32-80d9-4e31-9bcb-d316d83cf0b8",
            content_hash="sha256:" + "7" * 64,
            datasets=(),
            control_totals=(total,),
            failed_control_total_count=1,
            control_totals_passed=False,
        )
        report = MagicMock(
            status="READY",
            blocked_count=0,
            needs_review_count=0,
            ready_count=0,
            total_count=0,
            datasets=(),
            rows=(),
            checked_at=datetime.now(timezone.utc),
            run_id=str(uuid4()),
        )

        with (
            patch.object(
                context.preflight,
                "current_staging",
                return_value=staging,
            ),
            patch.object(
                context.preflight,
                "current_report",
                return_value=report,
            ),
        ):
            page = self.client.get(f"/workspaces/{workspace_id}/summary")
            package = self.client.post(
                f"/workspaces/{workspace_id}/summary/package",
                data={"csrf_token": self.csrf},
                headers=POST_HEADERS,
            )

        self.assertEqual(page.status_code, 200)
        self.assertIn("Some totals need attention", page.text)
        self.assertIn("Opening balance", page.text)
        self.assertIn("Difference -100 EUR", page.text)
        self.assertIn("<summary>Support details</summary>", page.text)
        self.assertNotIn("Save these rules as a Recipe", page.text)
        self.assertEqual(package.status_code, 422)
        self.assertIn("Resolve the named totals", package.text)

    def test_transformation_impact_uses_server_filters_and_100_row_pages(
        self,
    ) -> None:
        workspace_id, dataset, business_key = self._mapping_ready_workspace(
            scalar_field_count=1,
        )
        source_identity, source_value = dataset.columns
        context = self.app.state.context
        revision, validation = (
            context.mapping_workspace.check_definition(
                workspace_id,
                datasets=(
                    DatasetMapping(
                        dataset_id=dataset.dataset_id,
                        target_model="res.partner",
                        mode=MappingTargetMode.UPSERT,
                        source_identity_column_keys=(
                            source_identity.stable_key,
                        ),
                        target_identity=(
                            IdentityComponentMapping(
                                source_column_keys=(
                                    source_identity.stable_key,
                                ),
                                target_fields=business_key.key_fields,
                            ),
                        ),
                        fields=(
                            ScalarFieldMapping(
                                target_field="field_0000",
                                source_column_key=source_value.stable_key,
                                value_source=ScalarValueSource.SOURCE,
                            ),
                        ),
                    ),
                ),
                expected_parent_version=None,
                expected_working_draft_version=None,
                actor=context.actor,
            )
        )
        self.assertNotEqual(validation.status, MappingValidationStatus.INVALID)
        impact_rows = tuple(
            TransformationImpactRow(
                dataset=dataset.name,
                source_row=index + 2,
                source_column=source_value.source_name,
                target_field="field_0000",
                raw_value=f" raw {index} ",
                proposed_value=f"Raw {index}",
                rules="Trim",
                outcome="invalid" if index % 2 else "changed",
                message="Needs review" if index % 2 else "",
            )
            for index in range(205)
        )

        def fake_stage(*_args, **kwargs):
            sink = kwargs["transformation_impact_sink"]
            for row in impact_rows:
                sink(row)
            return MagicMock(
                transformation_impact=TransformationImpactReport(
                    mapping_content_hash=revision.definition.content_hash,
                    evaluated_count=205,
                    changed_count=103,
                    fallback_count=0,
                    null_count=0,
                    invalid_count=102,
                    provided_count=0,
                    unchanged_count=0,
                    rows=(),
                    detail_limit=0,
                )
            )

        impact_url = f"/workspaces/{workspace_id}/mapping/transformation-impact"
        first_visit = self.client.get(impact_url)
        self.assertIn("Prepare the comparison", first_visit.text)
        self.assertIn("data-transformation-impact-prepare", first_visit.text)
        self.assertIn("data-transformation-impact-status", first_visit.text)
        self.assertIn('aria-live="polite"', first_visit.text)
        impact_script = self.client.get("/static/app.js")
        self.assertIn("[data-transformation-impact-prepare]", impact_script.text)
        self.assertIn("Preparing the comparison…", impact_script.text)
        with patch(
            "impodo.application.transformation_impact_service.stage_browser_mapping",
            side_effect=fake_stage,
        ) as staged:
            prepared = self.client.post(
                f"{impact_url}/prepare",
                data={"csrf_token": self.csrf},
                headers=POST_HEADERS,
                follow_redirects=False,
            )
        self.assertEqual(prepared.status_code, 303)
        staged.assert_called_once()

        first_page = self.client.get(impact_url)
        self.assertEqual(first_page.text.count('class="impact-row'), 100)
        self.assertIn(
            "Contains 1 space before the value and 1 space after the value.",
            first_page.text,
        )
        self.assertIn(
            "Removed 1 space before the value and 1 space after the value.",
            first_page.text,
        )
        self.assertIn("Showing 1–100 of 205", first_page.text)
        self.assertIn("Next 100", first_page.text)
        next_match = re.search(
            r'href="([^"]+after=[^"]+)"[^>]*>Next 100</a>',
            first_page.text,
        )
        self.assertIsNotNone(next_match)
        second_page = self.client.get(unescape(next_match.group(1)))
        self.assertEqual(second_page.text.count('class="impact-row'), 100)
        self.assertIn("Showing 101–200 of 205", second_page.text)
        self.assertIn("Previous 100", second_page.text)

        invalid_page = self.client.get(f"{impact_url}?outcome=invalid")
        self.assertEqual(invalid_page.text.count('class="impact-row'), 100)
        self.assertIn("Showing 1–100 of 102", invalid_page.text)
        invalid_csv = self.client.post(
            f"{impact_url}.csv",
            data={"csrf_token": self.csrf, "outcome": "invalid"},
            headers=POST_HEADERS,
        )
        self.assertEqual(invalid_csv.status_code, 200)
        self.assertEqual(len(invalid_csv.text.splitlines()), 103)

    def test_schema_governance_keeps_duplicate_fields_out_of_one_rule(
        self,
    ) -> None:
        workspace_id, _dataset, original_key = self._mapping_ready_workspace(
            scalar_field_count=1,
        )
        context = self.app.state.context
        original_governance = (
            context.schema_workspace.schemas.get_schema_governance(workspace_id)
        )
        self.assertIsNotNone(original_governance)

        duplicate_simple = self.client.post(
            f"/workspaces/{workspace_id}/schema/govern",
            data={
                "csrf_token": self.csrf,
                "primary_key_field_0": "ref",
                "primary_scope_field_0": "ref",
                "key_fields_0": "ref",
                "scope_fields_0": "ref",
                "key_description_0": "Reference within reference",
            },
            headers=POST_HEADERS,
        )

        self.assertEqual(duplicate_simple.status_code, 422)
        self.assertIn(
            "Review the highlighted matching rule, then confirm it again.",
            duplicate_simple.text,
        )
        self.assertIn(
            (
                "For res.partner, choose each field only once. The matching fields "
                "and Within fields must be different."
            ),
            duplicate_simple.text,
        )
        self.assertRegex(
            duplicate_simple.text,
            (
                r'(?s)name="primary_key_field_0".*?'
                r'<option\s+value="ref"\s+selected'
            ),
        )
        self.assertRegex(
            duplicate_simple.text,
            (
                r'(?s)name="primary_scope_field_0".*?'
                r'<option\s+value="ref"\s+selected'
            ),
        )
        unchanged_governance = (
            context.schema_workspace.schemas.get_schema_governance(workspace_id)
        )
        self.assertIsNotNone(unchanged_governance)
        self.assertEqual(
            unchanged_governance.content_hash,
            original_governance.content_hash,
        )
        self.assertEqual(unchanged_governance.business_keys, (original_key,))

        duplicate_combined = self.client.post(
            f"/workspaces/{workspace_id}/schema/govern",
            data={
                "csrf_token": self.csrf,
                "primary_key_field_0": "",
                "primary_scope_field_0": "",
                "key_fields_0": "ref, ref",
                "scope_fields_0": "",
                "key_description_0": "Repeated combined reference",
            },
            headers=POST_HEADERS,
        )

        self.assertEqual(duplicate_combined.status_code, 422)
        self.assertIn('value="ref, ref"', duplicate_combined.text)
        self.assertIn("Repeated combined reference", duplicate_combined.text)

        valid = self.client.post(
            f"/workspaces/{workspace_id}/schema/govern",
            data={
                "csrf_token": self.csrf,
                "primary_key_field_0": "field_0000",
                "primary_scope_field_0": "ref",
                "key_fields_0": "field_0000",
                "scope_fields_0": "ref",
                "key_description_0": "Field within reference",
            },
            headers=POST_HEADERS,
            follow_redirects=False,
        )

        self.assertEqual(valid.status_code, 303)
        saved_governance = (
            context.schema_workspace.schemas.get_schema_governance(workspace_id)
        )
        self.assertIsNotNone(saved_governance)
        self.assertEqual(
            saved_governance.business_keys[0].key_fields,
            ("field_0000",),
        )
        self.assertEqual(
            saved_governance.business_keys[0].scope_fields,
            ("ref",),
        )

        schema_script = self.client.get("/static/app.js")
        self.assertIn("updateKeyFieldConflicts", schema_script.text)
        self.assertIn(
            "Matching fields and Within fields must be different.",
            schema_script.text,
        )

    def test_preparation_worker_and_progress_page_do_not_open_locked_databases(
        self,
    ) -> None:
        workspace_id, _dataset, _business_key = self._mapping_ready_workspace(
            scalar_field_count=1,
        )
        context = self.app.state.context
        workspace_state = context.queries.get(workspace_id)
        selection = context.queries.get_source_selection(workspace_id)
        assert selection is not None
        migration_workspace = context.migration_workspaces.get(
            workspace_id,
            actor=context.actor,
        )
        workspace = PreparationWorkspace.from_context(
            migration_workspace,
            context.data_versions.get(
                migration_workspace.data_version_id,
                actor=context.actor,
            ),
            context.migration_runs.get(
                migration_workspace.migration_run_id,
                actor=context.actor,
            ),
        )
        self.assertEqual(workspace.migration_run_purpose.value, "AUTHORING")
        manager = context.preparation_jobs
        assert manager is not None
        root = Path(self.temporary.name)
        registry_path = root / "registry.duckdb"
        project_path = (
            context.workspace_states.repository.workspace_directory(workspace_id)
            / "workspace-engine.duckdb"
        )

        with _spawned_duckdb_locks(registry_path):
            job = manager.enqueue(
                workspace_id,
                workspace_state.name,
                sum(item.row_count for item in selection.datasets),
                actor=context.actor,
                workspace=workspace,
            )
            progress_url = (
                f"/workspaces/{workspace_id}/preparation/{job.job_id}"
            )
            completed = _wait_for_preparation(
                self.client,
                progress_url,
                timeout=30,
            )

        self.assertEqual(completed["status"], PreparationJobStatus.FAILED.value)
        self.assertEqual(completed["failure_code"], "ReadinessError", completed)
        self.assertIn("Submit the mapping", str(completed["failure_message"]))
        worker_deadline = time.monotonic() + 2.0
        while manager.worker_alive(job.job_id) and time.monotonic() < worker_deadline:
            time.sleep(0.01)
        self.assertFalse(manager.worker_alive(job.job_id))

        with _spawned_duckdb_locks(registry_path, project_path):
            progress_page = self.client.get(progress_url)

        self.assertEqual(progress_page.status_code, 200, progress_page.text)
        self.assertIn("Stage 4 of 6", progress_page.text)
        self.assertIn("Preparation progress", progress_page.text)
        self.assertIn(
            f'href="/projects/{workspace.project_id}"',
            progress_page.text,
        )
        self.assertIn('aria-current="page"', progress_page.text)

    def _mapping_ready_workspace(
        self,
        *,
        scalar_field_count: int,
        relationship_field_count: int = 0,
        relationship_model: str = "res.partner",
        selection_field: bool = False,
        numeric_field: bool = False,
        readonly_scalar_indexes: tuple[int, ...] = (),
        readonly_relationship_indexes: tuple[int, ...] = (),
        required_scalar_indexes: tuple[int, ...] = (),
        verified_default_scalar_indexes: tuple[int, ...] = (),
        required_relationship_indexes: tuple[int, ...] = (),
        relationship_field_type: str = "many2one",
        business_key_description: str = "Unique reference",
        target_model: str = "res.partner",
        relationship_field_names: tuple[str, ...] | None = None,
    ):
        context = self.app.state.context
        created = _create_workspace(
            context,
            name="Large mapping",
            source_system="CSV",
        )
        now = datetime.now(timezone.utc)
        registered = replace(
            created,
            odoo_connection_mode=OdooConnectionMode.LOCAL,
            odoo_base_url="http://127.0.0.1:8069",
            odoo_database="odoo19_local",
            intended_models=(target_model,),
            status=WorkspaceStatus.REGISTERED,
            revision=2,
            updated_at=now,
            registered_at=now,
        )
        context.workspace_states.repository.save(
            registered,
            expected_revision=created.revision,
            event_type="WORKSPACE_REGISTERED",
            event_detail="",
            actor=context.actor,
        )
        _replace_run_target_setup(
            context,
            registered.workspace_id,
            connection_mode=OdooConnectionMode.LOCAL,
            base_url="http://127.0.0.1:8069",
            database="odoo19_local",
        )
        dataset = SourceDataset(
            dataset_id="dataset:large",
            name="large_contacts",
            source=FileSourceBinding(
                file_id="source:large",
                table_key="contacts",
                source_sha256="sha256:" + "1" * 64,
                catalog_hash="sha256:" + "2" * 64,
                encoding="utf-8",
                delimiter=",",
                header_row=1,
            ),
            row_count=1,
            columns=(
                SourceDatasetColumn(1, "Reference", "column:ref", "string"),
                SourceDatasetColumn(2, "Value", "column:value", "string"),
            ),
        )
        migration_workspace = context.migration_workspaces.get(
            registered.workspace_id,
            actor=context.actor,
        )
        data_version = context.data_versions.get(
            migration_workspace.data_version_id,
            actor=context.actor,
        )
        package_service = context.data_version_source_projection.packages
        current_package = package_service.repository.get_source_package(
            data_version.data_version_id
        )
        assert current_package is not None
        source_file_id = str(uuid4())
        source_hash = "sha256:" + "1" * 64
        source_catalog = SourceFileCatalog(
            contract_version=CATALOG_CONTRACT_VERSION,
            file_id=source_file_id,
            display_name="large-contacts.csv",
            source_sha256=source_hash,
            source_size_bytes=128,
            format="CSV",
            inspected_at=now,
            encoding="utf-8",
            delimiter=",",
            tables=(
                SourceTableCatalog(
                    table_key="contacts",
                    name="large_contacts",
                    kind="CSV",
                    hidden=False,
                    header_row=1,
                    row_count=1,
                    column_count=2,
                    columns=(
                        _source_column_profile(1, "Reference", "P001"),
                        _source_column_profile(2, "Value", "Example"),
                    ),
                    preview_rows=(("P001", "Example"),),
                ),
            ),
        )
        package_catalog = SourcePackageCatalog(
            file_id=source_file_id,
            source_sha256=source_hash,
            payload=json.loads(source_catalog.to_json()),
        )
        package_dataset = SourcePackageDataset(
            dataset_id=dataset.dataset_id,
            display_name=dataset.name,
            source_file_ids=(source_file_id,),
            source=FileSourceBinding(
                file_id=source_file_id,
                table_key="contacts",
                source_sha256=source_hash,
                catalog_hash=package_catalog.content_hash,
                encoding="utf-8",
                delimiter=",",
                header_row=1,
            ),
            row_count=dataset.row_count,
            columns=dataset.columns,
            schema_hash=source_column_contract_hash(dataset.columns),
            snapshot_hash="sha256:" + "4" * 64,
            snapshot_storage_key="snapshots/large_contacts.parquet",
            manifest={"logical_name": "large_contacts"},
        )
        dataset = package_dataset.to_mapping_dataset()
        draft_package = DataVersionSourcePackage(
            data_version_id=data_version.data_version_id,
            project_id=migration_workspace.project_id,
            revision=current_package.revision + 1,
            origin=SourcePackageOrigin.FILE,
            state=SourcePackageState.DRAFT,
            files=(
                SourcePackageFile(
                    file_id=source_file_id,
                    display_name="large-contacts.csv",
                    storage_key="source/large-contacts.csv",
                    size_bytes=128,
                    sha256=source_hash,
                    received_at=now,
                ),
            ),
            catalogs=(package_catalog,),
            configurations=(
                SourcePackageConfiguration(
                    file_id=source_file_id,
                    catalog_hash=package_catalog.content_hash,
                    payload={
                        "file_id": source_file_id,
                        "source_sha256": source_hash,
                        "catalog_hash": package_catalog.content_hash,
                        "encoding": "utf-8",
                        "delimiter": ",",
                        "selected_table_keys": ["contacts"],
                        "warnings_acknowledged": False,
                        "confirmed_at": now.isoformat(),
                        "confirmed_by": context.actor.identity.display_name,
                    },
                ),
            ),
            datasets=(package_dataset,),
            updated_at=now,
        )
        saved_package = package_service.replace_draft(
            draft_package,
            actor=context.actor,
            expected_package_revision=current_package.revision,
        )
        package_service.freeze(
            data_version.data_version_id,
            actor=context.actor,
            expected_data_version_revision=data_version.optimistic_revision,
            expected_package_revision=saved_package.revision,
            operation_id=str(uuid4()),
        )
        context.data_version_source_projection.projections.materialize(
            registered.workspace_id,
            actor=context.actor,
            dataset_ids=(dataset.dataset_id,),
            expected_workspace_revision=(
                migration_workspace.optimistic_revision
            ),
            operation_id=str(uuid4()),
        )
        selection = SourceSelection(
            selection_id=str(uuid4()),
            version=1,
            data_version_id=data_version.data_version_id,
            created_at=now,
            created_by=context.actor.identity.display_name,
            datasets=(dataset,),
            content_hash="sha256:" + "3" * 64,
        )
        context.sources.sources.save_source_selection(
            registered.workspace_id,
            selection,
            actor=context.actor,
        )
        fields = (
            SchemaField(
                name="ref",
                label="Reference",
                type="char",
                required=True,
                readonly=False,
                relation=None,
                relation_field=None,
                selection=(),
            ),
            *(
                SchemaField(
                    name=f"field_{index:04d}",
                    label=f"Field {index:04d}",
                    type=(
                        "selection"
                        if selection_field and index == 0
                        else (
                            "monetary"
                            if numeric_field and index == 0
                            else "char"
                        )
                    ),
                    required=index in required_scalar_indexes,
                    readonly=index in readonly_scalar_indexes,
                    relation=None,
                    relation_field=None,
                    selection=(
                        (
                            ("fr_FR", "French (France)"),
                            ("de_DE", "German"),
                        )
                        if selection_field and index == 0
                        else ()
                    ),
                    create_default_present=(
                        index in verified_default_scalar_indexes
                    ),
                    create_default_value=(
                        "fr_FR"
                        if selection_field and index == 0
                        else f"Odoo default {index:04d}"
                    )
                    if index in verified_default_scalar_indexes
                    else None,
                )
                for index in range(scalar_field_count)
            ),
            *(
                SchemaField(
                    name=(
                        relationship_field_names[index]
                        if relationship_field_names is not None
                        else f"relation_{index:04d}"
                    ),
                    label=f"Linked Field {index:04d}",
                    type=relationship_field_type,
                    required=index in required_relationship_indexes,
                    readonly=index in readonly_relationship_indexes,
                    relation=relationship_model,
                    relation_field=None,
                    selection=(),
                )
                for index in range(relationship_field_count)
            ),
        )
        schema = OdooSchemaCatalog(
            workspace_id=registered.workspace_id,
            policy_hash=ODOO_SOURCE_POLICY_HASH,
            captured_at=now,
            captured_by=context.actor.identity.display_name,
            connection_mode=registered.odoo_connection_mode.value,
            database=registered.odoo_database,
            odoo_version="19.0",
            models=(SchemaModel(target_model, target_model, fields),),
            content_hash="sha256:" + "4" * 64,
            origin=SchemaOrigin.LIVE_API,
            read_credential_binding_hash="sha256:" + "1" * 64,
            read_principal_hash="sha256:" + "1" * 64,
            read_permission_hash="sha256:" + "2" * 64,
            read_context_hash="sha256:" + "3" * 64,
            connection_target_hash=_browser_schema(registered).fingerprint.target_hash,
        )
        context.schema_workspace.schemas.save_odoo_schema_catalog(
            registered.workspace_id,
            schema,
            actor=context.actor,
        )
        business_key = BusinessKeyDefinition(
            key_id=f"{target_model}:ref",
            model=target_model,
            key_fields=("ref",),
            description=business_key_description,
            status=BusinessKeyStatus.CONFIRMED,
        )
        governance = SchemaGovernance(
            governance_id=str(uuid4()),
            version=1,
            workspace_id=registered.workspace_id,
            catalog_hash=schema.content_hash,
            permitted_models=(target_model,),
            business_keys=(business_key,),
            recorded_at=now,
            recorded_by=context.actor.identity.display_name,
        )
        context.schema_workspace.schemas.save_schema_governance(
            registered.workspace_id,
            governance,
            actor=context.actor,
        )
        return registered.workspace_id, dataset, business_key

    def _post(self, path: str, data: dict[str, str]):
        submitted = dict(data)
        if path == "/projects/new":
            submitted.setdefault("creation_request_id", str(uuid4()))
        return self.client.post(
            path,
            data=submitted,
            headers=POST_HEADERS,
            follow_redirects=False,
        )


def _csrf(html: str) -> str:
    matched = re.search(r'name="csrf_token" value="([^"]+)"', html)
    if matched is None:
        raise AssertionError("CSRF token not found")
    return matched.group(1)


def _browser_schema(workspace_state) -> MetadataSnapshot:
    return MetadataSnapshot(
        fingerprint=TargetFingerprint(
            target_hash=target_identity_hash(
                connection_mode=workspace_state.odoo_connection_mode.value,
                base_url=workspace_state.odoo_base_url,
                database=workspace_state.odoo_database,
            ),
            connection_mode=workspace_state.odoo_connection_mode.value,
            database=workspace_state.odoo_database,
            odoo_version="19.0",
            snapshot_timestamp="2026-07-29T12:00:00Z",
            module_versions={"base": "19.0.1.0"},
        ),
        models={
            "res.partner": ModelMetadata(
                model="res.partner",
                description=None,
                fields={
                    "name": FieldMetadata(
                        name="name",
                        type="char",
                        label="Name",
                        required=True,
                        stored=True,
                        computed=False,
                        has_inverse=False,
                        related=False,
                        translated=False,
                        company_dependent=False,
                        searchable=True,
                        sortable=True,
                        exportable=True,
                    ),
                    "write_date": FieldMetadata(
                        name="write_date",
                        type="datetime",
                        label="Last Updated on",
                        readonly=True,
                        stored=True,
                        computed=False,
                        has_inverse=False,
                        related=False,
                        translated=False,
                        company_dependent=False,
                        searchable=True,
                        sortable=True,
                        exportable=True,
                    ),
                    "ref": FieldMetadata(
                        name="ref",
                        type="char",
                        label="Reference",
                    ),
                    "display_name": FieldMetadata(
                        name="display_name",
                        type="char",
                        label="Display Name",
                        readonly=True,
                    ),
                    "company_id": FieldMetadata(
                        name="company_id",
                        type="many2one",
                        label="Company",
                        relation="res.company",
                    ),
                    "category_ids": FieldMetadata(
                        name="category_ids",
                        type="many2many",
                        label="Tags",
                        relation="res.partner.category",
                    ),
                    "child_ids": FieldMetadata(
                        name="child_ids",
                        type="one2many",
                        label="Contacts",
                        relation="res.partner",
                        relation_field="parent_id",
                    ),
                },
            )
        },
    )


class _BrowserOdooCaptureGateway:
    def __init__(self, workspace_state, schema: OdooSchemaCatalog) -> None:
        self.workspace_state = workspace_state
        self.schema = schema
        self.now = datetime.now(timezone.utc)
        self.calls: list[str] = []
        self.context = ProtectedOdooReadContext(
            language="en_US",
            timezone="UTC",
            primary_company_id=1,
            allowed_company_ids=(1,),
        )

    def probe_identity(self, request, *, cancellation=None):
        self.calls.append("identity")
        return (
            OdooReadIdentity(
                target_hash=self.schema.connection_target_hash,
                principal_hash=self.schema.read_principal_hash,
                permission_hash=self.schema.read_permission_hash,
                context_hash=self.schema.read_context_hash,
                readable_models=request.schema_model_names,
                observed_at=self.now.isoformat(),
            ),
            self.context,
        )

    def probe_schema(self, request, context, *, cancellation=None):
        self.calls.append("schema")
        return _browser_schema(self.workspace_state)

    def open_capture(self, request, context, *, cancellation=None):
        self.calls.append("open")
        return _BrowserOdooCaptureSession(request, self.now)

    def sample(self, request, context, *, limit, cancellation=None):
        raise AssertionError("Freeze action does not run a sample")


class _BrowserOdooCaptureSession:
    def __init__(self, request, now: datetime) -> None:
        self._page = OdooCapturePage(
            first_row_ordinal=1,
            odoo_ids=(11, 12),
            write_dates=(now, now + timedelta(seconds=1)),
            columns=(
                OdooCaptureValueColumn(
                    field_name="name",
                    field_type="char",
                    values=("Alice", "Bob"),
                ),
            ),
            response_bytes=100,
            normalized_bytes=20,
        )
        self._accounting = OdooCaptureAccounting(
            high_water_id=12,
            row_count=2,
            page_count=1,
            record_request_count=2,
            response_bytes=102,
            normalized_bytes=20,
            capture_started_at=now,
            capture_finished_at=now + timedelta(seconds=2),
            consistency=request.consistency,
            target_instance_assurance=request.target_instance_assurance,
            consistency_limitation=(
                "Native pages are not one database-wide point-in-time snapshot."
            ),
        )

    def pages(self):
        return iter((self._page,))

    @property
    def accounting(self):
        return self._accounting


def _browser_model_catalog(workspace_state) -> RecordSnapshot:
    fingerprint = TargetFingerprint(
        target_hash=target_identity_hash(
            connection_mode=workspace_state.odoo_connection_mode.value,
            base_url=workspace_state.odoo_base_url,
            database=workspace_state.odoo_database,
        ),
        connection_mode=workspace_state.odoo_connection_mode.value,
        database=workspace_state.odoo_database,
        odoo_version="19.0",
        snapshot_timestamp="2026-07-30T12:00:00Z",
        module_versions={"base": "19.0.1.0"},
    )
    fields = (
        "name",
        "model",
        "abstract",
        "transient",
        "modules",
        "state",
    )
    return RecordSnapshot(
        fingerprint=fingerprint,
        records={
            "ir.model": (
                TargetRecord(
                    model="ir.model",
                    odoo_id=1,
                    values={
                        "name": "Contact",
                        "model": "res.partner",
                        "abstract": False,
                        "transient": False,
                        "modules": "base, contacts",
                        "state": "base",
                    },
                ),
                TargetRecord(
                    model="ir.model",
                    odoo_id=2,
                    values={
                        "name": "Product",
                        "model": "product.template",
                        "abstract": False,
                        "transient": False,
                        "modules": "product, stock",
                        "state": "base",
                    },
                ),
                TargetRecord(
                    model="ir.model",
                    odoo_id=3,
                    values={
                        "name": "Company",
                        "model": "res.company",
                        "abstract": False,
                        "transient": False,
                        "modules": "base",
                        "state": "base",
                    },
                ),
            )
        },
        requested_fields={"ir.model": fields},
        complete=True,
    )


def _workbook_bytes() -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Products"
    worksheet.append(["Code", "Name"])
    worksheet.append(["P001", "Example product"])
    worksheet.add_table(Table(displayName="ProductTable", ref="A1:B2"))
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()

