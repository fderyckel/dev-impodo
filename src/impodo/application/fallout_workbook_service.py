"""Resolve authorized load fallout to frozen source cells and annotate a copy."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Protocol

from openpyxl.utils import get_column_letter, range_boundaries

from impodo.adapters.artifacts.fallout_workbook import (
    FalloutWorkbookCell,
    build_fallout_workbook,
)
from impodo.application.shared.artifacts import GovernedArtifactStores
from impodo.application.workspace.execution.reconciliation import (
    ReconciliationService,
)
from impodo.domain.execution_snapshot import ExecutionSnapshot
from impodo.domain.preparation.staging_contracts import CanonicalStagingRun
from impodo.domain.shared.access import Actor, Capability
from impodo.domain.shared.models import portable_value
from impodo.domain.source_binding import FileSourceBinding
from impodo.domain.workspace.contracts import SourceSelection
from impodo.domain.workspace.errors import WorkspaceError
from impodo.application.workspace.access import WorkspaceAccessService


class FalloutPreflightReader(Protocol):
    def execution_snapshot(
        self,
        workspace_id: str,
        preflight_run_id: str,
    ) -> ExecutionSnapshot: ...


class FalloutStagingReader(Protocol):
    def get_canonical_staging_run(
        self,
        workspace_id: str,
        run_id: str,
        *,
        expected_content_hash: str | None = None,
    ) -> CanonicalStagingRun | None: ...


class FalloutSourceReader(Protocol):
    def get_source_selection(self, workspace_id: str) -> SourceSelection | None: ...

    def get_source_catalogs(self, workspace_id: str) -> tuple[object, ...]: ...


class FalloutWorkspaceReader(Protocol):
    def get(self, workspace_id: str): ...


@dataclass(frozen=True, slots=True)
class GeneratedFalloutWorkbook:
    filename: str
    content: bytes
    affected_cell_count: int


class FalloutWorkbookService:
    """Create an authorized workbook from one internally consistent attempt."""

    def __init__(
        self,
        reconciliation: ReconciliationService,
        preflight: FalloutPreflightReader,
        staging: FalloutStagingReader,
        sources: FalloutSourceReader,
        workspaces: FalloutWorkspaceReader,
        artifacts: GovernedArtifactStores,
        authorization: WorkspaceAccessService,
    ) -> None:
        self._reconciliation = reconciliation
        self._preflight = preflight
        self._staging = staging
        self._sources = sources
        self._workspaces = workspaces
        self._artifacts = artifacts
        self._authorization = authorization

    def generate(
        self,
        workspace_id: str,
        *,
        actor: Actor,
    ) -> GeneratedFalloutWorkbook:
        self._authorization.require(
            actor,
            Capability.PROTECTED_EVIDENCE_READ,
            workspace_id=workspace_id,
        )
        report = self._reconciliation.current(workspace_id)
        detail = self._reconciliation.current_detail(workspace_id, actor=actor)
        if report is None:
            raise WorkspaceError("Verify the completed Odoo load first")
        if detail is None:
            raise WorkspaceError(
                "Re-check Odoo now to create exact source-cell fallout evidence"
            )
        execution_run = self._reconciliation.execution.get_run(
            workspace_id,
            report.execution_run_id,
        )
        if execution_run is None:
            raise WorkspaceError("The saved load journal is unavailable")
        snapshot = self._preflight.execution_snapshot(
            workspace_id,
            execution_run.preflight_run_id,
        )
        staging = self._staging.get_canonical_staging_run(
            workspace_id,
            snapshot.staging_run_id,
            expected_content_hash=snapshot.staging_content_hash,
        )
        selection = self._sources.get_source_selection(workspace_id)
        if staging is None or selection is None:
            raise WorkspaceError("The frozen source lineage is unavailable")
        if staging.source_selection_hash != selection.content_hash:
            raise WorkspaceError("The frozen source selection changed after the load")

        execution_rows = {item.source_trace_id: item for item in snapshot.rows}
        canonical_rows = {item.row_id: item for item in staging.rows}
        datasets = {item.dataset_id: item for item in selection.datasets}
        source_catalogs = {
            item.file_id: item
            for item in self._sources.get_source_catalogs(workspace_id)
        }
        workspace = self._workspaces.get(workspace_id)
        source_files = {item.file_id: item for item in workspace.source_files}
        workbook_cells: list[FalloutWorkbookCell] = []
        used_file_ids: set[str] = set()

        for difference in detail.differences:
            execution_row = execution_rows.get(difference.source_trace_id)
            canonical_row = canonical_rows.get(difference.source_trace_id)
            if execution_row is None or canonical_row is None:
                raise WorkspaceError("A fallout row no longer matches frozen lineage")
            field_sources = canonical_row.lineage.field_sources.get(
                difference.field,
                (),
            )
            resolved = []
            for physical_dataset_id, source_rows in (
                canonical_row.lineage.physical_sources.items()
            ):
                source_dataset = datasets.get(physical_dataset_id)
                if source_dataset is None or not isinstance(
                    source_dataset.source, FileSourceBinding
                ):
                    continue
                columns = {
                    item.stable_key: item for item in source_dataset.columns
                }
                binding = source_dataset.source
                for field_source in field_sources:
                    column = columns.get(field_source)
                    if column is None:
                        continue
                    for source_row in source_rows:
                        column_number = _source_column_number(
                            binding,
                            column.ordinal,
                            source_catalogs.get(binding.file_id),
                        )
                        resolved.append(
                            (
                                binding.file_id,
                                _worksheet_name(
                                    binding.table_key,
                                    source_files.get(binding.file_id),
                                ),
                                f"{get_column_letter(column_number)}{source_row}",
                            )
                        )
            if not resolved:
                primary = datasets.get(canonical_row.lineage.physical_dataset_id)
                if primary is not None and isinstance(
                    primary.source, FileSourceBinding
                ):
                    used_file_ids.add(primary.source.file_id)
                    source_file = source_files.get(primary.source.file_id)
                    source_name = source_file.display_name if source_file else "Source"
                else:
                    source_name = "Generated value"
                workbook_cells.append(
                    _workbook_cell(
                        difference,
                        business_key=_business_key(execution_row.business_identity),
                        source_file=source_name,
                        worksheet="",
                        coordinate="",
                    )
                )
                continue
            for file_id, worksheet, coordinate in sorted(set(resolved)):
                source_file = source_files.get(file_id)
                if source_file is None:
                    raise WorkspaceError("An accepted source file is unavailable")
                used_file_ids.add(file_id)
                workbook_cells.append(
                    _workbook_cell(
                        difference,
                        business_key=_business_key(execution_row.business_identity),
                        source_file=source_file.display_name,
                        worksheet=worksheet,
                        coordinate=coordinate,
                    )
                )

        if len(used_file_ids) != 1:
            raise WorkspaceError(
                "This verification spans multiple source files. Download each "
                "file-specific workbook from Support details."
            )
        file_id = next(iter(used_file_ids))
        source_file = source_files.get(file_id)
        if source_file is None:
            raise WorkspaceError("The accepted source workbook is unavailable")
        binding = next(
            (
                item.source
                for item in selection.datasets
                if isinstance(item.source, FileSourceBinding)
                and item.source.file_id == file_id
            ),
            None,
        )
        if binding is None or _canonical_hash(source_file.sha256) != binding.source_sha256:
            raise WorkspaceError("The accepted source workbook failed verification")
        with self._artifacts.materialize_source(
            selection.data_version_id,
            source_file.stored_name,
        ) as source_path:
            content = build_fallout_workbook(
                Path(source_path),
                source_display_name=source_file.display_name,
                source_header_row=binding.header_row,
                source_encoding=binding.encoding,
                source_delimiter=binding.delimiter,
                report=report,
                cells=workbook_cells,
            )
        return GeneratedFalloutWorkbook(
            filename=f"impodo-load-fallout-{report.reconciliation_id}.xlsx",
            content=content,
            affected_cell_count=sum(bool(item.coordinate) for item in workbook_cells),
        )

def _workbook_cell(
    difference,
    *,
    business_key: str,
    source_file: str,
    worksheet: str,
    coordinate: str,
) -> FalloutWorkbookCell:
    return FalloutWorkbookCell(
        source_file=source_file,
        worksheet=worksheet,
        coordinate=coordinate,
        business_key=business_key,
        odoo_record=f"{difference.target_model} · {difference.odoo_id}",
        field=difference.field.replace("_", " ").title(),
        source_value=None,
        prepared_value=difference.expected_value,
        odoo_value=difference.observed_value,
        result="Different after Odoo read-back",
        reason_code=difference.reason_code,
        recommended_action=_recommended_action(difference.reason_code),
    )


def _recommended_action(reason_code: str) -> str:
    if reason_code == "TARGET_NUMERIC_PRECISION_LOSS":
        return (
            "Change Odoo precision or approve an explicit rounding or unit "
            "conversion rule. Do not reload the records."
        )
    if reason_code == "HTML_CONTENT_DIFFERENT":
        return (
            "Review the meaningful HTML difference or approve an explicit "
            "normalization rule. Do not reload the records."
        )
    return "Review the source, mapping, and Odoo value. Do not reload the records."


def _worksheet_name(table_key: str, source_file) -> str:
    if table_key.startswith("sheet:"):
        return table_key.removeprefix("sheet:")
    if table_key.startswith("table:"):
        return table_key.removeprefix("table:").rsplit(":", 1)[0]
    if table_key == "csv":
        return _safe_sheet_title(Path(source_file.display_name).stem if source_file else "Source")
    return ""


def _source_column_number(binding: FileSourceBinding, ordinal: int, catalog) -> int:
    """Resolve a selected-table-relative ordinal to an Excel column number."""

    if not binding.table_key.startswith("table:"):
        return ordinal
    table = next(
        (
            item
            for item in getattr(catalog, "tables", ())
            if item.table_key == binding.table_key
        ),
        None,
    )
    if table is None or len(getattr(table, "named_tables", ())) != 1:
        raise WorkspaceError("The frozen Excel table bounds are unavailable")
    minimum_column, _minimum_row, maximum_column, _maximum_row = range_boundaries(
        table.named_tables[0].cell_range
    )
    column_number = minimum_column + ordinal - 1
    if ordinal < 1 or column_number > maximum_column:
        raise WorkspaceError("The frozen Excel table column is invalid")
    return column_number


def _safe_sheet_title(value: str) -> str:
    cleaned = "".join("_" if item in "[]:*?/\\" else item for item in value)
    return (cleaned or "Source")[:31]


def _business_key(value: tuple[object, ...]) -> str:
    portable = portable_value(value)
    if isinstance(portable, list):
        return " · ".join(str(item) for item in portable)
    return str(portable)


def _canonical_hash(value: str) -> str:
    return value if value.startswith("sha256:") else f"sha256:{value}"


__all__ = ["FalloutWorkbookService", "GeneratedFalloutWorkbook"]
