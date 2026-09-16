"""Create a source-preserving Excel view of exact load fallout cells."""

from __future__ import annotations

import csv
from dataclasses import replace
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO, TextIOWrapper
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from impodo.domain.reconciliation import ReconciliationRowStatus, ReconciliationRun
from impodo.application.fallout_workbook_port import FalloutWorkbookCell


FALLOUT_FILL = "FCE8E7"
FALLOUT_TEXT = "9F2F2F"
WARNING_FILL = "FFF5DF"
WARNING_TEXT = "7D4F00"
VERIFIED_FILL = "EDF7EF"
VERIFIED_TEXT = "4D7C5B"


def build_fallout_workbook(
    source_path: Path,
    *,
    source_display_name: str,
    source_header_row: int,
    source_encoding: str | None = None,
    source_delimiter: str | None = None,
    report: ReconciliationRun,
    cells: Iterable[FalloutWorkbookCell],
) -> bytes:
    """Return an annotated copy while leaving the accepted source unchanged."""

    issues = tuple(cells)
    if source_path.suffix.casefold() == ".xlsx":
        workbook = load_workbook(source_path, data_only=False, read_only=False)
    elif source_path.suffix.casefold() == ".csv":
        workbook = _csv_workbook(
            source_path,
            source_display_name=source_display_name,
            header_row=source_header_row,
            encoding=source_encoding,
            delimiter=source_delimiter,
        )
    else:
        raise ValueError("Fallout workbook requires an accepted CSV or XLSX source")

    issues = tuple(
        replace(
            issue,
            source_value=(
                workbook[issue.worksheet][issue.coordinate].value
                if issue.coordinate and issue.worksheet in workbook.sheetnames
                else issue.source_value
            ),
        )
        for issue in issues
    )
    title = _unique_summary_title(workbook)
    summary = workbook.create_sheet(title, 0)
    _write_summary(summary, report, issues)
    for issue in issues:
        if not issue.coordinate or not issue.worksheet:
            continue
        if issue.worksheet not in workbook.sheetnames:
            raise ValueError("A fallout cell names an unknown worksheet")
        worksheet = workbook[issue.worksheet]
        cell = worksheet[issue.coordinate]
        cell.fill = PatternFill("solid", fgColor=FALLOUT_FILL)
        cell.font = _merged_font(cell.font, color=FALLOUT_TEXT)
        red = Side(style="medium", color=FALLOUT_TEXT)
        cell.border = Border(left=red, right=red, top=red, bottom=red)
        cell.comment = Comment(
            "\n".join(
                (
                    f"Odoo field: {issue.field}",
                    f"Prepared value: {_display(issue.prepared_value)}",
                    f"Observed Odoo value: {_display(issue.odoo_value)}",
                    f"Odoo record: {issue.odoo_record}",
                    f"Next action: {issue.recommended_action}",
                )
            ),
            "Impodo",
        )

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _write_summary(worksheet, report, issues: tuple[FalloutWorkbookCell, ...]) -> None:
    written = sum(item.odoo_id is not None for item in report.rows)
    different = sum(
        item.status is ReconciliationRowStatus.DIFFERENT for item in report.rows
    )
    missing = sum(
        item.status is ReconciliationRowStatus.MISSING for item in report.rows
    )
    unwritten = sum(
        item.status
        in {ReconciliationRowStatus.NOT_WRITTEN, ReconciliationRowStatus.NOT_APPLIED}
        for item in report.rows
    )
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A8"
    worksheet["A1"] = "Impodo load fallout"
    worksheet["A1"].font = Font(bold=True, size=18, color="494D46")
    worksheet["A2"] = (
        f"Odoo accepted {written:,} written record(s), but {different:,} "
        "record(s) have a different final value. A read-back difference is "
        "not a rejected row."
    )
    worksheet["A3"] = (
        f"{report.verified_count:,} verified · {different:,} different · "
        f"{missing:,} missing · {unwritten:,} unwritten · "
        f"{report.unknown_count:,} unknown · "
        f"{len(issues):,} source-linked cell issue(s)"
    )
    worksheet["A4"] = (
        "Do not reload these records. Review the highlighted cells and the "
        "recommended action below."
    )
    worksheet["A5"] = f"Verification: {report.reconciliation_id}"
    headers = (
        "Source file",
        "Worksheet",
        "Cell",
        "Business key",
        "Odoo record",
        "Business field",
        "Source value",
        "Prepared value",
        "Odoo value",
        "Verification result",
        "Reason",
        "Recommended action",
    )
    for column, header in enumerate(headers, 1):
        cell = worksheet.cell(7, column, header)
        cell.fill = PatternFill("solid", fgColor="494D46")
        cell.font = Font(bold=True, color="FFFFFF")
    for row_number, issue in enumerate(issues, 8):
        values = (
            issue.source_file,
            issue.worksheet or "Generated by a rule — no source cell",
            issue.coordinate or "—",
            issue.business_key,
            issue.odoo_record,
            issue.field,
            issue.source_value,
            issue.prepared_value,
            issue.odoo_value,
            issue.result,
            issue.reason_code.replace("_", " ").title(),
            issue.recommended_action,
        )
        for column, value in enumerate(values, 1):
            target = worksheet.cell(row_number, column, _safe_generated_value(value))
            target.fill = PatternFill("solid", fgColor=FALLOUT_FILL)
            target.font = Font(color=FALLOUT_TEXT)
        if issue.coordinate and issue.worksheet:
            link = worksheet.cell(row_number, 3)
            escaped_sheet = issue.worksheet.replace("'", "''")
            link.hyperlink = f"#'{escaped_sheet}'!{issue.coordinate}"
            link.style = "Hyperlink"
    widths = (24, 24, 12, 28, 24, 24, 24, 24, 24, 22, 30, 58)
    for index, width in enumerate(widths, 1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    worksheet.auto_filter.ref = f"A7:L{max(7, 7 + len(issues))}"


def _csv_workbook(
    source_path: Path,
    *,
    source_display_name: str,
    header_row: int,
    encoding: str | None,
    delimiter: str | None,
):
    del header_row
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = _safe_sheet_title(Path(source_display_name).stem or "Source")
    with source_path.open("rb") as raw:
        text = TextIOWrapper(raw, encoding=encoding or "utf-8-sig", newline="")
        for row in csv.reader(text, delimiter=delimiter or ","):
            worksheet.append(row)
    return workbook


def _unique_summary_title(workbook) -> str:
    base = "Impodo Fallout"
    if base not in workbook.sheetnames:
        return base
    index = 2
    while f"{base} {index}" in workbook.sheetnames:
        index += 1
    return f"{base} {index}"


def _safe_sheet_title(value: str) -> str:
    cleaned = "".join("_" if item in "[]:*?/\\" else item for item in value)
    return (cleaned or "Source")[:31]


def _safe_generated_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (bool, int, float, date, datetime)):
        return value
    text = _display(value)
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def _display(value: Any) -> str:
    if value is None or value is False:
        return "(blank)"
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, (tuple, list, set)):
        return ", ".join(_display(item) for item in value)
    return str(value)


def _merged_font(font, *, color: str) -> Font:
    return Font(
        name=font.name,
        size=font.size,
        bold=font.bold,
        italic=font.italic,
        vertAlign=font.vertAlign,
        underline=font.underline,
        strike=font.strike,
        color=color,
    )


__all__ = ["FalloutWorkbookCell", "build_fallout_workbook"]
